import argparse
import csv
import json
import os
import platform
import socket
import sys
import threading
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
import tkinter as tk

from register_client import (
    collect_device_diagnostics,
    collect_inventory,
    merge_diagnostics,
    pc_auditor_optical_read_test,
    pc_auditor_usb_detection_test,
    powershell_value,
    read_sample_from_path,
    status_item,
    summarize_status,
    test_path_candidates,
)


APP_TITLE = "PC Audit Tool"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORT_DIR = os.path.join(BASE_DIR, "audit_reports")
SETTINGS_PATH = os.path.join(BASE_DIR, "pc_audit_settings.json")
DEFAULT_SETTINGS = {
    "expected_usb_devices": "",
    "dvd_test_path": "",
    "wifi_ssid": "",
}


def now_stamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def default_report_base():
    return f"pc_audit_{socket.gethostname()}_{now_stamp()}"


def ensure_report_dir():
    os.makedirs(REPORT_DIR, exist_ok=True)
    return REPORT_DIR


def load_settings():
    settings = dict(DEFAULT_SETTINGS)
    try:
        with open(SETTINGS_PATH, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        if isinstance(data, dict):
            settings.update({key: str(data.get(key, settings[key])) for key in settings})
    except (OSError, json.JSONDecodeError):
        pass
    return settings


def save_settings(settings):
    with open(SETTINGS_PATH, "w", encoding="utf-8") as handle:
        json.dump(settings, handle, indent=2, ensure_ascii=False)


def first_line(value):
    for line in str(value or "").splitlines():
        line = line.strip()
        if line:
            return line
    return ""


def dvd_read_test(dvd_test_path=""):
    path = str(dvd_test_path or "").strip()
    if not path:
        candidates = test_path_candidates("optical")
        path = candidates[0] if candidates else ""
    if not path:
        return status_item("Test DVD read", "UNKNOWN", "No DVD/optical drive path detected or configured")
    return read_sample_from_path(path, "Test DVD read")


def wifi_connection_test(expected_ssid=""):
    expected_ssid = str(expected_ssid or "").strip()
    system = platform.system().lower()

    if system.startswith("win"):
        adapter = first_line(
            powershell_value(
                "Get-CimInstance Win32_NetworkAdapter | "
                "Where-Object { $_.PhysicalAdapter -and $_.NetConnectionID -match 'Wi-Fi|Wireless|WLAN' } | "
                "Select-Object -First 1 | ForEach-Object { $_.Name + '|' + $_.NetConnectionStatus }"
            )
        )
        if not adapter:
            return status_item("Test Wi-Fi connection", "UNKNOWN", "No Wi-Fi adapter found")

        profile = first_line(
            powershell_value(
                "(netsh wlan show interfaces) | Select-String '^\\s*SSID\\s*:' | "
                "Select-Object -First 1 | ForEach-Object { ($_.ToString() -split ':',2)[1].Trim() }"
            )
        )
        connected = "|2" in adapter or bool(profile)
        details = adapter
        if profile:
            details += f" | SSID={profile}"

        if expected_ssid:
            details += f" | expected SSID={expected_ssid}"
            status = "OK" if connected and profile.lower() == expected_ssid.lower() else "FAIL"
        else:
            status = "OK" if connected else "WARN"
        return status_item("Test Wi-Fi connection", status, details)

    if system == "linux":
        ssid = ""
        try:
            import subprocess

            output = subprocess.run(
                ["iwgetid", "-r"],
                capture_output=True,
                text=True,
                errors="replace",
                timeout=5,
            )
            ssid = first_line(output.stdout)
        except Exception:
            ssid = ""
        if expected_ssid:
            status = "OK" if ssid.lower() == expected_ssid.lower() else "FAIL"
            return status_item("Test Wi-Fi connection", status, f"SSID={ssid or 'not connected'} | expected SSID={expected_ssid}")
        return status_item("Test Wi-Fi connection", "OK" if ssid else "UNKNOWN", f"SSID={ssid or 'not connected'}")

    return status_item("Test Wi-Fi connection", "UNKNOWN", f"Unsupported OS: {platform.system()}")


def collect_audit(
    include_extended=False,
    include_usb_test=True,
    expected_usb_devices=-1,
    include_dvd_test=True,
    dvd_test_path="",
    include_wifi_test=True,
    wifi_ssid="",
):
    inventory = collect_inventory()
    diagnostics = collect_device_diagnostics()

    extra_items = []
    if include_usb_test:
        try:
            extra_items.append(
                pc_auditor_usb_detection_test(
                    expected_usb_devices=expected_usb_devices,
                    non_interactive=True,
                )
            )
        except Exception as exc:
            extra_items.append({"name": "Test USB device detection", "status": "WARN", "details": str(exc)})

    if include_dvd_test:
        try:
            extra_items.append(dvd_read_test(dvd_test_path))
        except Exception as exc:
            extra_items.append({"name": "Test DVD read", "status": "WARN", "details": str(exc)})

    if include_wifi_test:
        try:
            extra_items.append(wifi_connection_test(wifi_ssid))
        except Exception as exc:
            extra_items.append({"name": "Test Wi-Fi connection", "status": "WARN", "details": str(exc)})

    if extra_items:
        diagnostics = merge_diagnostics(diagnostics, {"items": extra_items, "summary": summarize_status(extra_items)})

    if include_extended:
        try:
            from register_client import collect_auditor_style_tests

            class Args:
                server = "http://127.0.0.1:8765"
                test_attempts = 1
                timeout = 4
                memory_test_mb = 64
                expected_usb_devices = -1
                non_interactive_tests = True
                optical_test_path = ""
                optical_path = ""
                removable_path = ""
                usb_test_path = ""

            Args.expected_usb_devices = expected_usb_devices
            Args.optical_test_path = dvd_test_path

            diagnostics = merge_diagnostics(diagnostics, collect_auditor_style_tests(Args()))
        except Exception as exc:
            items = diagnostics.get("items", []) if isinstance(diagnostics, dict) else []
            items.append({"name": "Extended tests", "status": "WARN", "details": str(exc)})
            diagnostics = {"items": items, "summary": summarize_items(items)}

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "computer_name": socket.gethostname(),
        "inventory": inventory,
        "diagnostics": diagnostics,
    }


def summarize_items(items):
    summary = {"OK": 0, "WARN": 0, "FAIL": 0, "UNKNOWN": 0}
    for item in items:
        status = str(item.get("status", "UNKNOWN")).upper()
        summary[status] = summary.get(status, 0) + 1
    return summary


def flatten_audit(report):
    rows = []
    rows.append(("Generated At", report.get("generated_at", "")))
    rows.append(("Computer Name", report.get("computer_name", "")))
    rows.append(("", ""))
    rows.append(("Inventory", ""))
    for key, value in report.get("inventory", {}).items():
        rows.append((key, value))
    rows.append(("", ""))
    rows.append(("Diagnostics", ""))
    diagnostics = report.get("diagnostics") or {}
    for item in diagnostics.get("items", []):
        name = item.get("name", "")
        status = item.get("status", "")
        details = item.get("details", "")
        rows.append((name, f"{status} - {details}".strip(" -")))
    return rows


def report_to_text(report):
    lines = [APP_TITLE, "=" * len(APP_TITLE), ""]
    for key, value in flatten_audit(report):
        if not key and not value:
            lines.append("")
        elif value:
            lines.append(f"{key}: {value}")
        else:
            lines.append(str(key))
    return "\n".join(lines).rstrip() + "\n"


def save_json(report, path):
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=2, ensure_ascii=False)


def save_txt(report, path):
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(report_to_text(report))


def save_csv(report, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Field", "Value"])
        writer.writerows(flatten_audit(report))


def save_xlsx(report, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "PC Audit"
    ws.append(["Field", "Value"])
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    for row in flatten_audit(report):
        ws.append(list(row))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    wb.save(path)


def save_bundle(report):
    folder = ensure_report_dir()
    base = os.path.join(folder, default_report_base())
    paths = {
        "json": base + ".json",
        "txt": base + ".txt",
        "csv": base + ".csv",
    }
    save_json(report, paths["json"])
    save_txt(report, paths["txt"])
    save_csv(report, paths["csv"])
    try:
        paths["xlsx"] = base + ".xlsx"
        save_xlsx(report, paths["xlsx"])
    except Exception:
        paths.pop("xlsx", None)
    return paths


class AuditApp:
    def __init__(self, root):
        self.root = root
        self.report = None
        self.settings = load_settings()
        self.include_extended = tk.BooleanVar(value=False)
        self.include_usb_test = tk.BooleanVar(value=True)
        self.include_dvd_test = tk.BooleanVar(value=True)
        self.include_wifi_test = tk.BooleanVar(value=True)
        self.expected_usb_devices = tk.StringVar(value=self.settings.get("expected_usb_devices", ""))
        self.dvd_test_path = tk.StringVar(value=self.settings.get("dvd_test_path", ""))
        self.wifi_ssid = tk.StringVar(value=self.settings.get("wifi_ssid", ""))
        self.status_var = tk.StringVar(value="Ready")

        root.title(APP_TITLE)
        root.geometry("980x640")
        root.minsize(820, 520)

        self.build_ui()

    def build_ui(self):
        style = ttk.Style()
        style.configure("Header.TLabel", font=("Segoe UI", 15, "bold"))
        style.configure("Status.TLabel", foreground="#4b5563")

        outer = ttk.Frame(self.root, padding=16)
        outer.pack(fill="both", expand=True)

        header = ttk.Frame(outer)
        header.pack(fill="x")
        ttk.Label(header, text=APP_TITLE, style="Header.TLabel").pack(side="left")
        ttk.Label(header, textvariable=self.status_var, style="Status.TLabel").pack(side="right")

        notebook = ttk.Notebook(outer)
        notebook.pack(fill="both", expand=True, pady=(14, 0))

        audit_tab = ttk.Frame(notebook, padding=(0, 12, 0, 0))
        settings_tab = ttk.Frame(notebook, padding=16)
        notebook.add(audit_tab, text="Audit")
        notebook.add(settings_tab, text="Settings")

        toolbar = ttk.Frame(audit_tab)
        toolbar.pack(fill="x", pady=(14, 10))
        self.run_button = ttk.Button(toolbar, text="Run Audit", command=self.run_audit)
        self.run_button.pack(side="left", padx=(0, 8))
        ttk.Checkbutton(toolbar, text="USB test", variable=self.include_usb_test).pack(side="left", padx=(0, 8))
        ttk.Checkbutton(toolbar, text="DVD test", variable=self.include_dvd_test).pack(side="left", padx=(0, 8))
        ttk.Checkbutton(toolbar, text="WiFi test", variable=self.include_wifi_test).pack(side="left", padx=(0, 12))
        ttk.Checkbutton(toolbar, text="Include safe extended tests", variable=self.include_extended).pack(side="left")
        ttk.Button(toolbar, text="Save Reports", command=self.save_reports).pack(side="right")
        ttk.Button(toolbar, text="Export As...", command=self.export_as).pack(side="right", padx=(0, 8))

        body = ttk.PanedWindow(audit_tab, orient="horizontal")
        body.pack(fill="both", expand=True)

        inventory_frame = ttk.LabelFrame(body, text="Inventory", padding=8)
        diagnostic_frame = ttk.LabelFrame(body, text="Diagnostics", padding=8)
        body.add(inventory_frame, weight=1)
        body.add(diagnostic_frame, weight=1)

        self.inventory_tree = ttk.Treeview(inventory_frame, columns=("value",), show="tree headings")
        self.inventory_tree.heading("#0", text="Field")
        self.inventory_tree.heading("value", text="Value")
        self.inventory_tree.column("#0", width=180, anchor="w")
        self.inventory_tree.column("value", width=320, anchor="w")
        inv_scroll = ttk.Scrollbar(inventory_frame, orient="vertical", command=self.inventory_tree.yview)
        self.inventory_tree.configure(yscrollcommand=inv_scroll.set)
        self.inventory_tree.pack(side="left", fill="both", expand=True)
        inv_scroll.pack(side="right", fill="y")

        self.diagnostics_tree = ttk.Treeview(
            diagnostic_frame,
            columns=("status", "details"),
            show="tree headings",
        )
        self.diagnostics_tree.heading("#0", text="Check")
        self.diagnostics_tree.heading("status", text="Status")
        self.diagnostics_tree.heading("details", text="Details")
        self.diagnostics_tree.column("#0", width=190, anchor="w")
        self.diagnostics_tree.column("status", width=80, anchor="center")
        self.diagnostics_tree.column("details", width=360, anchor="w")
        diag_scroll = ttk.Scrollbar(diagnostic_frame, orient="vertical", command=self.diagnostics_tree.yview)
        self.diagnostics_tree.configure(yscrollcommand=diag_scroll.set)
        self.diagnostics_tree.pack(side="left", fill="both", expand=True)
        diag_scroll.pack(side="right", fill="y")

        self.build_settings_tab(settings_tab)

    def build_settings_tab(self, parent):
        ttk.Label(parent, text="PC Audit Settings", style="Header.TLabel").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 14))

        ttk.Label(parent, text="Expected USB").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=8)
        ttk.Entry(parent, textvariable=self.expected_usb_devices, width=18).grid(row=1, column=1, sticky="w", pady=8)
        ttk.Label(parent, text="Minimum visible USB devices/items. Leave blank to only detect.").grid(row=1, column=2, sticky="w", padx=(12, 0), pady=8)

        ttk.Label(parent, text="DVD path").grid(row=2, column=0, sticky="e", padx=(0, 8), pady=8)
        ttk.Entry(parent, textvariable=self.dvd_test_path, width=28).grid(row=2, column=1, sticky="ew", pady=8)
        ttk.Button(parent, text="Detect DVD", command=self.detect_dvd_path).grid(row=2, column=2, sticky="w", padx=(12, 0), pady=8)

        ttk.Label(parent, text="WiFi SSID").grid(row=3, column=0, sticky="e", padx=(0, 8), pady=8)
        ttk.Entry(parent, textvariable=self.wifi_ssid, width=28).grid(row=3, column=1, sticky="ew", pady=8)
        ttk.Label(parent, text="Expected WiFi network name. Leave blank to only check connection.").grid(row=3, column=2, sticky="w", padx=(12, 0), pady=8)

        ttk.Button(parent, text="Save Settings", command=self.save_tool_settings).grid(row=4, column=1, sticky="w", pady=(14, 0))
        ttk.Button(parent, text="Reload Settings", command=self.reload_tool_settings).grid(row=4, column=2, sticky="w", padx=(12, 0), pady=(14, 0))

        parent.columnconfigure(1, weight=1)

    def run_audit(self):
        self.status_var.set("Running audit...")
        self.run_button.configure(state="disabled")
        self.clear_trees()
        include_extended = self.include_extended.get()
        include_usb_test = self.include_usb_test.get()
        include_dvd_test = self.include_dvd_test.get()
        include_wifi_test = self.include_wifi_test.get()
        expected_usb_devices = self.parse_expected_usb_devices()
        if expected_usb_devices is None:
            self.run_button.configure(state="normal")
            return
        dvd_test_path = self.dvd_test_path.get().strip()
        wifi_ssid = self.wifi_ssid.get().strip()

        def worker():
            try:
                report = collect_audit(
                    include_extended=include_extended,
                    include_usb_test=include_usb_test,
                    expected_usb_devices=expected_usb_devices,
                    include_dvd_test=include_dvd_test,
                    dvd_test_path=dvd_test_path,
                    include_wifi_test=include_wifi_test,
                    wifi_ssid=wifi_ssid,
                )
                self.root.after(0, lambda: self.show_report(report))
            except Exception as exc:
                self.root.after(0, lambda: self.show_error(exc))

        threading.Thread(target=worker, daemon=True).start()

    def parse_expected_usb_devices(self):
        value = self.expected_usb_devices.get().strip()
        if not value:
            return -1
        try:
            count = int(value)
        except ValueError:
            messagebox.showerror(APP_TITLE, "Expected USB must be a number")
            return None
        if count < 0:
            messagebox.showerror(APP_TITLE, "Expected USB must be 0 or greater")
            return None
        return count

    def settings_payload(self):
        return {
            "expected_usb_devices": self.expected_usb_devices.get().strip(),
            "dvd_test_path": self.dvd_test_path.get().strip(),
            "wifi_ssid": self.wifi_ssid.get().strip(),
        }

    def save_tool_settings(self):
        if self.parse_expected_usb_devices() is None:
            return
        try:
            save_settings(self.settings_payload())
        except OSError as exc:
            messagebox.showerror(APP_TITLE, f"Could not save settings: {exc}")
            return
        self.status_var.set(f"Settings saved to {SETTINGS_PATH}")
        messagebox.showinfo(APP_TITLE, "Settings saved")

    def reload_tool_settings(self):
        self.settings = load_settings()
        self.expected_usb_devices.set(self.settings.get("expected_usb_devices", ""))
        self.dvd_test_path.set(self.settings.get("dvd_test_path", ""))
        self.wifi_ssid.set(self.settings.get("wifi_ssid", ""))
        self.status_var.set("Settings reloaded")

    def detect_dvd_path(self):
        candidates = test_path_candidates("optical")
        if not candidates:
            messagebox.showwarning(APP_TITLE, "No DVD/optical drive path detected")
            return
        self.dvd_test_path.set(candidates[0])
        self.status_var.set(f"DVD path detected: {candidates[0]}")

    def show_report(self, report):
        self.report = report
        self.clear_trees()
        for key, value in report.get("inventory", {}).items():
            self.inventory_tree.insert("", "end", text=key, values=(value,))
        diagnostics = report.get("diagnostics") or {}
        for item in diagnostics.get("items", []):
            self.diagnostics_tree.insert(
                "",
                "end",
                text=item.get("name", ""),
                values=(item.get("status", ""), item.get("details", "")),
            )
        summary = diagnostics.get("summary", {})
        self.status_var.set(
            "Done: "
            + ", ".join(f"{key} {value}" for key, value in summary.items() if value)
            if summary
            else "Done"
        )
        self.run_button.configure(state="normal")

    def show_error(self, exc):
        self.status_var.set("Audit failed")
        self.run_button.configure(state="normal")
        messagebox.showerror(APP_TITLE, str(exc))

    def clear_trees(self):
        for tree in (self.inventory_tree, self.diagnostics_tree):
            for item in tree.get_children():
                tree.delete(item)

    def require_report(self):
        if self.report:
            return True
        messagebox.showwarning(APP_TITLE, "Run an audit first.")
        return False

    def save_reports(self):
        if not self.require_report():
            return
        try:
            paths = save_bundle(self.report)
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not save reports: {exc}")
            return
        self.status_var.set(f"Saved to {REPORT_DIR}")
        messagebox.showinfo(APP_TITLE, "Saved:\n" + "\n".join(paths.values()))

    def export_as(self):
        if not self.require_report():
            return
        path = filedialog.asksaveasfilename(
            title="Export PC audit",
            defaultextension=".xlsx",
            initialfile=default_report_base() + ".xlsx",
            filetypes=[
                ("Excel workbook", "*.xlsx"),
                ("JSON", "*.json"),
                ("Text", "*.txt"),
                ("CSV", "*.csv"),
            ],
        )
        if not path:
            return
        try:
            ext = os.path.splitext(path)[1].lower()
            if ext == ".json":
                save_json(self.report, path)
            elif ext == ".txt":
                save_txt(self.report, path)
            elif ext == ".csv":
                save_csv(self.report, path)
            else:
                save_xlsx(self.report, path)
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not export report: {exc}")
            return
        self.status_var.set(f"Exported {path}")


def run_gui():
    root = tk.Tk()
    AuditApp(root)
    root.mainloop()


def parse_args():
    parser = argparse.ArgumentParser(description="Collect local PC inventory and diagnostics.")
    parser.add_argument("--cli", action="store_true", help="Run without the GUI and print a text report.")
    parser.add_argument("--save", action="store_true", help="Save JSON/TXT/CSV reports to audit_reports.")
    parser.add_argument("--json", action="store_true", help="Print JSON instead of text when using --cli.")
    parser.add_argument("--extended", action="store_true", help="Include safe non-interactive extended checks.")
    parser.add_argument("--no-usb-test", action="store_true", help="Do not include the USB detection test.")
    parser.add_argument("--no-dvd-test", action="store_true", help="Do not include the DVD/optical read test.")
    parser.add_argument("--dvd-test-path", default=None, help="Drive/path for DVD media test, e.g. D:")
    parser.add_argument("--no-wifi-test", action="store_true", help="Do not include the WiFi connection test.")
    parser.add_argument("--wifi-ssid", default=None, help="Expected WiFi SSID for the WiFi connection test.")
    parser.add_argument(
        "--expected-usb-devices",
        type=int,
        default=None,
        help="Expected minimum visible USB devices/items for the USB test.",
    )
    return parser.parse_args()


def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass

    args = parse_args()
    if not args.cli:
        run_gui()
        return 0

    settings = load_settings()
    expected_usb_devices = args.expected_usb_devices
    if expected_usb_devices is None:
        expected_text = settings.get("expected_usb_devices", "").strip()
        expected_usb_devices = int(expected_text) if expected_text.isdigit() else -1
    dvd_test_path = args.dvd_test_path if args.dvd_test_path is not None else settings.get("dvd_test_path", "")
    wifi_ssid = args.wifi_ssid if args.wifi_ssid is not None else settings.get("wifi_ssid", "")

    report = collect_audit(
        include_extended=args.extended,
        include_usb_test=not args.no_usb_test,
        expected_usb_devices=expected_usb_devices,
        include_dvd_test=not args.no_dvd_test,
        dvd_test_path=dvd_test_path,
        include_wifi_test=not args.no_wifi_test,
        wifi_ssid=wifi_ssid,
    )
    if args.save:
        paths = save_bundle(report)
        print("Saved reports:")
        for path in paths.values():
            print(path)
    if args.json:
        print(json.dumps(report, indent=2, ensure_ascii=False))
    else:
        print(report_to_text(report))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
