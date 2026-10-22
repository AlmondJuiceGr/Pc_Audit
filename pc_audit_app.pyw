import hashlib
import getpass
import html
import json
import locale
import os
import platform
import re
import subprocess
import shutil
import sqlite3
import socket
import threading
import time
import traceback
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
import tkinter as tk
from urllib import error as urlerror
from urllib import parse as urlparse
from urllib import request as urlrequest

from openpyxl import Workbook

try:
    import pyautogui
except ImportError:
    pyautogui = None

# =========================
# CONFIG
# =========================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_NAME = os.path.join(BASE_DIR, "app.db")
INVENTORY_DEBUG_LOG = os.path.join(BASE_DIR, "inventory_debug.log")

APP_TITLE = "PC Audit Dashboard"
APP_VERSION = "3.0.0"
LEGACY_BRAND = "Enter" + "prise"

# Professional color palette
BG = "#0f172a"
SURFACE = "#111827"
CARD = "#1f2937"
BORDER = "#334155"
TEXT = "#e5e7eb"
MUTED = "#94a3b8"
PRIMARY = "#2563eb"
PRIMARY_DARK = "#1d4ed8"
SUCCESS = "#16a34a"
DANGER = "#dc2626"
WARNING = "#f59e0b"
ENTRY_BG = "#0b1220"

DEFAULT_SETTINGS = {
    "system_name": APP_TITLE,
    "default_export_folder": BASE_DIR,
    "keyboard_delay": "80",
    "auto_refresh": "1",
    "network_server_url": "http://127.0.0.1:8765",
    "workstation_name": os.environ.get("COMPUTERNAME", "CLIENT-1"),
    "language": "el",
}


def load_saved_language():
    if not os.path.exists(DB_NAME):
        return DEFAULT_SETTINGS["language"]
    try:
        conn = sqlite3.connect(DB_NAME)
        c = conn.cursor()
        c.execute("SELECT value FROM app_settings WHERE key=?", ("language",))
        row = c.fetchone()
        conn.close()
    except sqlite3.Error:
        return DEFAULT_SETTINGS["language"]
    return row[0] if row and row[0] in {"en", "el"} else DEFAULT_SETTINGS["language"]


UI_LANGUAGE = load_saved_language()


def tr(en, el):
    return el if UI_LANGUAGE == "el" else en

AUDITOR_FIELDS = [
    "Computer name",
    "Current user",
    "OS",
    "OS version",
    "OS architecture",
    "Domain / Workgroup",
    "Product type",
    "Manufacturer",
    "Model",
    "Chassis type",
    "Part number",
    "Serial number",
    "Motherboard",
    "BIOS version",
    "Secure Boot",
    "TPM",
    "Processor",
    "RAM size",
    "RAM slots used",
    "Disk type",
    "Disk size",
    "Disk details",
    "Volumes",
    "Optical drive",
    "Video model",
    "LCD size",
    "LCD resolution",
    "LCD type",
    "Keyboard layout",
    "Network adapters",
    "Battery",
    "OS COA",
    "Grade",
]

AUDITOR_SAMPLE_VALUES = {
    "Computer name": "AUDIT-PC-001",
    "Current user": "technician",
    "OS": "Windows 11 Pro",
    "OS version": "10.0.22631 build 22631",
    "OS architecture": "64-bit",
    "Domain / Workgroup": "Workgroup: WORKGROUP",
    "Product type": "LAPTOP",
    "Manufacturer": "HP",
    "Model": "ELITEBOOK 8540W",
    "Chassis type": "LAPTOP",
    "Part number": "NU515AV",
    "Serial number": "CND0511LHM",
    "Motherboard": "HP 1521",
    "BIOS version": "F.60",
    "Secure Boot": "True",
    "TPM": "Present=True; Ready=True; Enabled=True",
    "Processor": "Intel Core i5 M 540 2.53 GHz",
    "RAM size": "2 GB",
    "RAM slots used": "1 / 2",
    "Disk type": "SATA SSD",
    "Disk size": "180 GB",
    "Disk details": "Samsung SSD | SATA | SSD | 180 GB",
    "Volumes": "C: 180 GB free 90 GB",
    "Optical drive": "DVD/RW",
    "Video model": "NVIDIA QUADRO FX 880M",
    "LCD size": "15.5\"",
    "LCD resolution": "1600x900",
    "LCD type": "Unknown",
    "Keyboard layout": "",
    "Network adapters": "Intel Ethernet | MAC 00:11:22:33:44:55 | IP 192.168.1.20",
    "Battery": "Primary Battery | 87% | status 2",
    "OS COA": "",
    "Grade": "",
}

LCD_TYPE_OPTIONS = [
    "",
    "Unknown",
    "IPS LCD",
    "TN LCD",
    "VA LCD",
    "OLED",
    "AMOLED",
    "Mini LED",
    "LCD",
]

OS_COA_OPTIONS = [
    "",
    "Windows 11 Pro",
    "Windows 11 Home",
    "Windows 10 Pro",
    "Windows 10 Home",
    "Windows 8.1 Pro",
    "Windows 8.1 Home",
    "Windows 7 Professional",
    "Windows 7 Home Premium",
    "Windows Vista Business",
    "Windows XP Professional",
    "Linux / No COA",
    "No COA",
]

FIELD_COUNT = len(AUDITOR_FIELDS)
SERIAL_NUMBER_INDEX = AUDITOR_FIELDS.index("Serial number")
LCD_TYPE_INDEX = AUDITOR_FIELDS.index("LCD type")
KEYBOARD_SKIP_FIELDS = {"LCD type"}
RECORD_SELECT_COLUMNS = "id, xenoglossi, timestamp, username, password, dropdown, checkbox, inputs"


# =========================
# DATABASE
# =========================
def get_conn():
    return sqlite3.connect(DB_NAME)


def hash_password(pw):
    return hashlib.sha256(pw.encode("utf-8")).hexdigest()


def remove_legacy_brand_word(value):
    cleaned = str(value or "").replace(LEGACY_BRAND, "").replace(LEGACY_BRAND.lower(), "")
    return " ".join(cleaned.split())


def init_db():
    conn = get_conn()
    c = conn.cursor()

    c.execute(
        """CREATE TABLE IF NOT EXISTS records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            xenoglossi TEXT NOT NULL DEFAULT '',
            timestamp TEXT,
            username TEXT,
            password TEXT,
            dropdown TEXT,
            checkbox INTEGER,
            inputs TEXT
        )"""
    )

    c.execute("PRAGMA table_info(records)")
    record_columns = [row[1] for row in c.fetchall()]
    if "xenoglossi" not in record_columns:
        c.execute("ALTER TABLE records ADD COLUMN xenoglossi TEXT NOT NULL DEFAULT ''")

    c.execute(
        "UPDATE records SET xenoglossi='LEGACY-' || id WHERE xenoglossi IS NULL OR TRIM(xenoglossi)=''"
    )

    c.execute(
        """CREATE TABLE IF NOT EXISTS app_users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'User',
            status TEXT NOT NULL DEFAULT 'Active',
            created_at TEXT,
            last_login TEXT
        )"""
    )

    c.execute(
        """CREATE TABLE IF NOT EXISTS app_settings (
            key TEXT PRIMARY KEY,
            value TEXT
        )"""
    )

    c.execute(
        """CREATE TABLE IF NOT EXISTS app_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            timestamp TEXT,
            username TEXT,
            action TEXT,
            details TEXT
        )"""
    )

    c.execute("SELECT id FROM app_users WHERE username=?", ("admin",))
    if c.fetchone() is None:
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        c.execute(
            "INSERT INTO app_users (username, password, role, status, created_at) VALUES (?, ?, ?, ?, ?)",
            ("admin", hash_password("1234"), "Admin", "Active", now),
        )

    for key, value in DEFAULT_SETTINGS.items():
        c.execute(
            "INSERT OR IGNORE INTO app_settings (key, value) VALUES (?, ?)",
            (key, value),
        )

    c.execute("SELECT value FROM app_settings WHERE key=?", ("system_name",))
    row = c.fetchone()
    if row:
        cleaned_system_name = remove_legacy_brand_word(row[0]) or APP_TITLE
        if cleaned_system_name != row[0]:
            c.execute("UPDATE app_settings SET value=? WHERE key=?", (cleaned_system_name, "system_name"))

    conn.commit()
    conn.close()


def log_action(action, details="", username=None):
    user = username or current_user or "system"
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "INSERT INTO app_logs (timestamp, username, action, details) VALUES (?, ?, ?, ?)",
        (datetime.now().strftime("%Y-%m-%d %H:%M:%S"), user, action, details),
    )
    conn.commit()
    conn.close()


def get_setting(key, default=""):
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT value FROM app_settings WHERE key=?", (key,))
    row = c.fetchone()
    conn.close()
    return row[0] if row else default


def set_setting(key, value):
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "INSERT OR REPLACE INTO app_settings (key, value) VALUES (?, ?)",
        (key, str(value)),
    )
    conn.commit()
    conn.close()


def selected_tree_id(tree):
    item = tree.focus()
    if not item:
        return None
    values = tree.item(item).get("values", [])
    return values[0] if values else None


# =========================
# APP ROOT
# =========================
root = tk.Tk()
root.title(APP_TITLE)
root.configure(bg=BG)


def maximize_window(window):
    try:
        window.state("zoomed")
    except tk.TclError:
        try:
            window.attributes("-zoomed", True)
        except tk.TclError:
            screen_width = window.winfo_screenwidth()
            screen_height = window.winfo_screenheight()
            window.geometry(f"{screen_width}x{screen_height}+0+0")


maximize_window(root)
SCREEN_RESOLUTION_CACHE = f"{root.winfo_screenwidth()}x{root.winfo_screenheight()}"

style = ttk.Style()
style.theme_use("clam")


# =========================
# STYLES
# =========================
style.configure("TNotebook", background=BG, borderwidth=0)
style.configure("TNotebook.Tab", padding=(18, 10), font=("Segoe UI", 10, "bold"))

style.configure("TFrame", background=BG)
style.configure("Card.TFrame", background=CARD)

style.configure("TLabel", background=BG, foreground=TEXT, font=("Segoe UI", 10))
style.configure("Muted.TLabel", background=BG, foreground=MUTED, font=("Segoe UI", 9))
style.configure("Card.TLabel", background=CARD, foreground=TEXT, font=("Segoe UI", 10))
style.configure("Header.TLabel", background=BG, foreground=TEXT, font=("Segoe UI", 20, "bold"))
style.configure("SubHeader.TLabel", background=BG, foreground=MUTED, font=("Segoe UI", 10))
style.configure("CardHeader.TLabel", background=CARD, foreground=TEXT, font=("Segoe UI", 13, "bold"))
style.configure("KpiValue.TLabel", background=CARD, foreground=TEXT, font=("Segoe UI", 22, "bold"))
style.configure("KpiTitle.TLabel", background=CARD, foreground=MUTED, font=("Segoe UI", 9, "bold"))

style.configure(
    "TEntry",
    fieldbackground=ENTRY_BG,
    background=ENTRY_BG,
    foreground=TEXT,
    insertcolor=TEXT,
    bordercolor=BORDER,
    lightcolor=BORDER,
    darkcolor=BORDER,
    padding=7,
)

style.configure(
    "TCombobox",
    fieldbackground=ENTRY_BG,
    background=ENTRY_BG,
    foreground=TEXT,
    arrowcolor=TEXT,
    bordercolor=BORDER,
    padding=6,
)

style.configure("TCheckbutton", background=BG, foreground=TEXT, font=("Segoe UI", 10))
style.configure("Card.TCheckbutton", background=CARD, foreground=TEXT, font=("Segoe UI", 10))

style.configure(
    "TButton",
    font=("Segoe UI", 10, "bold"),
    padding=(12, 7),
    background=PRIMARY,
    foreground="white",
    borderwidth=0,
)
style.map("TButton", background=[("active", PRIMARY_DARK)])

style.configure(
    "Danger.TButton",
    font=("Segoe UI", 10, "bold"),
    padding=(12, 7),
    background=DANGER,
    foreground="white",
)
style.map("Danger.TButton", background=[("active", "#b91c1c")])

style.configure(
    "Success.TButton",
    font=("Segoe UI", 10, "bold"),
    padding=(12, 7),
    background=SUCCESS,
    foreground="white",
)
style.map("Success.TButton", background=[("active", "#15803d")])

style.configure(
    "Treeview",
    rowheight=30,
    background=ENTRY_BG,
    foreground=TEXT,
    fieldbackground=ENTRY_BG,
    bordercolor=BORDER,
)
style.configure(
    "Treeview.Heading",
    background=CARD,
    foreground=TEXT,
    font=("Segoe UI", 10, "bold"),
)
style.map("Treeview", background=[("selected", PRIMARY)], foreground=[("selected", "white")])


# =========================
# GLOBAL STATE
# =========================
is_logged_in = False
current_user = None
current_role = None
selected_id = None
selected_user_id = None
pc_audit_batch_started = False
network_jobs = {}
network_clients = {}
audit_archive_records = {}


# =========================
# HELPERS
# =========================
def make_card(parent, title=None, padx=24, pady=20):
    card = tk.Frame(parent, bg=CARD, highlightbackground=BORDER, highlightthickness=1)
    if title:
        ttk.Label(card, text=title, style="CardHeader.TLabel").pack(anchor="w", padx=padx, pady=(pady, 8))
    return card


def create_scrollable_area(parent, bg=BG):
    container = tk.Frame(parent, bg=bg)
    container.pack(fill="both", expand=True)

    canvas = tk.Canvas(container, bg=bg, highlightthickness=0)
    scrollbar_y = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
    scrollbar_x = ttk.Scrollbar(container, orient="horizontal", command=canvas.xview)

    scrollable_frame = tk.Frame(canvas, bg=bg)
    window_id = canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")

    def update_scrollregion(event=None):
        canvas.configure(scrollregion=canvas.bbox("all"))

    def resize_inner_frame(event):
        canvas.itemconfig(window_id, width=event.width)

    def on_mousewheel(event):
        canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    scrollable_frame.bind("<Configure>", update_scrollregion)
    canvas.bind("<Configure>", resize_inner_frame)
    canvas.bind("<Enter>", lambda event: canvas.bind_all("<MouseWheel>", on_mousewheel))
    canvas.bind("<Leave>", lambda event: canvas.unbind_all("<MouseWheel>"))

    canvas.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

    canvas.grid(row=0, column=0, sticky="nsew")
    scrollbar_y.grid(row=0, column=1, sticky="ns")
    scrollbar_x.grid(row=1, column=0, sticky="ew")

    container.rowconfigure(0, weight=1)
    container.columnconfigure(0, weight=1)

    return scrollable_frame


def clear_tree(tree):
    tree.delete(*tree.get_children())


def clear_field(widget):
    if isinstance(widget, ttk.Combobox):
        widget.set("")
    else:
        widget.delete(0, tk.END)


def set_field(widget, value):
    clear_field(widget)
    value = "" if value is None else value
    if isinstance(widget, ttk.Combobox):
        widget.set(str(value))
    else:
        widget.insert(0, str(value))


def get_field(widget):
    return widget.get().strip()


def normalize_input_values(values):
    values = list(values)
    if len(values) == FIELD_COUNT - 1:
        values.insert(LCD_TYPE_INDEX, "")
    values += [""] * (FIELD_COUNT - len(values))
    return values[:FIELD_COUNT]


def set_protected_tabs(state):
    if "login_frame" in globals():
        notebook.tab(login_frame, state="normal")

    for tab_index in range(1, notebook.index("end")):
        notebook.tab(tab_index, state=state)

    if state == "hidden":
        return

    if current_role == "User":
        hidden_for_user = [
            login_frame,
            db_frame,
            analytics_frame,
            network_frame,
            users_frame,
            logs_frame,
        ]
        for frame in hidden_for_user:
            notebook.tab(frame, state="hidden")

    if state == "normal" and current_role != "Admin" and "reports_frame" in globals():
        notebook.tab(reports_frame, state="hidden")


def require_login():
    if not is_logged_in:
        messagebox.showwarning("Access denied", "Please login first")
        return False
    return True


def require_admin():
    if not require_login():
        return False
    if current_role != "Admin":
        messagebox.showwarning("Access denied", "Admin role required")
        return False
    return True


def parse_iso_date(value, field_name, end_of_day=False):
    value = value.strip()
    if not value:
        return None
    try:
        parsed = datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        messagebox.showerror("Date filter", f"{field_name} must be YYYY-MM-DD")
        return False
    suffix = "23:59:59" if end_of_day else "00:00:00"
    return parsed.strftime("%Y-%m-%d") + " " + suffix


def get_default_export_folder():
    folder = get_setting("default_export_folder", BASE_DIR)
    return folder if os.path.isdir(folder) else BASE_DIR


def ask_save_path(default_ext, filetypes, title, initialfile):
    return filedialog.asksaveasfilename(
        defaultextension=default_ext,
        filetypes=filetypes,
        title=title,
        initialdir=get_default_export_folder(),
        initialfile=initialfile,
    )


def refresh_all_data():
    load_data(search_var.get() if "search_var" in globals() else "")
    refresh_analytics()
    build_report(show_message=False)
    load_users()
    load_logs()


def language_label(code):
    return "Ελληνικά" if code == "el" else "English"


def language_code(label):
    return "el" if str(label).strip().lower().startswith("ε") else "en"


UI_TRANSLATIONS = {
    "Login": "Σύνδεση",
    "Records": "Εγγραφές",
    "Database": "Βάση Δεδομένων",
    "Analytics": "Στατιστικά",
    "Reports": "Αναφορές",
    "Network Jobs": "Δικτυακές Εργασίες",
    "Device Results": "Αποτελέσματα Συσκευών",
    "Audit Archive": "Αρχείο Audit",
    "Manual Tests": "Χειροκίνητοι Έλεγχοι",
    "Users": "Χρήστες",
    "Settings": "Ρυθμίσεις",
    "Logs": "Αρχεία Καταγραφής",
    "Record Form": "Φόρμα Εγγραφής",
    "Create or update database records": "Δημιουργία ή ενημέρωση εγγραφών",
    "Record Username": "Χρήστης Εγγραφής",
    "Record Password": "Κωδικός Εγγραφής",
    "Leave password blank while editing to keep the stored password.": "Αφήστε κενό τον κωδικό στην επεξεργασία για να μείνει ο αποθηκευμένος.",
    "Option": "Επιλογή",
    "Accept Terms": "Αποδοχή όρων",
    "Tab empty fields": "Tab σε κενά πεδία",
    "Load PC Auditor Example": "Φόρτωση Παραδείγματος",
    "Read This PC": "Ανάγνωση αυτού του PC",
    "Database Records": "Εγγραφές Βάσης",
    "Search, select, update, delete and export records": "Αναζήτηση, επιλογή, ενημέρωση, διαγραφή και εξαγωγή εγγραφών",
    "Search": "Αναζήτηση",
    "Refresh": "Ανανέωση",
    "Build filtered reports and export the exact result set": "Δημιουργία φιλτραρισμένων αναφορών και εξαγωγή αποτελεσμάτων",
    "Build Report": "Δημιουργία Αναφοράς",
    "Export Excel": "Εξαγωγή Excel",
    "Export JSON": "Εξαγωγή JSON",
    "Network Jobs": "Δικτυακές Εργασίες",
    "Share PC Auditor jobs through the network and execute them locally on each client": "Μοίρασμα εργασιών PC Auditor στο δίκτυο και εκτέλεση τοπικά σε κάθε client",
    "Connection": "Σύνδεση",
    "Server URL": "URL Server",
    "This Client": "Αυτός ο Client",
    "Status": "Κατάσταση",
    "Save": "Αποθήκευση",
    "Test": "Δοκιμή",
    "Available Clients": "Διαθέσιμοι Clients",
    "Register This Client": "Εγγραφή αυτού του Client",
    "Refresh Clients": "Ανανέωση Clients",
    "Load Client Info To Form": "Φόρτωση Client στη Φόρμα",
    "Select a row to use it as the publish assignee.": "Επιλέξτε γραμμή για ανάθεση δημοσίευσης.",
    "Publish Current Form": "Δημοσίευση Τρέχουσας Φόρμας",
    "Job Title": "Τίτλος Εργασίας",
    "Assignee": "Ανάθεση",
    "Publish Job": "Δημοσίευση Εργασίας",
    "Fetch Jobs": "Λήψη Εργασιών",
    "Load Selected To Form": "Φόρτωση Επιλεγμένου στη Φόρμα",
    "Apply Selected": "Εφαρμογή Επιλεγμένου",
    "Mark Completed": "Σήμανση Ολοκληρωμένου",
    "Mark Failed": "Σήμανση Αποτυχίας",
    "Selected Job": "Επιλεγμένη Εργασία",
    "Remote client hardware diagnostics reported by the registration heartbeat": "Διαγνωστικά hardware από απομακρυσμένους clients",
    "Refresh Results": "Ανανέωση Αποτελεσμάτων",
    "Apply Filter": "Εφαρμογή Φίλτρου",
    "Selected Result": "Επιλεγμένο Αποτέλεσμα",
    "Historical PC audit snapshots from every registered network client": "Ιστορικά snapshots audit από κάθε δικτυακό client",
    "Refresh Archive": "Ανανέωση Αρχείου",
    "Client:": "Client:",
    "Limit:": "Όριο:",
    "Load To Form": "Φόρτωση στη Φόρμα",
    "Print HTML": "Εκτύπωση HTML",
    "Selected Audit": "Επιλεγμένο Audit",
    "All technician and auditor-style device/port checks in one view": "Όλοι οι τεχνικοί έλεγχοι συσκευών/θυρών σε μία προβολή",
    "Refresh Manual Tests": "Ανανέωση Χειροκίνητων Ελέγχων",
    "Manual tests required": "Απαιτούνται χειροκίνητοι έλεγχοι",
    "Save Requirement": "Αποθήκευση Απαίτησης",
    "Load Requirement": "Φόρτωση Απαίτησης",
    "Run Windows Tests": "Εκτέλεση Windows Tests",
    "Manage application users, roles and account status": "Διαχείριση χρηστών, ρόλων και κατάστασης λογαριασμών",
    "Username": "Χρήστης",
    "Password": "Κωδικός",
    "Role": "Ρόλος",
    "Active": "Ενεργός",
    "Add User": "Προσθήκη Χρήστη",
    "Update": "Ενημέρωση",
    "Clear": "Καθαρισμός",
    "Delete": "Διαγραφή",
    "Configure application preferences and database maintenance": "Ρυθμίσεις εφαρμογής και συντήρηση βάσης",
    "Application Settings": "Ρυθμίσεις Εφαρμογής",
    "System Name": "Όνομα Συστήματος",
    "Export Folder": "Φάκελος Εξαγωγών",
    "Keyboard Delay (ms)": "Καθυστέρηση Πληκτρολογίου (ms)",
    "Auto refresh dashboards after changes": "Αυτόματη ανανέωση μετά από αλλαγές",
    "Save Settings": "Αποθήκευση Ρυθμίσεων",
    "Reset Defaults": "Επαναφορά Προεπιλογών",
    "Windows License": "Άδεια Windows",
    "Edition": "Έκδοση",
    "Product Name": "Όνομα Προϊόντος",
    "Product Key": "Κλειδί Προϊόντος",
    "Channel": "Κανάλι",
    "Read Windows License": "Ανάγνωση Άδειας Windows",
    "Digital licenses may not expose a full product key.": "Οι ψηφιακές άδειες μπορεί να μην εμφανίζουν πλήρες product key.",
    "Database Maintenance": "Συντήρηση Βάσης",
    "Backup Database": "Backup Βάσης",
    "Batch Launcher": "Εκκίνηση Batch",
    "Refresh Batch Files": "Ανανέωση Batch Αρχείων",
    "Run Selected Background": "Εκτέλεση Επιλεγμένων στο Background",
    "Run All Background": "Εκτέλεση Όλων στο Background",
    "Audit trail for sessions, records, exports and administration": "Ιστορικό συνεδριών, εγγραφών, εξαγωγών και διαχείρισης",
    "User:": "Χρήστης:",
    "Action:": "Ενέργεια:",
    "Clear Logs": "Καθαρισμός Logs",
    "New / Clear": "Νέα / Καθαρισμός",
    "Save / Update": "Αποθήκευση / Ενημέρωση",
    "Apply Keys (F1)": "Εφαρμογή Πλήκτρων (F1)",
    "Logout": "Αποσύνδεση",
    "Internal administration dashboard": "Πίνακας διαχείρισης",
    "Status: Logged out": "Κατάσταση: Αποσυνδεδεμένος",
    "Demo credentials: admin / 1234": "Στοιχεία δοκιμής: admin / 1234",
}


def ui_text(value):
    if UI_LANGUAGE != "el":
        return value
    return UI_TRANSLATIONS.get(value, value)


def apply_ui_language(widget):
    if UI_LANGUAGE != "el":
        return
    try:
        text = widget.cget("text")
        translated = ui_text(text)
        if translated != text:
            widget.configure(text=translated)
    except tk.TclError:
        pass
    if isinstance(widget, ttk.Notebook):
        for tab_id in widget.tabs():
            text = widget.tab(tab_id, "text")
            widget.tab(tab_id, text=ui_text(text))
    if isinstance(widget, ttk.Treeview):
        for column in ("#0",) + tuple(widget["columns"]):
            try:
                text = widget.heading(column).get("text", "")
                translated = ui_text(text)
                if translated != text:
                    widget.heading(column, text=translated)
            except tk.TclError:
                pass
    for child in widget.winfo_children():
        apply_ui_language(child)


# =========================
# MAIN LAYOUT
# =========================
header = tk.Frame(root, bg=BG)
header.pack(fill="x", padx=18, pady=(14, 8))

left_header = tk.Frame(header, bg=BG)
left_header.pack(side="left")

title_label = ttk.Label(left_header, text=APP_TITLE, style="Header.TLabel")
title_label.pack(anchor="w")
ttk.Label(left_header, text=tr("Internal administration dashboard", "Πίνακας διαχείρισης"), style="SubHeader.TLabel").pack(anchor="w")

right_header = tk.Frame(header, bg=BG)
right_header.pack(side="right")

session_label = ttk.Label(right_header, text=tr("Status: Logged out", "Κατάσταση: Αποσυνδεδεμένος"), style="Muted.TLabel")
session_label.pack(side="left", padx=10)

logout_btn = ttk.Button(right_header, text=tr("Logout", "Αποσύνδεση"), command=lambda: logout(), style="Danger.TButton")
ttk.Label(right_header, text=f"Version {APP_VERSION}", style="Muted.TLabel").pack(side="left", padx=5)

notebook = ttk.Notebook(root)
notebook.pack(fill="both", expand=True, padx=18, pady=8)


# =========================
# LOGIN LOGIC
# =========================
def do_login():
    global is_logged_in, current_user, current_role

    user = login_username_entry.get().strip()
    pw = login_password_entry.get().strip()

    conn = get_conn()
    c = conn.cursor()
    c.execute(
        "SELECT id, username, role, status FROM app_users WHERE username=? AND password=?",
        (user, hash_password(pw)),
    )
    row = c.fetchone()

    if not row:
        conn.close()
        messagebox.showerror("Login", "Wrong credentials")
        return

    if row[3] != "Active":
        conn.close()
        messagebox.showerror("Login", "This account is disabled")
        return

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    c.execute("UPDATE app_users SET last_login=? WHERE id=?", (now, row[0]))
    conn.commit()
    conn.close()

    is_logged_in = True
    current_user = row[1]
    current_role = row[2]

    set_protected_tabs("normal")
    notebook.select(form_frame)
    logout_btn.pack(side="left", padx=10)
    session_label.config(text=f"Status: {current_user} ({current_role})")
    log_action("LOGIN", "User logged in", current_user)
    refresh_all_data()
    messagebox.showinfo("Login", "Login successful")


def logout():
    global is_logged_in, current_user, current_role

    if current_user:
        log_action("LOGOUT", "User logged out", current_user)

    is_logged_in = False
    current_user = None
    current_role = None
    clear_form()
    set_protected_tabs("hidden")
    notebook.select(login_frame)
    login_username_entry.delete(0, tk.END)
    login_password_entry.delete(0, tk.END)
    logout_btn.pack_forget()
    session_label.config(text="Status: Logged out")
    messagebox.showinfo("Logout", "Logged out")


# =========================
# LOGIN TAB
# =========================
login_frame = tk.Frame(notebook, bg=BG)
notebook.add(login_frame, text=tr("Login", "Σύνδεση"))

login_card = make_card(login_frame, "Secure Login")
login_card.pack(expand=True)

login_body = tk.Frame(login_card, bg=CARD)
login_body.pack(padx=34, pady=(8, 30))

ttk.Label(login_body, text=tr("Username", "Χρήστης"), style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=12, pady=8)
login_username_entry = ttk.Entry(login_body, width=34)
login_username_entry.grid(row=0, column=1, pady=8)

ttk.Label(login_body, text=tr("Password", "Κωδικός"), style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=12, pady=8)
login_password_entry = ttk.Entry(login_body, show="*", width=34)
login_password_entry.grid(row=1, column=1, pady=8)

ttk.Button(login_body, text=tr("Login", "Σύνδεση"), command=do_login, style="Success.TButton").grid(row=2, column=0, columnspan=2, pady=18)
ttk.Label(login_body, text=tr("Demo credentials: admin / 1234", "Στοιχεία δοκιμής: admin / 1234"), style="Card.TLabel").grid(row=3, column=0, columnspan=2)

login_password_entry.bind("<Return>", lambda event: do_login())


# =========================
# FORM TAB
# =========================
form_frame = tk.Frame(notebook, bg=BG)
notebook.add(form_frame, text=tr("Records", "Εγγραφές"), state="hidden")

form_scroll = create_scrollable_area(form_frame, bg=BG)

form_header = tk.Frame(form_scroll, bg=BG)
form_header.pack(fill="x", padx=10, pady=(12, 4))

ttk.Label(form_header, text=tr("Record Form", "Φόρμα Εγγραφής"), style="Header.TLabel").pack(anchor="w")
ttk.Label(form_header, text=tr("Create or update database records", "Δημιουργία ή ενημέρωση εγγραφών"), style="SubHeader.TLabel").pack(anchor="w")

form_card = make_card(form_scroll, "Data Entry")
form_card.pack(fill="x", padx=10, pady=10)

identity_container = tk.Frame(form_card, bg=CARD)
identity_container.pack(padx=24, pady=(8, 8), fill="x")

ttk.Label(identity_container, text="Asset Tag", style="Card.TLabel").grid(
    row=0, column=0, sticky="e", padx=(8, 10), pady=7
)
record_xenoglossi_entry = ttk.Entry(identity_container, width=34)
record_xenoglossi_entry.grid(row=0, column=1, sticky="ew", padx=(0, 28), pady=7)

ttk.Label(identity_container, text=tr("Record Username", "Χρήστης Εγγραφής"), style="Card.TLabel").grid(
    row=0, column=2, sticky="e", padx=(8, 10), pady=7
)
record_username_entry = ttk.Entry(identity_container, width=34)
record_username_entry.grid(row=0, column=3, sticky="ew", padx=(0, 28), pady=7)

ttk.Label(identity_container, text=tr("Record Password", "Κωδικός Εγγραφής"), style="Card.TLabel").grid(
    row=1, column=0, sticky="e", padx=(8, 10), pady=7
)
record_password_entry = ttk.Entry(identity_container, show="*", width=34)
record_password_entry.grid(row=1, column=1, sticky="ew", padx=(0, 28), pady=7)

ttk.Label(
    identity_container,
    text=tr("Leave password blank while editing to keep the stored password.", "Αφήστε κενό τον κωδικό στην επεξεργασία για να μείνει ο αποθηκευμένος."),
    style="Card.TLabel",
).grid(row=1, column=3, sticky="w", pady=(0, 7))

identity_container.columnconfigure(1, weight=1)
identity_container.columnconfigure(3, weight=1)

form_container = tk.Frame(form_card, bg=CARD)
form_container.pack(padx=24, pady=(8, 22), fill="x")

entries = []

for i, field_name in enumerate(AUDITOR_FIELDS):
    row = i % 16
    col_block = i // 16
    label_col = col_block * 2
    entry_col = label_col + 1

    ttk.Label(form_container, text=field_name, style="Card.TLabel").grid(
        row=row, column=label_col, sticky="e", padx=(8, 10), pady=7
    )

    if field_name == "OS COA":
        e = ttk.Combobox(form_container, values=OS_COA_OPTIONS, width=32, state="readonly")
    elif field_name == "LCD type":
        e = ttk.Combobox(form_container, values=LCD_TYPE_OPTIONS, width=32, state="readonly")
    else:
        e = ttk.Entry(form_container, width=34)
    e.grid(row=row, column=entry_col, sticky="ew", padx=(0, 28), pady=7)
    entries.append(e)

for col in range(1, ((len(AUDITOR_FIELDS) + 15) // 16) * 2, 2):
    form_container.columnconfigure(col, weight=1)

options_row = tk.Frame(form_card, bg=CARD)
options_row.pack(fill="x", padx=24, pady=(0, 22))

ttk.Label(options_row, text=tr("Option", "Επιλογή"), style="Card.TLabel").pack(side="left", padx=(8, 10))
dropdown = ttk.Combobox(options_row, values=[f"Option {i}" for i in range(1, 11)], width=28, state="readonly")
dropdown.pack(side="left", padx=(0, 25))

checkbox_var = tk.BooleanVar()
ttk.Checkbutton(options_row, text=tr("Accept Terms", "Αποδοχή όρων"), variable=checkbox_var, style="Card.TCheckbutton").pack(side="left")

tab_empty_var = tk.BooleanVar(value=True)
ttk.Checkbutton(
    options_row,
    text="Tab empty fields",
    variable=tab_empty_var,
    style="Card.TCheckbutton",
).pack(side="left", padx=(20, 0))

ttk.Button(
    options_row,
    text=tr("Load PC Auditor Example", "Φόρτωση Παραδείγματος"),
    command=lambda: load_auditor_sample(),
).pack(side="left", padx=(18, 0))

ttk.Button(
    options_row,
    text=tr("Read This PC", "Ανάγνωση αυτού του PC"),
    command=lambda: start_read_current_pc_info(),
).pack(side="left", padx=(12, 0))

record_actions_row = tk.Frame(form_card, bg=CARD)
record_actions_row.pack(fill="x", padx=24, pady=(0, 22))

ttk.Button(
    record_actions_row,
    text="New / Clear",
    command=lambda: clear_form(),
).pack(side="left", padx=(8, 6), pady=4)
ttk.Button(
    record_actions_row,
    text="Save / Update",
    command=lambda: submit(),
    style="Success.TButton",
).pack(side="left", padx=6, pady=4)
ttk.Button(
    record_actions_row,
    text="Delete",
    command=lambda: delete(),
    style="Danger.TButton",
).pack(side="left", padx=6, pady=4)
ttk.Button(
    record_actions_row,
    text="Export Excel",
    command=lambda: export_excel(),
).pack(side="left", padx=6, pady=4)
ttk.Button(
    record_actions_row,
    text="Export JSON",
    command=lambda: export_json(),
).pack(side="left", padx=6, pady=4)
ttk.Button(
    record_actions_row,
    text="Apply Keys (F1)",
    command=lambda: apply_keyboard_actions(),
).pack(side="left", padx=6, pady=4)
ttk.Button(
    record_actions_row,
    text="Run PC Audit",
    command=lambda: run_pc_audit_batch(),
).pack(side="left", padx=6, pady=4)


# =========================
# DATABASE TAB
# =========================
db_frame = tk.Frame(notebook, bg=BG)
notebook.add(db_frame, text="Database", state="hidden")

db_header = tk.Frame(db_frame, bg=BG)
db_header.pack(fill="x", padx=10, pady=(12, 4))

ttk.Label(db_header, text="Database Records", style="Header.TLabel").pack(anchor="w")
ttk.Label(db_header, text="Search, select, update, delete and export records", style="SubHeader.TLabel").pack(anchor="w")

toolbar = tk.Frame(db_frame, bg=BG)
toolbar.pack(fill="x", padx=10, pady=8)

search_var = tk.StringVar()
ttk.Label(toolbar, text="Search username / Asset Tag:").pack(side="left")
search_entry = ttk.Entry(toolbar, textvariable=search_var, width=30)
search_entry.pack(side="left", padx=8)
ttk.Button(toolbar, text="Search", command=lambda: load_data(search_var.get())).pack(side="left", padx=4)
ttk.Button(toolbar, text="Refresh", command=lambda: load_data()).pack(side="left", padx=4)

columns = ("RecordID", "Asset Tag", "Timestamp", "Username", "Dropdown", "Accepted")
tree_container = tk.Frame(db_frame, bg=BG)
tree_container.pack(fill="both", expand=True, padx=10, pady=10)

tree = ttk.Treeview(tree_container, columns=columns, show="headings")
tree_scroll_y = ttk.Scrollbar(tree_container, orient="vertical", command=tree.yview)
tree_scroll_x = ttk.Scrollbar(tree_container, orient="horizontal", command=tree.xview)
tree.configure(yscrollcommand=tree_scroll_y.set, xscrollcommand=tree_scroll_x.set)

for col in columns:
    tree.heading(col, text=col)
    tree.column(col, anchor="center", width=160, stretch=True)
tree.configure(displaycolumns=("Asset Tag", "Timestamp", "Username", "Dropdown", "Accepted"))

tree.grid(row=0, column=0, sticky="nsew")
tree_scroll_y.grid(row=0, column=1, sticky="ns")
tree_scroll_x.grid(row=1, column=0, sticky="ew")

tree_container.rowconfigure(0, weight=1)
tree_container.columnconfigure(0, weight=1)


# =========================
# ANALYTICS TAB
# =========================
analytics_frame = tk.Frame(notebook, bg=BG)
notebook.add(analytics_frame, text="Analytics", state="hidden")

analytics_content = tk.Frame(analytics_frame, bg=BG)
analytics_content.pack(fill="both", expand=True, padx=24, pady=24)

analytics_top = tk.Frame(analytics_content, bg=BG)
analytics_top.pack(fill="x")
ttk.Label(analytics_top, text="Analytics", style="Header.TLabel").pack(side="left", anchor="w")
ttk.Button(analytics_top, text="Refresh", command=lambda: refresh_analytics()).pack(side="right")
ttk.Label(
    analytics_content,
    text="Live overview of records, users and system activity",
    style="SubHeader.TLabel",
).pack(anchor="w", pady=(0, 18))

kpi_frame = tk.Frame(analytics_content, bg=BG)
kpi_frame.pack(fill="x", pady=(0, 16))

kpi_vars = {}
for idx, key in enumerate(["Total Records", "Accepted Terms", "Active Users", "Exports"]):
    card = make_card(kpi_frame)
    card.grid(row=0, column=idx, sticky="ew", padx=(0 if idx == 0 else 8, 0))
    ttk.Label(card, text=key.upper(), style="KpiTitle.TLabel").pack(anchor="w", padx=18, pady=(14, 2))
    value_label = ttk.Label(card, text="0", style="KpiValue.TLabel")
    value_label.pack(anchor="w", padx=18, pady=(0, 14))
    kpi_vars[key] = value_label
    kpi_frame.columnconfigure(idx, weight=1)

analytics_bottom = tk.Frame(analytics_content, bg=BG)
analytics_bottom.pack(fill="both", expand=True)

option_card = make_card(analytics_bottom, "Option Distribution")
option_card.pack(side="left", fill="both", expand=True, padx=(0, 8))

option_tree = ttk.Treeview(option_card, columns=("Option", "Records"), show="headings", height=10)
for col in ("Option", "Records"):
    option_tree.heading(col, text=col)
    option_tree.column(col, anchor="center", width=160, stretch=True)
option_tree.pack(fill="both", expand=True, padx=18, pady=(8, 18))

activity_card = make_card(analytics_bottom, "Recent Activity")
activity_card.pack(side="left", fill="both", expand=True, padx=(8, 0))

activity_tree = ttk.Treeview(activity_card, columns=("Time", "User", "Action"), show="headings", height=10)
for col in ("Time", "User", "Action"):
    activity_tree.heading(col, text=col)
    activity_tree.column(col, anchor="center", width=150, stretch=True)
activity_tree.pack(fill="both", expand=True, padx=18, pady=(8, 18))


# =========================
# REPORTS TAB
# =========================
reports_frame = tk.Frame(notebook, bg=BG)
notebook.add(reports_frame, text="Reports", state="hidden")

reports_content = tk.Frame(reports_frame, bg=BG)
reports_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(reports_content, text="Reports", style="Header.TLabel").pack(anchor="w")
ttk.Label(
    reports_content,
    text="Build filtered reports and export the exact result set",
    style="SubHeader.TLabel",
).pack(anchor="w", pady=(0, 18))

report_filter_card = make_card(reports_content, "Report Filters")
report_filter_card.pack(fill="x", pady=(0, 14))

report_filters = tk.Frame(report_filter_card, bg=CARD)
report_filters.pack(fill="x", padx=24, pady=(8, 20))

report_username_var = tk.StringVar()
report_asset_tag_var = tk.StringVar()
report_option_var = tk.StringVar()
report_from_var = tk.StringVar()
report_to_var = tk.StringVar()

ttk.Label(report_filters, text="Username", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(report_filters, textvariable=report_username_var, width=24).grid(row=0, column=1, sticky="ew", padx=(0, 18), pady=6)
ttk.Label(report_filters, text="Asset Tag", style="Card.TLabel").grid(row=0, column=2, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(report_filters, textvariable=report_asset_tag_var, width=24).grid(row=0, column=3, sticky="ew", padx=(0, 18), pady=6)
ttk.Label(report_filters, text="Option", style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=6)
report_option_combo = ttk.Combobox(
    report_filters,
    textvariable=report_option_var,
    values=[""] + [f"Option {i}" for i in range(1, 11)],
    width=22,
    state="readonly",
)
report_option_combo.grid(row=1, column=1, sticky="ew", padx=(0, 18), pady=6)
ttk.Label(report_filters, text="From", style="Card.TLabel").grid(row=1, column=2, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(report_filters, textvariable=report_from_var, width=24).grid(row=1, column=3, sticky="ew", padx=(0, 18), pady=6)
ttk.Label(report_filters, text="To", style="Card.TLabel").grid(row=2, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(report_filters, textvariable=report_to_var, width=24).grid(row=2, column=1, sticky="ew", padx=(0, 18), pady=6)
ttk.Button(report_filters, text="Build Report", command=lambda: build_report()).grid(row=0, column=4, padx=4, pady=6)
ttk.Button(report_filters, text="Export Excel", command=lambda: export_report_excel()).grid(row=1, column=4, padx=4, pady=6)

for idx in (1, 3):
    report_filters.columnconfigure(idx, weight=1)

report_summary_var = tk.StringVar(value="No report built yet")
ttk.Label(reports_content, textvariable=report_summary_var, style="SubHeader.TLabel").pack(anchor="w", pady=(0, 8))

report_tree_container = tk.Frame(reports_content, bg=BG)
report_tree_container.pack(fill="both", expand=True)

report_columns = ("Asset Tag", "Timestamp", "Username", "Dropdown", "Accepted", "Filled Inputs")
report_tree = ttk.Treeview(report_tree_container, columns=report_columns, show="headings")
report_y = ttk.Scrollbar(report_tree_container, orient="vertical", command=report_tree.yview)
report_x = ttk.Scrollbar(report_tree_container, orient="horizontal", command=report_tree.xview)
report_tree.configure(yscrollcommand=report_y.set, xscrollcommand=report_x.set)

for col in report_columns:
    report_tree.heading(col, text=col)
    report_tree.column(col, anchor="center", width=150, stretch=True)

report_tree.grid(row=0, column=0, sticky="nsew")
report_y.grid(row=0, column=1, sticky="ns")
report_x.grid(row=1, column=0, sticky="ew")
report_tree_container.rowconfigure(0, weight=1)
report_tree_container.columnconfigure(0, weight=1)


# =========================
# NETWORK JOBS TAB
# =========================
network_frame = tk.Frame(notebook, bg=BG)
notebook.add(network_frame, text="Network Jobs", state="hidden")

network_scroll = create_scrollable_area(network_frame, bg=BG)
network_content = tk.Frame(network_scroll, bg=BG)
network_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(network_content, text="Network Jobs", style="Header.TLabel").pack(anchor="w")
ttk.Label(
    network_content,
    text="Share PC Auditor jobs through the network and execute them locally on each client",
    style="SubHeader.TLabel",
).pack(anchor="w", pady=(0, 18))

network_config_card = make_card(network_content, "Connection")
network_config_card.pack(fill="x", pady=(0, 14))

network_config = tk.Frame(network_config_card, bg=CARD)
network_config.pack(fill="x", padx=24, pady=(8, 20))

network_server_url_var = tk.StringVar()
workstation_name_var = tk.StringVar()
job_assignee_var = tk.StringVar(value="ALL")
job_title_var = tk.StringVar()
job_status_filter_var = tk.StringVar(value="Queued")

ttk.Label(network_config, text="Server URL", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(network_config, textvariable=network_server_url_var, width=36).grid(
    row=0, column=1, columnspan=3, sticky="ew", padx=(0, 18), pady=6
)
ttk.Label(network_config, text="This Client", style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(network_config, textvariable=workstation_name_var, width=24).grid(
    row=1, column=1, sticky="ew", padx=(0, 18), pady=6
)
ttk.Label(network_config, text="Status", style="Card.TLabel").grid(row=1, column=2, sticky="e", padx=(0, 8), pady=6)
ttk.Combobox(
    network_config,
    textvariable=job_status_filter_var,
    values=["Queued", "Loaded", "Completed", "Cancelled", "Failed", "All"],
    state="readonly",
    width=22,
).grid(row=1, column=3, sticky="ew", padx=(0, 18), pady=6)
ttk.Button(network_config, text="Save", command=lambda: save_network_settings()).grid(row=0, column=4, padx=4, pady=6)
ttk.Button(network_config, text="Test", command=lambda: test_network_connection()).grid(row=1, column=4, padx=4, pady=6)

for idx in (1, 3):
    network_config.columnconfigure(idx, weight=1)

network_clients_card = make_card(network_content, "Available Clients")
network_clients_card.pack(fill="x", pady=(0, 14))

network_clients_body = tk.Frame(network_clients_card, bg=CARD)
network_clients_body.pack(fill="x", padx=24, pady=(8, 20))

network_clients_toolbar = tk.Frame(network_clients_body, bg=CARD)
network_clients_toolbar.pack(fill="x", pady=(0, 8))
ttk.Button(
    network_clients_toolbar,
    text="Register This Client",
    command=lambda: register_this_client(),
).pack(side="left", padx=(0, 6))
ttk.Button(
    network_clients_toolbar,
    text="Refresh Clients",
    command=lambda: fetch_network_clients(),
).pack(side="left", padx=6)
ttk.Button(
    network_clients_toolbar,
    text="Load Client Info To Form",
    command=lambda: load_selected_client_info_to_form(),
).pack(side="left", padx=6)
ttk.Label(
    network_clients_toolbar,
    text="Select a row to use it as the publish assignee.",
    style="Card.TLabel",
).pack(side="left", padx=12)

network_client_columns = ("Name", "IP", "MAC", "User", "Role", "Last Seen", "Platform", "App")
network_client_tree = ttk.Treeview(network_clients_body, columns=network_client_columns, show="headings", height=5)
client_y = ttk.Scrollbar(network_clients_body, orient="vertical", command=network_client_tree.yview)
client_x = ttk.Scrollbar(network_clients_body, orient="horizontal", command=network_client_tree.xview)
network_client_tree.configure(yscrollcommand=client_y.set, xscrollcommand=client_x.set)

for col in network_client_columns:
    network_client_tree.heading(col, text=col)
    network_client_tree.column(col, anchor="center", width=140, stretch=True)
network_client_tree.column("Name", width=180)
network_client_tree.column("IP", width=130)
network_client_tree.column("MAC", width=150)
network_client_tree.column("Platform", width=180)

network_client_tree.pack(side="left", fill="both", expand=True)
client_y.pack(side="left", fill="y")
client_x.pack(fill="x")

network_publish_card = make_card(network_content, "Publish Current Form")
network_publish_card.pack(fill="x", pady=(0, 14))

network_publish = tk.Frame(network_publish_card, bg=CARD)
network_publish.pack(fill="x", padx=24, pady=(8, 20))

ttk.Label(network_publish, text="Job Title", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(network_publish, textvariable=job_title_var, width=36).grid(
    row=0, column=1, sticky="ew", padx=(0, 18), pady=6
)
ttk.Label(network_publish, text="Assignee", style="Card.TLabel").grid(row=0, column=2, sticky="e", padx=(0, 8), pady=6)
job_assignee_combo = ttk.Combobox(network_publish, textvariable=job_assignee_var, values=["ALL"], width=24)
job_assignee_combo.grid(
    row=0, column=3, sticky="ew", padx=(0, 18), pady=6
)
ttk.Button(network_publish, text="Publish Job", command=lambda: publish_current_form_job(), style="Success.TButton").grid(
    row=0, column=4, padx=4, pady=6
)
network_publish.columnconfigure(1, weight=1)
network_publish.columnconfigure(3, weight=1)

network_toolbar = tk.Frame(network_content, bg=BG)
network_toolbar.pack(fill="x", pady=(0, 8))
ttk.Button(network_toolbar, text="Fetch Jobs", command=lambda: fetch_network_jobs()).pack(side="left", padx=(0, 6))
ttk.Button(network_toolbar, text="Load Selected To Form", command=lambda: load_selected_network_job()).pack(side="left", padx=6)
ttk.Button(network_toolbar, text="Apply Selected", command=lambda: apply_selected_network_job()).pack(side="left", padx=6)
ttk.Button(network_toolbar, text="Mark Completed", command=lambda: complete_selected_network_job(), style="Success.TButton").pack(
    side="left", padx=6
)
ttk.Button(network_toolbar, text="Mark Failed", command=lambda: fail_selected_network_job(), style="Danger.TButton").pack(
    side="left", padx=6
)

network_body = tk.Frame(network_content, bg=BG)
network_body.pack(fill="both", expand=True)

network_tree_container = tk.Frame(network_body, bg=BG)
network_tree_container.pack(side="left", fill="both", expand=True, padx=(0, 8))

network_columns = ("ID", "Title", "Assignee", "Status", "Created By", "Created At", "Updated At")
network_tree = ttk.Treeview(network_tree_container, columns=network_columns, show="headings")
network_y = ttk.Scrollbar(network_tree_container, orient="vertical", command=network_tree.yview)
network_x = ttk.Scrollbar(network_tree_container, orient="horizontal", command=network_tree.xview)
network_tree.configure(yscrollcommand=network_y.set, xscrollcommand=network_x.set)

for col in network_columns:
    network_tree.heading(col, text=col)
    network_tree.column(col, anchor="center", width=130, stretch=True)
network_tree.column("Title", width=220)

network_tree.grid(row=0, column=0, sticky="nsew")
network_y.grid(row=0, column=1, sticky="ns")
network_x.grid(row=1, column=0, sticky="ew")
network_tree_container.rowconfigure(0, weight=1)
network_tree_container.columnconfigure(0, weight=1)

network_details_card = make_card(network_body, "Selected Job")
network_details_card.pack(side="left", fill="both", expand=True, padx=(8, 0))

network_details_text = tk.Text(
    network_details_card,
    bg=ENTRY_BG,
    fg=TEXT,
    insertbackground=TEXT,
    relief="flat",
    height=12,
    wrap="word",
)
network_details_text.pack(fill="both", expand=True, padx=18, pady=(8, 18))
network_details_text.configure(state="disabled")


# =========================
# DEVICE RESULTS TAB
# =========================
device_results_frame = tk.Frame(notebook, bg=BG)
notebook.add(device_results_frame, text="Device Results", state="hidden")

device_results_content = tk.Frame(device_results_frame, bg=BG)
device_results_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(device_results_content, text="Device Results", style="Header.TLabel").pack(anchor="w")
ttk.Label(
    device_results_content,
    text="Remote client hardware diagnostics reported by the registration heartbeat",
    style="SubHeader.TLabel",
).pack(anchor="w", pady=(0, 18))

device_results_toolbar = tk.Frame(device_results_content, bg=BG)
device_results_toolbar.pack(fill="x", pady=(0, 10))

device_status_filter_var = tk.StringVar(value="All")
ttk.Button(device_results_toolbar, text="Refresh Results", command=lambda: refresh_device_results()).pack(side="left", padx=(0, 6))
ttk.Label(device_results_toolbar, text="Status:").pack(side="left", padx=(12, 6))
ttk.Combobox(
    device_results_toolbar,
    textvariable=device_status_filter_var,
    values=["All", "OK", "WARN", "FAIL", "UNKNOWN"],
    state="readonly",
    width=14,
).pack(side="left", padx=(0, 8))
ttk.Button(device_results_toolbar, text="Apply Filter", command=lambda: populate_device_results()).pack(side="left", padx=6)

device_summary_var = tk.StringVar(value="No diagnostics loaded")
ttk.Label(device_results_content, textvariable=device_summary_var, style="SubHeader.TLabel").pack(anchor="w", pady=(0, 8))

device_results_body = tk.Frame(device_results_content, bg=BG)
device_results_body.pack(fill="both", expand=True)

device_results_tree_container = tk.Frame(device_results_body, bg=BG)
device_results_tree_container.pack(side="left", fill="both", expand=True, padx=(0, 8))

device_result_columns = ("Client", "Device", "Status", "Details", "Last Seen")
device_results_tree = ttk.Treeview(device_results_tree_container, columns=device_result_columns, show="headings")
device_results_y = ttk.Scrollbar(device_results_tree_container, orient="vertical", command=device_results_tree.yview)
device_results_x = ttk.Scrollbar(device_results_tree_container, orient="horizontal", command=device_results_tree.xview)
device_results_tree.configure(yscrollcommand=device_results_y.set, xscrollcommand=device_results_x.set)

for col in device_result_columns:
    device_results_tree.heading(col, text=col)
    device_results_tree.column(col, anchor="center", width=150, stretch=True)
device_results_tree.column("Details", width=420)

device_results_tree.grid(row=0, column=0, sticky="nsew")
device_results_y.grid(row=0, column=1, sticky="ns")
device_results_x.grid(row=1, column=0, sticky="ew")
device_results_tree_container.rowconfigure(0, weight=1)
device_results_tree_container.columnconfigure(0, weight=1)

device_details_card = make_card(device_results_body, "Selected Result")
device_details_card.pack(side="left", fill="both", expand=True, padx=(8, 0))

device_details_text = tk.Text(
    device_details_card,
    bg=ENTRY_BG,
    fg=TEXT,
    insertbackground=TEXT,
    relief="flat",
    height=12,
    wrap="word",
)
device_details_text.pack(fill="both", expand=True, padx=18, pady=(8, 18))
device_details_text.configure(state="disabled")


# =========================
# AUDIT ARCHIVE TAB
# =========================
audit_archive_frame = tk.Frame(notebook, bg=BG)
notebook.add(audit_archive_frame, text="Audit Archive", state="hidden")

audit_archive_content = tk.Frame(audit_archive_frame, bg=BG)
audit_archive_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(audit_archive_content, text="Audit Archive", style="Header.TLabel").pack(anchor="w")
ttk.Label(
    audit_archive_content,
    text="Historical PC audit snapshots from every registered network client",
    style="SubHeader.TLabel",
).pack(anchor="w", pady=(0, 18))

audit_archive_toolbar = tk.Frame(audit_archive_content, bg=BG)
audit_archive_toolbar.pack(fill="x", pady=(0, 10))

audit_archive_client_var = tk.StringVar()
audit_archive_limit_var = tk.StringVar(value="500")
ttk.Button(audit_archive_toolbar, text="Refresh Archive", command=lambda: fetch_audit_archive()).pack(side="left", padx=(0, 6))
ttk.Label(audit_archive_toolbar, text="Client:").pack(side="left", padx=(12, 6))
audit_archive_client_combo = ttk.Combobox(audit_archive_toolbar, textvariable=audit_archive_client_var, values=[""], width=24)
audit_archive_client_combo.pack(side="left", padx=(0, 8))
ttk.Label(audit_archive_toolbar, text="Limit:").pack(side="left", padx=(12, 6))
ttk.Entry(audit_archive_toolbar, textvariable=audit_archive_limit_var, width=8).pack(side="left", padx=(0, 8))
ttk.Button(audit_archive_toolbar, text="Load To Form", command=lambda: load_selected_audit_to_form()).pack(side="left", padx=6)
ttk.Button(audit_archive_toolbar, text="Print HTML", command=lambda: print_audit_archive_html()).pack(side="left", padx=6)

audit_archive_summary_var = tk.StringVar(value="No audit history loaded")
ttk.Label(audit_archive_content, textvariable=audit_archive_summary_var, style="SubHeader.TLabel").pack(anchor="w", pady=(0, 8))

audit_archive_body = tk.Frame(audit_archive_content, bg=BG)
audit_archive_body.pack(fill="both", expand=True)

audit_archive_tree_container = tk.Frame(audit_archive_body, bg=BG)
audit_archive_tree_container.pack(side="left", fill="both", expand=True, padx=(0, 8))

audit_archive_columns = ("ID", "Client", "Created At", "Platform", "OK", "WARN", "FAIL", "UNKNOWN")
audit_archive_tree = ttk.Treeview(audit_archive_tree_container, columns=audit_archive_columns, show="headings")
audit_archive_y = ttk.Scrollbar(audit_archive_tree_container, orient="vertical", command=audit_archive_tree.yview)
audit_archive_x = ttk.Scrollbar(audit_archive_tree_container, orient="horizontal", command=audit_archive_tree.xview)
audit_archive_tree.configure(yscrollcommand=audit_archive_y.set, xscrollcommand=audit_archive_x.set)

for col in audit_archive_columns:
    audit_archive_tree.heading(col, text=col)
    audit_archive_tree.column(col, anchor="center", width=115, stretch=True)
audit_archive_tree.column("Client", width=180)
audit_archive_tree.column("Created At", width=160)
audit_archive_tree.column("Platform", width=220)

audit_archive_tree.grid(row=0, column=0, sticky="nsew")
audit_archive_y.grid(row=0, column=1, sticky="ns")
audit_archive_x.grid(row=1, column=0, sticky="ew")
audit_archive_tree_container.rowconfigure(0, weight=1)
audit_archive_tree_container.columnconfigure(0, weight=1)

audit_archive_details_card = make_card(audit_archive_body, "Selected Audit")
audit_archive_details_card.pack(side="left", fill="both", expand=True, padx=(8, 0))

audit_archive_details_text = tk.Text(
    audit_archive_details_card,
    bg=ENTRY_BG,
    fg=TEXT,
    insertbackground=TEXT,
    relief="flat",
    height=12,
    wrap="word",
)
audit_archive_details_text.pack(fill="both", expand=True, padx=18, pady=(8, 18))
audit_archive_details_text.configure(state="disabled")


# =========================
# MANUAL TESTS TAB
# =========================
manual_tests_frame = tk.Frame(notebook, bg=BG)
notebook.add(manual_tests_frame, text="Manual Tests", state="hidden")

manual_tests_content = tk.Frame(manual_tests_frame, bg=BG)
manual_tests_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(manual_tests_content, text="Manual Tests", style="Header.TLabel").pack(anchor="w")
ttk.Label(
    manual_tests_content,
    text="All technician and auditor-style device/port checks in one view",
    style="SubHeader.TLabel",
).pack(anchor="w", pady=(0, 18))

manual_tests_toolbar = tk.Frame(manual_tests_content, bg=BG)
manual_tests_toolbar.pack(fill="x", pady=(0, 10))

manual_status_filter_var = tk.StringVar(value="All")
manual_required_var = tk.BooleanVar(value=False)
manual_windows_usb_expected_var = tk.StringVar()
manual_windows_optical_path_var = tk.StringVar()
manual_windows_usb_path_var = tk.StringVar()
manual_windows_usb_ports_var = tk.BooleanVar(value=True)
manual_windows_fast_usb_var = tk.BooleanVar(value=True)
manual_windows_optical_var = tk.BooleanVar(value=True)
manual_windows_removable_var = tk.BooleanVar(value=True)
manual_windows_sound_var = tk.BooleanVar(value=True)
manual_windows_upload_var = tk.BooleanVar(value=True)
manual_windows_test_status_var = tk.StringVar(value="Windows device tests not run")
ttk.Button(manual_tests_toolbar, text="Refresh Manual Tests", command=lambda: refresh_manual_tests()).pack(side="left", padx=(0, 6))
ttk.Label(manual_tests_toolbar, text="Status:").pack(side="left", padx=(12, 6))
ttk.Combobox(
    manual_tests_toolbar,
    textvariable=manual_status_filter_var,
    values=["All", "OK", "WARN", "FAIL", "UNKNOWN"],
    state="readonly",
    width=14,
).pack(side="left", padx=(0, 8))
ttk.Button(manual_tests_toolbar, text="Apply Filter", command=lambda: populate_manual_tests()).pack(side="left", padx=6)
ttk.Checkbutton(
    manual_tests_toolbar,
    text="Manual tests required",
    variable=manual_required_var,
).pack(side="left", padx=(18, 6))
ttk.Button(manual_tests_toolbar, text="Save Requirement", command=lambda: save_manual_requirement()).pack(side="left", padx=6)
ttk.Button(manual_tests_toolbar, text="Load Requirement", command=lambda: load_manual_requirement()).pack(side="left", padx=6)

manual_windows_card = make_card(manual_tests_content, "Windows Device Tests")
manual_windows_card.pack(fill="x", pady=(0, 10))

manual_windows_body = tk.Frame(manual_windows_card, bg=CARD)
manual_windows_body.pack(fill="x", padx=24, pady=(8, 18))

ttk.Label(manual_windows_body, text="USB ports", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(manual_windows_body, textvariable=manual_windows_usb_expected_var, width=10).grid(row=0, column=1, sticky="w", padx=(0, 18), pady=6)
ttk.Label(manual_windows_body, text="Optical path", style="Card.TLabel").grid(row=0, column=2, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(manual_windows_body, textvariable=manual_windows_optical_path_var, width=18).grid(row=0, column=3, sticky="ew", padx=(0, 8), pady=6)
ttk.Button(manual_windows_body, text="Browse", command=lambda: browse_manual_test_path(manual_windows_optical_path_var)).grid(row=0, column=4, sticky="w", padx=(0, 18), pady=6)

ttk.Label(manual_windows_body, text="USB/SD path", style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(manual_windows_body, textvariable=manual_windows_usb_path_var, width=18).grid(row=1, column=1, sticky="ew", padx=(0, 8), pady=6)
ttk.Button(manual_windows_body, text="Browse", command=lambda: browse_manual_test_path(manual_windows_usb_path_var)).grid(row=1, column=2, sticky="w", padx=(0, 18), pady=6)
ttk.Button(manual_windows_body, text="Detect Paths", command=lambda: detect_windows_test_paths()).grid(row=1, column=3, sticky="w", padx=(0, 8), pady=6)

manual_windows_options = tk.Frame(manual_windows_body, bg=CARD)
manual_windows_options.grid(row=2, column=0, columnspan=4, sticky="w", pady=(2, 6))
ttk.Checkbutton(manual_windows_options, text="USB ports", variable=manual_windows_usb_ports_var, style="Card.TCheckbutton").pack(side="left", padx=(0, 12))
ttk.Checkbutton(manual_windows_options, text="Fast USB", variable=manual_windows_fast_usb_var, style="Card.TCheckbutton").pack(side="left", padx=(0, 12))
ttk.Checkbutton(manual_windows_options, text="DVD read", variable=manual_windows_optical_var, style="Card.TCheckbutton").pack(side="left", padx=(0, 12))
ttk.Checkbutton(manual_windows_options, text="USB/SD read", variable=manual_windows_removable_var, style="Card.TCheckbutton").pack(side="left", padx=(0, 12))
ttk.Checkbutton(manual_windows_options, text="Sound L/R", variable=manual_windows_sound_var, style="Card.TCheckbutton").pack(side="left", padx=(0, 12))
ttk.Checkbutton(manual_windows_options, text="Upload", variable=manual_windows_upload_var, style="Card.TCheckbutton").pack(side="left", padx=(0, 12))

manual_windows_run_btn = ttk.Button(
    manual_windows_body,
    text="Run Selected Tests",
    command=lambda: run_windows_device_tests_from_tab(),
    style="Success.TButton",
)
manual_windows_run_btn.grid(row=2, column=4, sticky="w", padx=(0, 18), pady=6)
ttk.Button(
    manual_windows_body,
    text="Run BAT",
    command=lambda: run_windows_device_tests_batch(),
).grid(row=3, column=4, sticky="w", padx=(0, 18), pady=6)

ttk.Label(manual_windows_body, textvariable=manual_windows_test_status_var, style="Card.TLabel").grid(
    row=4, column=0, columnspan=5, sticky="w", pady=(6, 0)
)
manual_windows_body.columnconfigure(1, weight=1)
manual_windows_body.columnconfigure(3, weight=1)

manual_summary_var = tk.StringVar(value="No manual tests loaded")
ttk.Label(manual_tests_content, textvariable=manual_summary_var, style="SubHeader.TLabel").pack(anchor="w", pady=(0, 8))

manual_tests_body = tk.Frame(manual_tests_content, bg=BG)
manual_tests_body.pack(fill="both", expand=True)

manual_tree_container = tk.Frame(manual_tests_body, bg=BG)
manual_tree_container.pack(side="left", fill="both", expand=True, padx=(0, 8))

manual_columns = ("Client", "Test", "Status", "Note", "Last Seen")
manual_tests_tree = ttk.Treeview(manual_tree_container, columns=manual_columns, show="headings")
manual_y = ttk.Scrollbar(manual_tree_container, orient="vertical", command=manual_tests_tree.yview)
manual_x = ttk.Scrollbar(manual_tree_container, orient="horizontal", command=manual_tests_tree.xview)
manual_tests_tree.configure(yscrollcommand=manual_y.set, xscrollcommand=manual_x.set)

for col in manual_columns:
    manual_tests_tree.heading(col, text=col)
    manual_tests_tree.column(col, anchor="center", width=150, stretch=True)
manual_tests_tree.column("Note", width=440)

manual_tests_tree.grid(row=0, column=0, sticky="nsew")
manual_y.grid(row=0, column=1, sticky="ns")
manual_x.grid(row=1, column=0, sticky="ew")
manual_tree_container.rowconfigure(0, weight=1)
manual_tree_container.columnconfigure(0, weight=1)

manual_details_card = make_card(manual_tests_body, "Selected Test")
manual_details_card.pack(side="left", fill="both", expand=True, padx=(8, 0))

manual_details_text = tk.Text(
    manual_details_card,
    bg=ENTRY_BG,
    fg=TEXT,
    insertbackground=TEXT,
    relief="flat",
    height=12,
    wrap="word",
)
manual_details_text.pack(fill="both", expand=True, padx=18, pady=(8, 18))
manual_details_text.configure(state="disabled")


# =========================
# USERS TAB
# =========================
users_frame = tk.Frame(notebook, bg=BG)
notebook.add(users_frame, text="Users", state="hidden")

users_content = tk.Frame(users_frame, bg=BG)
users_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(users_content, text="Users", style="Header.TLabel").pack(anchor="w")
ttk.Label(users_content, text="Manage application users, roles and account status", style="SubHeader.TLabel").pack(
    anchor="w", pady=(0, 18)
)

user_form_card = make_card(users_content, "User Details")
user_form_card.pack(fill="x", pady=(0, 14))

user_form = tk.Frame(user_form_card, bg=CARD)
user_form.pack(fill="x", padx=24, pady=(8, 20))

user_username_var = tk.StringVar()
user_role_var = tk.StringVar(value="User")
user_active_var = tk.BooleanVar(value=True)

ttk.Label(user_form, text="Username", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Entry(user_form, textvariable=user_username_var, width=24).grid(row=0, column=1, sticky="ew", padx=(0, 18), pady=6)
ttk.Label(user_form, text="Password", style="Card.TLabel").grid(row=0, column=2, sticky="e", padx=(0, 8), pady=6)
user_password_entry = ttk.Entry(user_form, show="*", width=24)
user_password_entry.grid(row=0, column=3, sticky="ew", padx=(0, 18), pady=6)
ttk.Label(user_form, text="Role", style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=6)
ttk.Combobox(user_form, textvariable=user_role_var, values=["Admin", "User", "Viewer"], state="readonly", width=22).grid(
    row=1, column=1, sticky="ew", padx=(0, 18), pady=6
)
ttk.Checkbutton(user_form, text="Active", variable=user_active_var, style="Card.TCheckbutton").grid(
    row=1, column=2, sticky="w", padx=(0, 8), pady=6
)
ttk.Button(user_form, text="Add User", command=lambda: add_user(), style="Success.TButton").grid(row=0, column=4, padx=4)
ttk.Button(user_form, text="Update", command=lambda: update_user()).grid(row=1, column=4, padx=4)
ttk.Button(user_form, text="Clear", command=lambda: clear_user_form()).grid(row=0, column=5, padx=4)
ttk.Button(user_form, text="Delete", command=lambda: delete_user(), style="Danger.TButton").grid(row=1, column=5, padx=4)

for idx in (1, 3):
    user_form.columnconfigure(idx, weight=1)

user_tree_container = tk.Frame(users_content, bg=BG)
user_tree_container.pack(fill="both", expand=True)

user_columns = ("ID", "Username", "Role", "Status", "Created", "Last Login")
user_tree = ttk.Treeview(user_tree_container, columns=user_columns, show="headings")
user_y = ttk.Scrollbar(user_tree_container, orient="vertical", command=user_tree.yview)
user_x = ttk.Scrollbar(user_tree_container, orient="horizontal", command=user_tree.xview)
user_tree.configure(yscrollcommand=user_y.set, xscrollcommand=user_x.set)

for col in user_columns:
    user_tree.heading(col, text=col)
    user_tree.column(col, anchor="center", width=150, stretch=True)

user_tree.grid(row=0, column=0, sticky="nsew")
user_y.grid(row=0, column=1, sticky="ns")
user_x.grid(row=1, column=0, sticky="ew")
user_tree_container.rowconfigure(0, weight=1)
user_tree_container.columnconfigure(0, weight=1)


# =========================
# SETTINGS TAB
# =========================
settings_frame = tk.Frame(notebook, bg=BG)
notebook.add(settings_frame, text="Settings", state="hidden")

settings_scroll = create_scrollable_area(settings_frame, bg=BG)
settings_content = tk.Frame(settings_scroll, bg=BG)
settings_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(settings_content, text="Settings", style="Header.TLabel").pack(anchor="w")
ttk.Label(settings_content, text="Configure application preferences and database maintenance", style="SubHeader.TLabel").pack(
    anchor="w", pady=(0, 18)
)

settings_card = make_card(settings_content, "Application Settings")
settings_card.pack(fill="x", pady=(0, 14))

settings_form = tk.Frame(settings_card, bg=CARD)
settings_form.pack(fill="x", padx=24, pady=(8, 20))

setting_system_name_var = tk.StringVar()
setting_export_folder_var = tk.StringVar()
setting_keyboard_delay_var = tk.StringVar()
setting_auto_refresh_var = tk.BooleanVar()
setting_language_var = tk.StringVar(value=language_label(UI_LANGUAGE))
windows_license_edition_var = tk.StringVar(value="Not read")
windows_license_product_var = tk.StringVar(value="Not read")
windows_license_key_var = tk.StringVar(value="Not read")
windows_license_status_var = tk.StringVar(value="Not read")
windows_license_channel_var = tk.StringVar(value="Not read")

ttk.Label(settings_form, text="System Name", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(settings_form, textvariable=setting_system_name_var, width=42).grid(
    row=0, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Label(settings_form, text="Export Folder", style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(settings_form, textvariable=setting_export_folder_var, width=42).grid(
    row=1, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Button(settings_form, text="Browse", command=lambda: choose_export_folder()).grid(row=1, column=2, padx=4, pady=7)
ttk.Label(settings_form, text="Keyboard Delay (ms)", style="Card.TLabel").grid(
    row=2, column=0, sticky="e", padx=(0, 8), pady=7
)
ttk.Entry(settings_form, textvariable=setting_keyboard_delay_var, width=18).grid(
    row=2, column=1, sticky="w", padx=(0, 18), pady=7
)
ttk.Label(settings_form, text="Language", style="Card.TLabel").grid(
    row=3, column=0, sticky="e", padx=(0, 8), pady=7
)
ttk.Combobox(
    settings_form,
    textvariable=setting_language_var,
    values=["Ελληνικά", "English"],
    state="readonly",
    width=18,
).grid(row=3, column=1, sticky="w", padx=(0, 18), pady=7)
ttk.Checkbutton(settings_form, text="Auto refresh dashboards after changes", variable=setting_auto_refresh_var, style="Card.TCheckbutton").grid(
    row=4, column=1, sticky="w", pady=7
)
ttk.Button(settings_form, text="Save Settings", command=lambda: save_settings(), style="Success.TButton").grid(
    row=5, column=1, sticky="w", pady=(14, 0)
)
ttk.Button(settings_form, text="Reset Defaults", command=lambda: reset_settings()).grid(row=5, column=1, sticky="e", pady=(14, 0))

settings_form.columnconfigure(1, weight=1)

license_card = make_card(settings_content, "Windows License")
license_card.pack(fill="x", pady=(0, 14))

license_form = tk.Frame(license_card, bg=CARD)
license_form.pack(fill="x", padx=24, pady=(8, 20))

ttk.Label(license_form, text="Edition", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(license_form, textvariable=windows_license_edition_var, width=70, state="readonly").grid(
    row=0, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Label(license_form, text="Product Name", style="Card.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(license_form, textvariable=windows_license_product_var, width=70, state="readonly").grid(
    row=1, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Label(license_form, text="Product Key", style="Card.TLabel").grid(row=2, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(license_form, textvariable=windows_license_key_var, width=70, state="readonly").grid(
    row=2, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Label(license_form, text="Status", style="Card.TLabel").grid(row=3, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(license_form, textvariable=windows_license_status_var, width=70, state="readonly").grid(
    row=3, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Label(license_form, text="Channel", style="Card.TLabel").grid(row=4, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(license_form, textvariable=windows_license_channel_var, width=70, state="readonly").grid(
    row=4, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Button(license_form, text="Read Windows License", command=lambda: read_windows_license_to_settings()).grid(
    row=0, column=2, rowspan=2, sticky="nsew", padx=4, pady=7
)
ttk.Label(
    license_form,
    text="Digital licenses may not expose a full product key.",
    style="Card.TLabel",
).grid(row=5, column=1, sticky="w", pady=(4, 0))
license_form.columnconfigure(1, weight=1)

maintenance_card = make_card(settings_content, "Database Maintenance")
maintenance_card.pack(fill="x")

maintenance_body = tk.Frame(maintenance_card, bg=CARD)
maintenance_body.pack(fill="x", padx=24, pady=(8, 20))

db_path_var = tk.StringVar(value=DB_NAME)
ttk.Label(maintenance_body, text="Database", style="Card.TLabel").grid(row=0, column=0, sticky="e", padx=(0, 8), pady=7)
ttk.Entry(maintenance_body, textvariable=db_path_var, width=70, state="readonly").grid(
    row=0, column=1, sticky="ew", padx=(0, 18), pady=7
)
ttk.Button(maintenance_body, text="Backup Database", command=lambda: backup_database()).grid(row=0, column=2, padx=4, pady=7)
maintenance_body.columnconfigure(1, weight=1)

batch_launcher_card = make_card(settings_content, "Batch Launcher")
batch_launcher_card.pack(fill="both", expand=True, pady=(14, 0))

batch_launcher_body = tk.Frame(batch_launcher_card, bg=CARD)
batch_launcher_body.pack(fill="both", expand=True, padx=24, pady=(8, 20))

batch_toolbar = tk.Frame(batch_launcher_body, bg=CARD)
batch_toolbar.pack(fill="x", pady=(0, 8))

ttk.Button(batch_toolbar, text="Refresh Batch Files", command=lambda: load_batch_files()).pack(side="left", padx=(0, 6))
ttk.Button(batch_toolbar, text="Run Selected Background", command=lambda: run_selected_batch_files(), style="Success.TButton").pack(
    side="left", padx=6
)
ttk.Button(batch_toolbar, text="Run All Background", command=lambda: run_all_batch_files(), style="Danger.TButton").pack(side="left", padx=6)

batch_status_var = tk.StringVar(value="No batch files loaded")
ttk.Label(batch_toolbar, textvariable=batch_status_var, style="Card.TLabel").pack(side="left", padx=14)

batch_tree_container = tk.Frame(batch_launcher_body, bg=CARD)
batch_tree_container.pack(fill="both", expand=True)

batch_columns = ("File", "Modified", "Size")
batch_tree = ttk.Treeview(batch_tree_container, columns=batch_columns, show="headings", height=6, selectmode="extended")
batch_y = ttk.Scrollbar(batch_tree_container, orient="vertical", command=batch_tree.yview)
batch_tree.configure(yscrollcommand=batch_y.set)

for col in batch_columns:
    batch_tree.heading(col, text=col)
    batch_tree.column(col, anchor="center", width=160, stretch=True)
batch_tree.column("File", anchor="w", width=280)

batch_tree.pack(side="left", fill="both", expand=True)
batch_y.pack(side="left", fill="y")


# =========================
# LOGS TAB
# =========================
logs_frame = tk.Frame(notebook, bg=BG)
notebook.add(logs_frame, text="Logs", state="hidden")

logs_content = tk.Frame(logs_frame, bg=BG)
logs_content.pack(fill="both", expand=True, padx=24, pady=24)

ttk.Label(logs_content, text="Logs", style="Header.TLabel").pack(anchor="w")
ttk.Label(logs_content, text="Audit trail for sessions, records, exports and administration", style="SubHeader.TLabel").pack(
    anchor="w", pady=(0, 18)
)

log_toolbar = tk.Frame(logs_content, bg=BG)
log_toolbar.pack(fill="x", pady=(0, 10))

log_user_var = tk.StringVar()
log_action_var = tk.StringVar()
ttk.Label(log_toolbar, text="User:").pack(side="left")
ttk.Entry(log_toolbar, textvariable=log_user_var, width=22).pack(side="left", padx=8)
ttk.Label(log_toolbar, text="Action:").pack(side="left")
ttk.Entry(log_toolbar, textvariable=log_action_var, width=22).pack(side="left", padx=8)
ttk.Button(log_toolbar, text="Refresh", command=lambda: load_logs()).pack(side="left", padx=4)
ttk.Button(log_toolbar, text="Export Excel", command=lambda: export_logs_excel()).pack(side="left", padx=4)
ttk.Button(log_toolbar, text="Clear Logs", command=lambda: clear_logs(), style="Danger.TButton").pack(side="left", padx=4)

log_tree_container = tk.Frame(logs_content, bg=BG)
log_tree_container.pack(fill="both", expand=True)

log_columns = ("ID", "Timestamp", "User", "Action", "Details")
log_tree = ttk.Treeview(log_tree_container, columns=log_columns, show="headings")
log_y = ttk.Scrollbar(log_tree_container, orient="vertical", command=log_tree.yview)
log_x = ttk.Scrollbar(log_tree_container, orient="horizontal", command=log_tree.xview)
log_tree.configure(yscrollcommand=log_y.set, xscrollcommand=log_x.set)

for col in log_columns:
    log_tree.heading(col, text=col)
    log_tree.column(col, anchor="center", width=150, stretch=True)
log_tree.column("Details", width=320)

log_tree.grid(row=0, column=0, sticky="nsew")
log_y.grid(row=0, column=1, sticky="ns")
log_x.grid(row=1, column=0, sticky="ew")
log_tree_container.rowconfigure(0, weight=1)
log_tree_container.columnconfigure(0, weight=1)


# =========================
# SYSTEM INVENTORY
# =========================
def run_command(command, timeout=6):
    try:
        flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" and hasattr(subprocess, "CREATE_NO_WINDOW") else 0
        kwargs = {
            "capture_output": True,
            "text": True,
            "errors": "replace",
            "timeout": timeout,
        }
        if os.name == "nt":
            kwargs["creationflags"] = flags
        result = subprocess.run(
            command,
            **kwargs,
        )
    except Exception:
        return ""

    if result.returncode != 0:
        return ""
    return result.stdout.strip()


def powershell_value(script):
    command = [
        "powershell",
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-Command",
        "$ErrorActionPreference='SilentlyContinue'; " + script,
    ]
    value = run_command(command)
    if value:
        return value
    command[0] = "pwsh"
    return run_command(command)


def read_text_file(path):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            return handle.read().strip()
    except OSError:
        return ""


def read_binary_file(path):
    try:
        with open(path, "rb") as handle:
            return handle.read()
    except OSError:
        return b""


def first_line(value):
    for line in str(value).splitlines():
        line = line.strip()
        if line:
            return line
    return ""


def clean_serial_number(value):
    serial = first_line(value)
    if not serial:
        return ""
    normalized = serial.strip().lower()
    invalid_values = {
        "default string",
        "to be filled by o.e.m.",
        "to be filled by oem",
        "system serial number",
        "serial number",
        "none",
        "unknown",
        "not specified",
        "not available",
        "n/a",
        "na",
        "0",
    }
    if normalized in invalid_values:
        return ""
    return serial


def format_gb(bytes_value):
    try:
        gb = int(float(bytes_value)) / (1024 ** 3)
    except (TypeError, ValueError):
        return ""

    rounded = round(gb)
    if rounded <= 0:
        return ""
    return f"{rounded} GB"


def join_lines(value, limit=6):
    lines = [line.strip() for line in str(value).splitlines() if line.strip()]
    return "; ".join(lines[:limit])


def windows_domain_or_workgroup():
    return first_line(
        powershell_value(
            "$cs=Get-CimInstance Win32_ComputerSystem; "
            "if ($cs.PartOfDomain) { 'Domain: ' + $cs.Domain } else { 'Workgroup: ' + $cs.Workgroup }"
        )
    )


def windows_secure_boot():
    return first_line(powershell_value("try { (Confirm-SecureBootUEFI).ToString() } catch { '' }"))


def windows_tpm_status():
    return first_line(
        powershell_value(
            "try { $t=Get-Tpm; if ($null -ne $t.TpmPresent) { "
            "'Present=' + $t.TpmPresent + '; Ready=' + $t.TpmReady + '; Enabled=' + $t.TpmEnabled } } catch { '' }"
        )
    )


def windows_disk_details():
    return join_lines(
        powershell_value(
            "Get-CimInstance Win32_DiskDrive | Select-Object -First 8 | "
            "ForEach-Object { $_.Model + ' | ' + $_.InterfaceType + ' | ' + $_.MediaType + ' | ' + "
            "([math]::Round($_.Size / 1GB).ToString() + ' GB') }"
        ),
        limit=8,
    )


def windows_volume_details():
    return join_lines(
        powershell_value(
            "Get-CimInstance Win32_LogicalDisk -Filter \"DriveType=3\" | "
            "ForEach-Object { $_.DeviceID + ' ' + ([math]::Round($_.Size / 1GB).ToString() + ' GB') + "
            "' free ' + ([math]::Round($_.FreeSpace / 1GB).ToString() + ' GB') }"
        ),
        limit=8,
    )


def windows_network_details():
    return join_lines(
        powershell_value(
            "Get-CimInstance Win32_NetworkAdapterConfiguration | Where-Object { $_.IPEnabled } | "
            "Select-Object -First 6 | ForEach-Object { $_.Description + ' | MAC ' + $_.MACAddress + "
            "' | IP ' + (($_.IPAddress | Where-Object { $_ -notlike '*:*' }) -join ', ') }"
        ),
        limit=6,
    )


def windows_battery_details():
    return first_line(
        powershell_value(
            "$b=Get-CimInstance Win32_Battery | Select-Object -First 1; "
            "if ($b) { $b.Name + ' | ' + $b.EstimatedChargeRemaining + '% | status ' + $b.BatteryStatus }"
        )
    )


def windows_ram_slots():
    used = first_line(powershell_value("(Get-CimInstance Win32_PhysicalMemory | Measure-Object).Count"))
    total = first_line(
        powershell_value(
            "$arrays=Get-CimInstance Win32_PhysicalMemoryArray; "
            "if ($arrays) { ($arrays | Measure-Object -Property MemoryDevices -Sum).Sum }"
        )
    )
    if used and total:
        return f"{used} / {total}"
    return used


def linux_command_lines(command, timeout=6, limit=6):
    return join_lines(run_command(command, timeout=timeout), limit=limit)


def linux_ram_slots():
    dmidecode = shutil.which("dmidecode")
    if not dmidecode:
        return ""
    output = run_command([dmidecode, "-t", "memory"], timeout=8)
    if not output:
        return ""
    slots = 0
    used = 0
    in_device = False
    current_has_size = False
    for raw_line in output.splitlines():
        line = raw_line.strip()
        if line == "Memory Device":
            if in_device:
                used += 1 if current_has_size else 0
            slots += 1
            in_device = True
            current_has_size = False
        elif in_device and line.startswith("Size:") and "No Module Installed" not in line:
            current_has_size = True
    if in_device:
        used += 1 if current_has_size else 0
    return f"{used} / {slots}" if slots else ""


def chassis_label(chassis_value):
    try:
        chassis = int(str(chassis_value).strip().split()[0])
    except (TypeError, ValueError, IndexError):
        return ""

    laptop_codes = {8, 9, 10, 11, 12, 14, 18, 21, 30, 31, 32}
    desktop_codes = {3, 4, 5, 6, 7, 15, 16}
    tablet_codes = {13, 17, 23, 24}

    if chassis in laptop_codes:
        return "LAPTOP"
    if chassis in desktop_codes:
        return "DESKTOP"
    if chassis in tablet_codes:
        return "TABLET"
    return ""


def screen_resolution():
    return SCREEN_RESOLUTION_CACHE


def locale_label():
    lang = locale.getlocale()[0] or os.environ.get("LANG", "")
    return lang.replace("_", "-") if lang else ""


def infer_screen_type_from_text(text):
    haystack = f" {text} ".lower()
    if "amoled" in haystack:
        return "AMOLED"
    if "oled" in haystack:
        return "OLED"
    if "mini led" in haystack or "miniled" in haystack:
        return "Mini LED"
    if "ips" in haystack:
        return "IPS LCD"
    if " va " in haystack or "mva" in haystack or "pva" in haystack:
        return "VA LCD"
    if " tn " in haystack or "twisted nematic" in haystack:
        return "TN LCD"
    if "lcd" in haystack or "led" in haystack:
        return "LCD"
    return "Unknown"


def edid_descriptor_text(edid_bytes):
    if len(edid_bytes) < 128:
        return ""

    chunks = []
    for offset in range(54, 126, 18):
        block = edid_bytes[offset:offset + 18]
        if len(block) == 18 and block[:3] == b"\x00\x00\x00" and block[3] in {0xfc, 0xfe, 0xff}:
            text = block[5:18].decode("ascii", errors="ignore").strip().replace("\n", " ")
            if text:
                chunks.append(text)
    return " ".join(chunks)


def windows_monitor_text():
    monitor_name = powershell_value(
        "Get-CimInstance -Namespace root/wmi -ClassName WmiMonitorID | "
        "ForEach-Object { "
        "$name=($_.UserFriendlyName | Where-Object { $_ -ne 0 } | ForEach-Object { [char]$_ }) -join ''; "
        "$serial=($_.SerialNumberID | Where-Object { $_ -ne 0 } | ForEach-Object { [char]$_ }) -join ''; "
        "$name + ' ' + $serial }"
    )
    desktop_monitor = powershell_value(
        "Get-CimInstance Win32_DesktopMonitor | ForEach-Object { $_.Name + ' ' + $_.MonitorType }"
    )
    video = powershell_value(
        "Get-CimInstance Win32_VideoController | Select-Object -First 1 -ExpandProperty VideoModeDescription"
    )
    return " ".join([monitor_name, desktop_monitor, video])


def linux_monitor_text():
    parts = []
    drm_path = "/sys/class/drm"
    if os.path.isdir(drm_path):
        for name in os.listdir(drm_path):
            edid_path = os.path.join(drm_path, name, "edid")
            edid = read_binary_file(edid_path)
            text = edid_descriptor_text(edid)
            if text:
                parts.append(text)

    xrandr = run_command(["xrandr", "--verbose"], timeout=4)
    if xrandr:
        parts.append(xrandr)

    return " ".join(parts)


def windows_inventory():
    manufacturer = first_line(powershell_value("(Get-CimInstance Win32_ComputerSystem).Manufacturer"))
    model = first_line(powershell_value("(Get-CimInstance Win32_ComputerSystem).Model"))
    chassis = first_line(powershell_value("(Get-CimInstance Win32_SystemEnclosure).ChassisTypes[0]"))
    product_type = chassis_label(chassis)
    part_number = first_line(powershell_value("(Get-CimInstance Win32_ComputerSystemProduct).SKUNumber"))
    if not part_number:
        part_number = first_line(powershell_value("(Get-CimInstance Win32_ComputerSystemProduct).Name"))
    serial = clean_serial_number(powershell_value("(Get-CimInstance Win32_BIOS).SerialNumber"))
    if not serial:
        serial = clean_serial_number(powershell_value("(Get-CimInstance Win32_ComputerSystemProduct).IdentifyingNumber"))
    if not serial:
        serial = clean_serial_number(powershell_value("(Get-CimInstance Win32_BaseBoard).SerialNumber"))
    os_caption = first_line(powershell_value("(Get-CimInstance Win32_OperatingSystem).Caption"))
    os_version = first_line(
        powershell_value("$os=Get-CimInstance Win32_OperatingSystem; $os.Version + ' build ' + $os.BuildNumber")
    )
    motherboard = first_line(powershell_value("$bb=Get-CimInstance Win32_BaseBoard; $bb.Manufacturer + ' ' + $bb.Product"))
    bios_version = first_line(powershell_value("(Get-CimInstance Win32_BIOS).SMBIOSBIOSVersion"))
    processor = first_line(powershell_value("(Get-CimInstance Win32_Processor | Select-Object -First 1 -ExpandProperty Name)"))
    ram_size = first_line(
        powershell_value("[math]::Round((Get-CimInstance Win32_ComputerSystem).TotalPhysicalMemory / 1GB).ToString() + ' GB'")
    )
    disk_info = first_line(
        powershell_value(
            "Get-PhysicalDisk | Select-Object -First 1 | "
            "ForEach-Object { ($_.MediaType.ToString()) + '|' + ([math]::Round($_.Size / 1GB).ToString() + ' GB') }"
        )
    )
    disk_type = ""
    disk_size = ""
    if "|" in disk_info:
        disk_type, disk_size = [part.strip() for part in disk_info.split("|", 1)]
        if disk_type.upper() == "SSD":
            disk_type = "SATA SSD"
    optical = first_line(powershell_value("(Get-CimInstance Win32_CDROMDrive | Select-Object -First 1 -ExpandProperty Name)"))
    video = first_line(powershell_value("(Get-CimInstance Win32_VideoController | Select-Object -First 1 -ExpandProperty Name)"))
    resolution = first_line(
        powershell_value(
            "Get-CimInstance Win32_VideoController | Select-Object -First 1 | "
            "ForEach-Object { if ($_.CurrentHorizontalResolution) { "
            "'{0}x{1}' -f $_.CurrentHorizontalResolution, $_.CurrentVerticalResolution } }"
        )
    ) or screen_resolution()

    return {
        "Computer name": socket.gethostname(),
        "Current user": getpass.getuser(),
        "OS": os_caption,
        "OS version": os_version,
        "OS architecture": first_line(powershell_value("(Get-CimInstance Win32_OperatingSystem).OSArchitecture")),
        "Domain / Workgroup": windows_domain_or_workgroup(),
        "Product type": product_type,
        "Manufacturer": manufacturer,
        "Model": model,
        "Chassis type": product_type,
        "Part number": part_number,
        "Serial number": serial,
        "Motherboard": motherboard,
        "BIOS version": bios_version,
        "Secure Boot": windows_secure_boot(),
        "TPM": windows_tpm_status(),
        "Processor": processor,
        "RAM size": ram_size,
        "RAM slots used": windows_ram_slots(),
        "Disk type": disk_type,
        "Disk size": disk_size,
        "Disk details": windows_disk_details(),
        "Volumes": windows_volume_details(),
        "Optical drive": optical,
        "Video model": video,
        "LCD resolution": resolution,
        "LCD type": infer_screen_type_from_text(windows_monitor_text()),
        "Keyboard layout": locale_label(),
        "Network adapters": windows_network_details(),
        "Battery": windows_battery_details(),
    }


def linux_disk_info():
    output = run_command(["lsblk", "-bdno", "NAME,SIZE,ROTA,TYPE"])
    if output:
        for line in output.splitlines():
            parts = line.split()
            if len(parts) >= 4 and parts[3] == "disk":
                disk_type = "SATA HDD" if parts[2] == "1" else "SATA SSD"
                return disk_type, format_gb(parts[1])

    for name in os.listdir("/sys/block") if os.path.isdir("/sys/block") else []:
        if name.startswith(("loop", "ram", "fd")):
            continue
        size = read_text_file(f"/sys/block/{name}/size")
        rotational = read_text_file(f"/sys/block/{name}/queue/rotational")
        if size:
            disk_type = "SATA HDD" if rotational == "1" else "SATA SSD"
            try:
                return disk_type, format_gb(int(size) * 512)
            except ValueError:
                continue

    return "", ""


def linux_cpu_name():
    cpuinfo = read_text_file("/proc/cpuinfo")
    for line in cpuinfo.splitlines():
        if line.lower().startswith("model name"):
            return line.split(":", 1)[1].strip()
    return ""


def linux_ram_size():
    meminfo = read_text_file("/proc/meminfo")
    for line in meminfo.splitlines():
        if line.startswith("MemTotal:"):
            parts = line.split()
            if len(parts) >= 2:
                return format_gb(int(parts[1]) * 1024)
    return ""


def linux_video_model():
    output = run_command(["lspci"])
    for line in output.splitlines():
        lower = line.lower()
        if "vga compatible controller" in lower or "3d controller" in lower or "display controller" in lower:
            return line.split(":", 2)[-1].strip()
    return ""


def linux_optical_drive():
    info = read_text_file("/proc/sys/dev/cdrom/info")
    for line in info.splitlines():
        if line.startswith("drive name:"):
            value = line.split(":", 1)[1].strip()
            return value or ""
    return ""


def linux_inventory():
    manufacturer = read_text_file("/sys/class/dmi/id/sys_vendor")
    model = read_text_file("/sys/class/dmi/id/product_name")
    part_number = read_text_file("/sys/class/dmi/id/product_sku")
    serial = clean_serial_number(read_text_file("/sys/class/dmi/id/product_serial"))
    if not serial:
        serial = clean_serial_number(read_text_file("/sys/class/dmi/id/board_serial"))
    chassis = read_text_file("/sys/class/dmi/id/chassis_type")
    product_type = chassis_label(chassis)
    disk_type, disk_size = linux_disk_info()
    os_release = read_text_file("/etc/os-release")
    os_pretty = ""
    for line in os_release.splitlines():
        if line.startswith("PRETTY_NAME="):
            os_pretty = line.split("=", 1)[1].strip().strip('"')
            break

    return {
        "Computer name": socket.gethostname(),
        "Current user": getpass.getuser(),
        "OS": linux_command_lines(["uname", "-srmo"], limit=1),
        "OS version": os_pretty,
        "Product type": product_type,
        "Manufacturer": manufacturer,
        "Model": model,
        "Chassis type": product_type,
        "Part number": part_number,
        "Serial number": serial,
        "Motherboard": (read_text_file("/sys/class/dmi/id/board_vendor") + " " + read_text_file("/sys/class/dmi/id/board_name")).strip(),
        "BIOS version": read_text_file("/sys/class/dmi/id/bios_version"),
        "Processor": linux_cpu_name(),
        "RAM size": linux_ram_size(),
        "RAM slots used": linux_ram_slots(),
        "Disk type": disk_type,
        "Disk size": disk_size,
        "Disk details": linux_command_lines(["lsblk", "-o", "NAME,SIZE,TYPE,MODEL"], limit=10),
        "Volumes": linux_command_lines(["df", "-h", "--output=source,size,used,avail,pcent,target"], limit=10),
        "Optical drive": linux_optical_drive(),
        "Video model": linux_video_model(),
        "LCD resolution": screen_resolution(),
        "LCD type": infer_screen_type_from_text(linux_monitor_text()),
        "Keyboard layout": locale_label(),
        "Network adapters": linux_command_lines(["ip", "-brief", "addr"], limit=10),
    }


def write_inventory_debug_log(title):
    try:
        with open(INVENTORY_DEBUG_LOG, "a", encoding="utf-8") as handle:
            handle.write(f"\n[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {title}\n")
            handle.write(traceback.format_exc())
    except OSError:
        pass


def collect_system_inventory():
    try:
        if platform.system().lower().startswith("win"):
            data = windows_inventory()
        elif platform.system().lower() == "linux":
            data = linux_inventory()
        else:
            data = {"LCD resolution": screen_resolution(), "Keyboard layout": locale_label()}
        return {key: value for key, value in data.items() if value and str(value).strip()}
    except Exception:
        write_inventory_debug_log("inventory error")
        return {}


# =========================
# RECORD FUNCTIONS
# =========================
def clear_form():
    global selected_id, pc_audit_batch_started
    selected_id = None
    pc_audit_batch_started = False

    record_xenoglossi_entry.delete(0, tk.END)
    record_username_entry.delete(0, tk.END)
    record_password_entry.delete(0, tk.END)
    dropdown.set("")
    checkbox_var.set(False)

    for entry in entries:
        clear_field(entry)

    if tree.selection():
        tree.selection_remove(tree.selection())


def load_auditor_sample():
    for entry in entries:
        clear_field(entry)

    for index, field_name in enumerate(AUDITOR_FIELDS):
        set_field(entries[index], AUDITOR_SAMPLE_VALUES.get(field_name, ""))

    if not record_xenoglossi_entry.get().strip():
        record_xenoglossi_entry.insert(0, AUDITOR_SAMPLE_VALUES["Serial number"])

    if not record_username_entry.get().strip():
        record_username_entry.insert(0, AUDITOR_SAMPLE_VALUES["Serial number"])

    if not dropdown.get().strip():
        dropdown.set("Option 1")


def load_current_pc_info():
    if not require_login():
        return

    inventory = collect_system_inventory()

    if not inventory:
        messagebox.showwarning("Read This PC", "No system details could be detected on this machine")
        return

    updated = 0
    for index, field_name in enumerate(AUDITOR_FIELDS):
        value = inventory.get(field_name, "")
        if value:
            set_field(entries[index], value)
            updated += 1

    serial = inventory.get("Serial number", "")
    if serial and not record_xenoglossi_entry.get().strip():
        record_xenoglossi_entry.insert(0, serial)
    if serial and not record_username_entry.get().strip():
        record_username_entry.insert(0, serial)

    log_action("READ_THIS_PC", f"Detected {updated} field(s) from local system")
    message = f"Detected and filled {updated} field(s)"
    if not serial:
        message += (
            "\n\nNo valid machine serial number was reported by BIOS/SystemProduct/BaseBoard. "
            "The Serial number field was left blank."
        )
    messagebox.showinfo("Read This PC", message)


def finish_read_current_pc_info(inventory, error_message):
    root.configure(cursor="")

    if error_message:
        messagebox.showerror(
            "Read This PC",
            f"Could not read this PC. Details were written to inventory_debug.log.\n\n{error_message}",
        )
        return

    if not inventory:
        messagebox.showwarning("Read This PC", "No system details could be detected on this machine")
        return

    updated = 0
    for index, field_name in enumerate(AUDITOR_FIELDS):
        value = inventory.get(field_name, "")
        if value:
            set_field(entries[index], value)
            updated += 1

    serial = inventory.get("Serial number", "")
    if serial and not record_xenoglossi_entry.get().strip():
        record_xenoglossi_entry.insert(0, serial)
    if serial and not record_username_entry.get().strip():
        record_username_entry.insert(0, serial)

    log_action("READ_THIS_PC", f"Detected {updated} field(s) from local system")
    message = f"Detected and filled {updated} field(s)"
    if not serial:
        message += (
            "\n\nNo valid machine serial number was reported by BIOS/SystemProduct/BaseBoard. "
            "The Serial number field was left blank."
        )
    messagebox.showinfo("Read This PC", message)


def start_read_current_pc_info():
    if not require_login():
        return

    root.configure(cursor="watch")
    root.update_idletasks()

    def worker():
        try:
            inventory = collect_system_inventory()
            error_message = ""
        except Exception as exc:
            inventory = {}
            error_message = str(exc)
            write_inventory_debug_log("read this pc error")
        root.after(0, lambda: finish_read_current_pc_info(inventory, error_message))

    threading.Thread(target=worker, daemon=True).start()


def load_data(filter_text=""):
    clear_tree(tree)

    conn = get_conn()
    c = conn.cursor()

    if filter_text:
        c.execute(
            f"SELECT {RECORD_SELECT_COLUMNS} FROM records WHERE username LIKE ? OR xenoglossi LIKE ? ORDER BY id DESC",
            (f"%{filter_text}%", f"%{filter_text}%"),
        )
    else:
        c.execute(f"SELECT {RECORD_SELECT_COLUMNS} FROM records ORDER BY id DESC")

    for row in c.fetchall():
        tree.insert("", "end", values=(row[0], row[1], row[2], row[3], row[5], "Yes" if row[6] else "No"))

    conn.close()


def on_select(event):
    global selected_id, pc_audit_batch_started

    item = tree.focus()
    if not item:
        return

    selected_id = tree.item(item)["values"][0]
    pc_audit_batch_started = False

    conn = get_conn()
    c = conn.cursor()
    c.execute(f"SELECT {RECORD_SELECT_COLUMNS} FROM records WHERE id=?", (selected_id,))
    row = c.fetchone()
    conn.close()

    if not row:
        return

    record_xenoglossi_entry.delete(0, tk.END)
    record_xenoglossi_entry.insert(0, row[1])

    record_username_entry.delete(0, tk.END)
    record_username_entry.insert(0, row[3])

    record_password_entry.delete(0, tk.END)

    dropdown.set(row[5])
    checkbox_var.set(bool(row[6]))

    for entry in entries:
        clear_field(entry)

    values = normalize_input_values(row[7].split("|") if row[7] else [])
    for i, val in enumerate(values):
        set_field(entries[i], val)

    notebook.select(form_frame)


tree.bind("<<TreeviewSelect>>", on_select)


def validate_form():
    if not require_login():
        return False

    if not pc_audit_batch_started:
        messagebox.showerror("Validation", "Run PC Audit is required before saving")
        return False

    xenoglossi = record_xenoglossi_entry.get().strip()
    if not xenoglossi:
        messagebox.showerror("Validation", "Asset Tag is required")
        return False

    if not record_username_entry.get().strip():
        messagebox.showerror("Validation", "Record username is required")
        return False

    if selected_id is None and not record_password_entry.get().strip():
        messagebox.showerror("Validation", "Password is required for a new record")
        return False

    serial_number = get_field(entries[SERIAL_NUMBER_INDEX]).strip()
    if not serial_number:
        messagebox.showerror("Validation", "Serial number is required")
        entries[SERIAL_NUMBER_INDEX].focus_set()
        return False

    if not dropdown.get().strip():
        messagebox.showerror("Validation", "Please select an option")
        return False

    return True


def submit():
    global selected_id

    if not validate_form():
        return

    conn = get_conn()
    c = conn.cursor()

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    xenoglossi = record_xenoglossi_entry.get().strip()
    username = record_username_entry.get().strip()
    dropdown_value = dropdown.get().strip()
    checkbox_value = int(checkbox_var.get())
    inputs_value = "|".join([get_field(entry) for entry in entries])

    if selected_id:
        if record_password_entry.get().strip():
            password_value = hash_password(record_password_entry.get().strip())
            c.execute(
                "UPDATE records SET xenoglossi=?, timestamp=?, username=?, password=?, dropdown=?, checkbox=?, inputs=? WHERE id=?",
                (xenoglossi, timestamp, username, password_value, dropdown_value, checkbox_value, inputs_value, selected_id),
            )
        else:
            c.execute(
                "UPDATE records SET xenoglossi=?, timestamp=?, username=?, dropdown=?, checkbox=?, inputs=? WHERE id=?",
                (xenoglossi, timestamp, username, dropdown_value, checkbox_value, inputs_value, selected_id),
            )
        action = "RECORD_UPDATE"
        details = f"Updated record {xenoglossi} ({username})"
    else:
        password_value = hash_password(record_password_entry.get().strip())
        c.execute(
            "INSERT INTO records (xenoglossi, timestamp, username, password, dropdown, checkbox, inputs) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (xenoglossi, timestamp, username, password_value, dropdown_value, checkbox_value, inputs_value),
        )
        selected_id = c.lastrowid
        action = "RECORD_CREATE"
        details = f"Created record {xenoglossi} ({username})"

    conn.commit()
    conn.close()

    log_action(action, details)
    load_data()
    populate_audit_client_filter()
    if get_setting("auto_refresh", "1") == "1":
        refresh_analytics()
        build_report(show_message=False)
    clear_form()
    messagebox.showinfo("Success", "Record saved successfully")


def delete():
    global selected_id

    if not require_login():
        return

    if not selected_id:
        messagebox.showwarning("Delete", "Please select a record first")
        return

    confirm = messagebox.askyesno("Delete", "Are you sure you want to delete this record?")
    if not confirm:
        return

    deleted_id = selected_id
    conn = get_conn()
    c = conn.cursor()
    c.execute("DELETE FROM records WHERE id=?", (selected_id,))
    conn.commit()
    conn.close()

    log_action("RECORD_DELETE", f"Deleted record #{deleted_id}")
    load_data()
    populate_audit_client_filter()
    if get_setting("auto_refresh", "1") == "1":
        refresh_analytics()
        build_report(show_message=False)
    clear_form()
    messagebox.showinfo("Delete", "Record deleted")


def export_excel():
    if not require_login():
        return

    conn = get_conn()
    c = conn.cursor()
    c.execute(f"SELECT {RECORD_SELECT_COLUMNS} FROM records ORDER BY id DESC")
    data = c.fetchall()
    conn.close()

    if not data:
        messagebox.showwarning("Export", "No data to export")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Records"

    headers = ["Asset Tag", "Timestamp", "Username", "Dropdown", "Checkbox"] + AUDITOR_FIELDS
    ws.append(headers)

    for row in data:
        input_values = normalize_input_values(row[7].split("|") if row[7] else [])
        ws.append([row[1], row[2], row[3], row[5], row[6]] + input_values)

    path = ask_save_path(
        ".xlsx",
        [("Excel files", "*.xlsx")],
        "Save Excel File",
        f"records_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
    )

    if path:
        wb.save(path)
        log_action("EXPORT_EXCEL", f"Exported records to {path}")
        refresh_analytics()
        messagebox.showinfo("Export", "Excel file exported successfully")


def export_json():
    if not require_login():
        return

    conn = get_conn()
    c = conn.cursor()
    c.execute(f"SELECT {RECORD_SELECT_COLUMNS} FROM records ORDER BY id DESC")
    data = c.fetchall()
    conn.close()

    if not data:
        messagebox.showwarning("Export", "No data to export")
        return

    json_data = []

    for row in data:
        input_values = normalize_input_values(row[7].split("|") if row[7] else [])

        record = {
            "internal_id": row[0],
            "asset_tag": row[1],
            "timestamp": row[2],
            "username": row[3],
            "password": "hidden",
            "dropdown": row[5],
            "checkbox": bool(row[6]),
            "fields": {},
        }

        for i, field_name in enumerate(AUDITOR_FIELDS):
            record["fields"][field_name] = input_values[i]

        json_data.append(record)

    path = ask_save_path(
        ".json",
        [("JSON files", "*.json")],
        "Save JSON File",
        f"records_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json",
    )

    if path:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(json_data, f, indent=4, ensure_ascii=False)

        log_action("EXPORT_JSON", f"Exported records to {path}")
        refresh_analytics()
        messagebox.showinfo("Export", "JSON file exported successfully")


# =========================
# ANALYTICS FUNCTIONS
# =========================
def refresh_analytics():
    if "kpi_vars" not in globals():
        return

    conn = get_conn()
    c = conn.cursor()

    c.execute("SELECT COUNT(*) FROM records")
    total_records = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM records WHERE checkbox=1")
    accepted_terms = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM app_users WHERE status='Active'")
    active_users = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM app_logs WHERE action LIKE 'EXPORT_%'")
    exports = c.fetchone()[0]

    kpi_vars["Total Records"].config(text=str(total_records))
    kpi_vars["Accepted Terms"].config(text=str(accepted_terms))
    kpi_vars["Active Users"].config(text=str(active_users))
    kpi_vars["Exports"].config(text=str(exports))

    clear_tree(option_tree)
    c.execute("SELECT dropdown, COUNT(*) FROM records GROUP BY dropdown ORDER BY COUNT(*) DESC, dropdown")
    option_rows = c.fetchall()
    if option_rows:
        for row in option_rows:
            option_tree.insert("", "end", values=(row[0] or "(blank)", row[1]))
    else:
        option_tree.insert("", "end", values=("No data", 0))

    clear_tree(activity_tree)
    c.execute("SELECT timestamp, username, action FROM app_logs ORDER BY id DESC LIMIT 12")
    for row in c.fetchall():
        activity_tree.insert("", "end", values=row)

    conn.close()


# =========================
# REPORT FUNCTIONS
# =========================
def get_report_rows():
    date_from = parse_iso_date(report_from_var.get(), "From")
    if date_from is False:
        return None
    date_to = parse_iso_date(report_to_var.get(), "To", end_of_day=True)
    if date_to is False:
        return None

    query = f"SELECT {RECORD_SELECT_COLUMNS} FROM records"
    clauses = []
    params = []

    username = report_username_var.get().strip()
    if username:
        clauses.append("username LIKE ?")
        params.append(f"%{username}%")

    asset_tag = report_asset_tag_var.get().strip()
    if asset_tag:
        clauses.append("xenoglossi LIKE ?")
        params.append(f"%{asset_tag}%")

    option = report_option_var.get().strip()
    if option:
        clauses.append("dropdown = ?")
        params.append(option)

    if date_from:
        clauses.append("timestamp >= ?")
        params.append(date_from)

    if date_to:
        clauses.append("timestamp <= ?")
        params.append(date_to)

    if clauses:
        query += " WHERE " + " AND ".join(clauses)

    query += " ORDER BY id DESC"

    conn = get_conn()
    c = conn.cursor()
    c.execute(query, params)
    rows = c.fetchall()
    conn.close()
    return rows


def build_report(show_message=True):
    if "report_tree" not in globals():
        return
    if not is_logged_in or current_role != "Admin":
        return

    rows = get_report_rows()
    if rows is None:
        return

    clear_tree(report_tree)

    accepted = 0
    filled_inputs = 0
    for row in rows:
        input_values = normalize_input_values(row[7].split("|") if row[7] else [])
        filled = len([value for value in input_values if value.strip()])
        filled_inputs += filled
        accepted += 1 if row[6] else 0
        report_tree.insert(
            "",
            "end",
            values=(row[1], row[2], row[3], row[5], "Yes" if row[6] else "No", filled),
        )

    report_summary_var.set(
        f"Rows: {len(rows)} | Accepted: {accepted} | Filled inputs: {filled_inputs}"
    )

    if show_message:
        log_action("REPORT_BUILD", f"Built report with {len(rows)} rows")


def export_report_excel():
    if not require_admin():
        return

    rows = get_report_rows()
    if rows is None:
        return

    if not rows:
        messagebox.showwarning("Report", "No report data to export")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Asset Tag", "Timestamp", "Username", "Dropdown", "Checkbox"] + AUDITOR_FIELDS)

    for row in rows:
        input_values = normalize_input_values(row[7].split("|") if row[7] else [])
        ws.append([row[1], row[2], row[3], row[5], row[6]] + input_values)

    path = ask_save_path(
        ".xlsx",
        [("Excel files", "*.xlsx")],
        "Save Report",
        f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
    )

    if path:
        wb.save(path)
        log_action("EXPORT_REPORT", f"Exported report to {path}")
        refresh_analytics()
        messagebox.showinfo("Report", "Report exported successfully")


# =========================
# NETWORK JOB FUNCTIONS
# =========================
def load_network_settings_ui():
    if "network_server_url_var" not in globals():
        return

    network_server_url_var.set(get_setting("network_server_url", "http://127.0.0.1:8765"))
    workstation_name_var.set(get_setting("workstation_name", os.environ.get("COMPUTERNAME", "CLIENT-1")))
    if not job_assignee_var.get().strip():
        job_assignee_var.set("ALL")


def save_network_settings(show_message=True):
    if not require_login():
        return False

    server_url = network_server_url_var.get().strip().rstrip("/")
    workstation = workstation_name_var.get().strip()

    if not server_url:
        messagebox.showerror("Network Jobs", "Server URL is required")
        return False

    if not server_url.startswith(("http://", "https://")):
        server_url = "http://" + server_url
        network_server_url_var.set(server_url)

    if not workstation:
        messagebox.showerror("Network Jobs", "This Client is required")
        return False

    set_setting("network_server_url", server_url)
    set_setting("workstation_name", workstation)

    if show_message:
        log_action("NETWORK_SETTINGS_SAVE", f"Saved network settings for {workstation}")
        messagebox.showinfo("Network Jobs", "Network settings saved")

    return True


def network_base_url():
    value = network_server_url_var.get().strip() or get_setting("network_server_url", "http://127.0.0.1:8765")
    value = value.rstrip("/")
    if not value.startswith(("http://", "https://")):
        value = "http://" + value
    return value


def network_request(method, path, data=None, params=None, auth=False, timeout=8):
    url = network_base_url() + path
    if params:
        url += "?" + urlparse.urlencode(params)

    body = None
    headers = {"Accept": "application/json"}
    if data is not None:
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json; charset=utf-8"

    req = urlrequest.Request(url, data=body, headers=headers, method=method)

    try:
        with urlrequest.urlopen(req, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urlerror.HTTPError as exc:
        raw = exc.read().decode("utf-8", errors="replace")
        try:
            payload = json.loads(raw)
            message = payload.get("error", raw)
        except json.JSONDecodeError:
            message = raw or str(exc)
        raise RuntimeError(message) from exc
    except urlerror.URLError as exc:
        raise RuntimeError(str(exc.reason)) from exc
    except TimeoutError as exc:
        raise RuntimeError("Network request timed out") from exc


def set_network_details(text):
    if "network_details_text" not in globals():
        return
    network_details_text.configure(state="normal")
    network_details_text.delete("1.0", tk.END)
    network_details_text.insert("1.0", text)
    network_details_text.configure(state="disabled")


def test_network_connection():
    if not save_network_settings(show_message=False):
        return

    try:
        result = network_request("GET", "/health", auth=False, timeout=5)
    except RuntimeError as exc:
        messagebox.showerror("Network Jobs", f"Connection failed: {exc}")
        return

    if result.get("ok"):
        register_this_client(show_message=False)
        fetch_network_clients(show_message=False)
        messagebox.showinfo("Network Jobs", "Connection successful")
    else:
        messagebox.showerror("Network Jobs", "Server did not return a healthy response")


def client_platform_label():
    return f"{platform.system()} {platform.release()} ({platform.machine()})"


def client_network_identity(client):
    inventory = client.get("inventory", {}) if isinstance(client, dict) else {}
    network_text = str(inventory.get("Network adapters", "") if isinstance(inventory, dict) else "")
    mac_match = re.search(r"\b(?:[0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}\b", network_text)
    ip_match = re.search(r"\b(?:\d{1,3}\.){3}\d{1,3}\b", network_text)
    return (
        ip_match.group(0) if ip_match else "",
        mac_match.group(0).upper() if mac_match else "",
    )


def register_this_client(show_message=True):
    if not require_login():
        return False

    if not save_network_settings(show_message=False):
        return False

    client_name = workstation_name_var.get().strip()
    inventory = collect_system_inventory()
    try:
        from register_client import collect_device_diagnostics

        inventory["Diagnostics"] = collect_device_diagnostics()
    except Exception as exc:
        inventory["Diagnostics"] = {
            "items": [
                {
                    "name": "Diagnostics",
                    "status": "FAIL",
                    "details": f"Could not collect diagnostics from app registration: {exc}",
                }
            ],
            "summary": {"FAIL": 1},
        }

    try:
        result = network_request(
            "POST",
            "/clients/register",
            data={
                "name": client_name,
                "username": current_user or "",
                "role": current_role or "",
                "app_version": APP_VERSION,
                "platform": client_platform_label(),
                "inventory": inventory,
            },
            auth=True,
            timeout=8,
        )
    except RuntimeError as exc:
        if show_message:
            messagebox.showerror("Network Jobs", f"Could not register client: {exc}")
        return False

    if result.get("ok"):
        log_action("NETWORK_CLIENT_REGISTER", f"Registered client {client_name}")
        if show_message:
            messagebox.showinfo("Network Jobs", f"Client registered: {client_name}")
        return True

    if show_message:
        messagebox.showerror("Network Jobs", "Server did not accept the client registration")
    return False


def fetch_network_clients(show_message=True):
    global network_clients

    if not require_login():
        return

    if not save_network_settings(show_message=False):
        return

    try:
        result = network_request("GET", "/clients", auth=True, timeout=8)
    except RuntimeError as exc:
        if show_message:
            messagebox.showerror("Network Jobs", f"Could not fetch clients: {exc}")
        return

    clients = result.get("clients", [])
    network_clients = {client.get("name", ""): client for client in clients}

    clear_tree(network_client_tree)
    client_names = ["ALL"]

    for client in clients:
        name = client.get("name", "")
        client_ip, client_mac = client_network_identity(client)
        if name:
            client_names.append(name)
        network_client_tree.insert(
            "",
            "end",
            values=(
                name,
                client_ip,
                client_mac,
                client.get("username", ""),
                client.get("role", ""),
                client.get("last_seen", ""),
                client.get("platform", ""),
                client.get("app_version", ""),
            ),
        )

    job_assignee_combo.configure(values=client_names)

    if job_assignee_var.get().strip() not in client_names:
        job_assignee_var.set("ALL")

    populate_device_results()
    populate_manual_tests()
    populate_audit_client_filter()

    if show_message:
        messagebox.showinfo("Network Jobs", f"Fetched {len(clients)} client(s)")


def diagnostic_rows_from_clients():
    rows = []

    for client_name, client in sorted(network_clients.items()):
        inventory = client.get("inventory", {})
        diagnostics = inventory.get("Diagnostics", {}) if isinstance(inventory, dict) else {}
        items = diagnostics.get("items", []) if isinstance(diagnostics, dict) else []

        if not items:
            rows.append(
                {
                    "client": client_name,
                    "device": "Diagnostics",
                    "status": "UNKNOWN",
                    "details": "No diagnostics reported by this client",
                    "last_seen": client.get("last_seen", ""),
                    "raw": client,
                }
            )
            continue

        for item in items:
            if not isinstance(item, dict):
                continue
            rows.append(
                {
                    "client": client_name,
                    "device": item.get("name", ""),
                    "status": item.get("status", "UNKNOWN"),
                    "details": item.get("details", ""),
                    "last_seen": client.get("last_seen", ""),
                    "raw": item,
                }
            )

    return rows


def set_device_details(text):
    if "device_details_text" not in globals():
        return
    device_details_text.configure(state="normal")
    device_details_text.delete("1.0", tk.END)
    device_details_text.insert("1.0", text)
    device_details_text.configure(state="disabled")


def set_audit_archive_details(text):
    if "audit_archive_details_text" not in globals():
        return
    audit_archive_details_text.configure(state="normal")
    audit_archive_details_text.delete("1.0", tk.END)
    audit_archive_details_text.insert("1.0", text)
    audit_archive_details_text.configure(state="disabled")


def record_row_to_local_audit(row):
    input_values = normalize_input_values(row[7].split("|") if row[7] else [])
    inventory = {
        field_name: value
        for field_name, value in zip(AUDITOR_FIELDS, input_values)
        if str(value or "").strip()
    }
    inventory["Record ID"] = row[0]
    inventory["Record Username"] = row[3]
    inventory["Record Dropdown"] = row[5]
    inventory["Record Accepted"] = "Yes" if row[6] else "No"

    client_base = (
        inventory.get("Computer name")
        or inventory.get("Serial number")
        or row[1]
        or row[3]
        or f"Record {row[0]}"
    )
    return {
        "id": f"local-{row[0]}",
        "source": "Local Record",
        "record_id": row[0],
        "client_name": f"Local - {client_base}",
        "username": row[3],
        "role": "Record",
        "app_version": APP_VERSION,
        "platform": "Local records database",
        "inventory": inventory,
        "diagnostics_summary": {},
        "created_at": row[2],
    }


def local_record_audits(limit=500, client_filter=""):
    conn = get_conn()
    c = conn.cursor()
    c.execute(f"SELECT {RECORD_SELECT_COLUMNS} FROM records ORDER BY id DESC")
    rows = c.fetchall()
    conn.close()

    audits = []
    client_filter = str(client_filter or "").strip()
    for row in rows:
        audit = record_row_to_local_audit(row)
        if client_filter and audit["client_name"] != client_filter:
            continue
        audits.append(audit)
        if limit and len(audits) >= limit:
            break
    return audits


def local_record_client_names():
    return sorted({audit["client_name"] for audit in local_record_audits(limit=0) if audit.get("client_name")})


def populate_audit_client_filter():
    if "audit_archive_client_combo" not in globals():
        return
    names = [""] + sorted(set(name for name in network_clients if name) | set(local_record_client_names()))
    audit_archive_client_combo.configure(values=names)


def fetch_audit_archive(show_message=True):
    global audit_archive_records

    if not require_login():
        return

    try:
        limit = int(audit_archive_limit_var.get().strip() or "500")
    except ValueError:
        messagebox.showerror("Audit Archive", "Limit must be a number")
        return
    limit = max(1, min(limit, 5000))
    audit_archive_limit_var.set(str(limit))

    params = {"limit": str(limit)}
    client = audit_archive_client_var.get().strip()
    if client:
        params["client"] = client

    local_audits = local_record_audits(limit=limit, client_filter=client)
    network_audits = []
    network_error = ""

    if save_network_settings(show_message=False):
        try:
            result = network_request("GET", "/audits", params=params, auth=True, timeout=10)
            network_audits = result.get("audits", [])
        except RuntimeError as exc:
            network_error = str(exc)
    else:
        network_error = "Network settings are incomplete"

    audits = local_audits + network_audits
    audit_archive_records = {str(audit.get("id")): audit for audit in audits}
    clear_tree(audit_archive_tree)

    totals = {"OK": 0, "WARN": 0, "FAIL": 0, "UNKNOWN": 0}
    for audit in audits:
        summary = audit.get("diagnostics_summary", {}) if isinstance(audit, dict) else {}
        if not isinstance(summary, dict):
            summary = {}
        for key in totals:
            totals[key] += int(summary.get(key, 0) or 0)
        audit_archive_tree.insert(
            "",
            "end",
            values=(
                audit.get("id", ""),
                audit.get("client_name", ""),
                audit.get("created_at", ""),
                audit.get("platform", ""),
                summary.get("OK", 0),
                summary.get("WARN", 0),
                summary.get("FAIL", 0),
                summary.get("UNKNOWN", 0),
            ),
        )

    audit_archive_summary_var.set(
        f"Audits: {len(audits)} | Local: {len(local_audits)} | Network: {len(network_audits)} | "
        f"OK: {totals['OK']} | WARN: {totals['WARN']} | "
        f"FAIL: {totals['FAIL']} | UNKNOWN: {totals['UNKNOWN']}"
    )
    if show_message:
        note = f"\n\nNetwork unavailable: {network_error}" if network_error else ""
        messagebox.showinfo("Audit Archive", f"Loaded {len(audits)} audit record(s){note}")


def selected_audit_id():
    item = audit_archive_tree.focus()
    if not item:
        return ""
    values = audit_archive_tree.item(item).get("values", [])
    if not values:
        return ""
    return str(values[0])


def fetch_audit_detail(audit_id):
    if not audit_id:
        return None
    audit = audit_archive_records.get(str(audit_id), {})
    if audit.get("inventory"):
        return audit
    result = network_request("GET", f"/audits/{audit_id}", auth=True, timeout=10)
    detail = result.get("audit", {})
    if detail:
        audit_archive_records[str(audit_id)] = detail
    return detail


def on_audit_archive_select(event=None):
    audit_id = selected_audit_id()
    if not audit_id:
        return
    try:
        audit = fetch_audit_detail(audit_id)
    except RuntimeError as exc:
        messagebox.showerror("Audit Archive", f"Could not load audit detail: {exc}")
        return
    set_audit_archive_details(json.dumps(audit or {}, ensure_ascii=False, indent=2))


def load_selected_audit_to_form():
    audit_id = selected_audit_id()
    if not audit_id:
        messagebox.showwarning("Audit Archive", "Select an audit first")
        return
    try:
        audit = fetch_audit_detail(audit_id)
    except RuntimeError as exc:
        messagebox.showerror("Audit Archive", f"Could not load audit detail: {exc}")
        return
    inventory = audit.get("inventory", {}) if isinstance(audit, dict) else {}
    if not inventory:
        messagebox.showwarning("Audit Archive", "Selected audit has no inventory details")
        return

    clear_form()
    for index, field_name in enumerate(AUDITOR_FIELDS):
        value = inventory.get(field_name, "")
        if value:
            set_field(entries[index], value)

    serial = inventory.get("Serial number", "")
    record_xenoglossi_entry.insert(0, serial or audit.get("client_name", ""))
    record_username_entry.insert(0, serial or audit.get("client_name", ""))
    notebook.select(form_frame)


def audit_ids_for_report():
    selected = selected_audit_id()
    if selected:
        return [selected]
    ids = []
    for item in audit_archive_tree.get_children():
        values = audit_archive_tree.item(item).get("values", [])
        if values:
            ids.append(str(values[0]))
    return ids


def audit_diagnostic_items(audit):
    inventory = audit.get("inventory", {}) if isinstance(audit, dict) else {}
    diagnostics = inventory.get("Diagnostics", {}) if isinstance(inventory, dict) else {}
    items = diagnostics.get("items", []) if isinstance(diagnostics, dict) else []
    return [item for item in items if isinstance(item, dict)]


def html_escape(value):
    return html.escape(str(value or ""), quote=True)


def status_class(status):
    status = str(status or "UNKNOWN").upper()
    if status in {"OK", "WARN", "FAIL"}:
        return status.lower()
    return "unknown"


def printable_report_folder():
    folder = os.path.join(get_default_export_folder(), "print_reports")
    os.makedirs(folder, exist_ok=True)
    return folder


def build_audit_html_report(audits):
    generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sections = []

    for audit in audits:
        inventory = audit.get("inventory", {}) if isinstance(audit.get("inventory", {}), dict) else {}
        summary = audit.get("diagnostics_summary", {}) if isinstance(audit.get("diagnostics_summary", {}), dict) else {}
        diagnostics = audit_diagnostic_items(audit)
        serial_number = inventory.get("Serial number", "")
        part_number = inventory.get("Part number", "")
        manufacturer = inventory.get("Manufacturer", "")
        model = inventory.get("Model", "")

        inventory_rows = []
        for field, value in inventory.items():
            if field == "Diagnostics":
                continue
            inventory_rows.append(
                f"<tr><th>{html_escape(field)}</th><td>{html_escape(value)}</td></tr>"
            )

        diagnostic_rows = []
        if diagnostics:
            for item in diagnostics:
                status = str(item.get("status", "UNKNOWN")).upper()
                diagnostic_rows.append(
                    "<tr>"
                    f"<td>{html_escape(item.get('name', ''))}</td>"
                    f"<td><span class=\"badge {status_class(status)}\">{html_escape(status)}</span></td>"
                    f"<td>{html_escape(item.get('details', ''))}</td>"
                    "</tr>"
                )
        else:
            diagnostic_rows.append("<tr><td colspan=\"3\">No diagnostics recorded</td></tr>")

        sections.append(
            f"""
            <section class="audit">
              <div class="audit-header">
                <div>
                  <h2>{html_escape(audit.get('client_name', 'Audit'))}</h2>
                  <p>Audit #{html_escape(audit.get('id', ''))} | {html_escape(audit.get('created_at', ''))}</p>
                </div>
                <div class="summary">
                  <span class="badge ok">OK {html_escape(summary.get('OK', 0))}</span>
                  <span class="badge warn">WARN {html_escape(summary.get('WARN', 0))}</span>
                  <span class="badge fail">FAIL {html_escape(summary.get('FAIL', 0))}</span>
                  <span class="badge unknown">UNKNOWN {html_escape(summary.get('UNKNOWN', 0))}</span>
                </div>
              </div>
              <div class="meta">
                <div><strong>Platform</strong><br>{html_escape(audit.get('platform', ''))}</div>
                <div><strong>User</strong><br>{html_escape(audit.get('username', ''))}</div>
                <div><strong>Role</strong><br>{html_escape(audit.get('role', ''))}</div>
                <div><strong>App</strong><br>{html_escape(audit.get('app_version', ''))}</div>
                <div><strong>Manufacturer</strong><br>{html_escape(manufacturer)}</div>
                <div><strong>Model</strong><br>{html_escape(model)}</div>
                <div><strong>Serial number</strong><br>{html_escape(serial_number or 'Not available')}</div>
                <div><strong>Part number</strong><br>{html_escape(part_number or 'Not available')}</div>
              </div>
              <h3>Inventory</h3>
              <table class="kv"><tbody>{''.join(inventory_rows)}</tbody></table>
              <h3>Diagnostics</h3>
              <table>
                <thead><tr><th>Check</th><th>Status</th><th>Details</th></tr></thead>
                <tbody>{''.join(diagnostic_rows)}</tbody>
              </table>
            </section>
            """
        )

    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>PC Audit Report</title>
  <style>
    :root {{
      --text: #111827;
      --muted: #4b5563;
      --border: #d1d5db;
      --surface: #f8fafc;
      --ok: #15803d;
      --warn: #b45309;
      --fail: #b91c1c;
      --unknown: #475569;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      color: var(--text);
      font-family: "Segoe UI", Arial, sans-serif;
      background: white;
      font-size: 12px;
    }}
    .page {{
      width: 210mm;
      margin: 0 auto;
      padding: 16mm;
    }}
    .report-title {{
      border-bottom: 2px solid var(--text);
      padding-bottom: 10px;
      margin-bottom: 18px;
      display: flex;
      justify-content: space-between;
      gap: 20px;
      align-items: flex-end;
    }}
    h1 {{ font-size: 24px; margin: 0; }}
    h2 {{ font-size: 18px; margin: 0 0 4px; }}
    h3 {{ font-size: 14px; margin: 18px 0 8px; }}
    p {{ margin: 0; color: var(--muted); }}
    .audit {{ page-break-after: always; margin-bottom: 24px; }}
    .audit:last-child {{ page-break-after: auto; }}
    .audit-header, .summary, .meta {{ display: flex; gap: 8px; }}
    .audit-header {{ justify-content: space-between; align-items: flex-start; margin-bottom: 12px; }}
    .summary {{ flex-wrap: wrap; justify-content: flex-end; }}
    .meta {{
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      border: 1px solid var(--border);
      background: var(--surface);
      margin: 12px 0 14px;
    }}
    .meta div {{ padding: 8px; border-right: 1px solid var(--border); }}
    .meta div:last-child {{ border-right: 0; }}
    table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
    th, td {{ border: 1px solid var(--border); padding: 6px 8px; vertical-align: top; word-wrap: break-word; }}
    thead th {{ background: var(--surface); text-align: left; }}
    .kv th {{ width: 32%; text-align: left; background: var(--surface); }}
    .badge {{
      display: inline-block;
      min-width: 54px;
      padding: 3px 7px;
      border-radius: 4px;
      color: white;
      text-align: center;
      font-weight: 700;
      font-size: 11px;
    }}
    .ok {{ background: var(--ok); }}
    .warn {{ background: var(--warn); }}
    .fail {{ background: var(--fail); }}
    .unknown {{ background: var(--unknown); }}
    .actions {{ position: sticky; top: 0; padding: 8px 0; background: white; text-align: right; }}
    button {{ padding: 7px 12px; font: inherit; }}
    @media print {{
      .actions {{ display: none; }}
      body {{ background: white; }}
      .page {{ width: auto; margin: 0; padding: 0; }}
    }}
  </style>
</head>
<body>
  <div class="page">
    <div class="actions"><button onclick="window.print()">Print / Save PDF</button></div>
    <header class="report-title">
      <div>
        <h1>PC Audit Report</h1>
        <p>Generated {html_escape(generated)}</p>
      </div>
      <p>{len(audits)} audit record(s)</p>
    </header>
    {''.join(sections)}
  </div>
</body>
</html>
"""


def print_audit_archive_html():
    if not require_login():
        return
    ids = audit_ids_for_report()
    if not ids:
        messagebox.showwarning("Print Report", "Select an audit or refresh the archive first")
        return

    audits = []
    try:
        for audit_id in ids:
            audit = fetch_audit_detail(audit_id)
            if audit:
                audits.append(audit)
    except RuntimeError as exc:
        messagebox.showerror("Print Report", f"Could not load audit details: {exc}")
        return

    if not audits:
        messagebox.showwarning("Print Report", "No audit details found")
        return

    folder = printable_report_folder()
    path = os.path.join(folder, f"pc_audit_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html")
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(build_audit_html_report(audits))

    try:
        os.startfile(path)
    except Exception as exc:
        messagebox.showerror("Print Report", f"Report saved but could not be opened:\n{path}\n\n{exc}")
        return

    log_action("HTML_REPORT", f"Created printable audit report {path}")


def browse_manual_test_path(target_var):
    path = filedialog.askdirectory(title="Select drive or mounted media")
    if path:
        target_var.set(path)


def detect_windows_test_paths():
    if not platform.system().lower().startswith("win"):
        messagebox.showwarning("Manual Tests", "Windows device tests can only auto-detect Windows drive paths")
        return

    try:
        from register_client import test_path_candidates

        optical_candidates = test_path_candidates("optical")
        removable_candidates = test_path_candidates("removable")
    except Exception as exc:
        messagebox.showerror("Manual Tests", f"Could not detect paths: {exc}")
        return

    if optical_candidates and not manual_windows_optical_path_var.get().strip():
        manual_windows_optical_path_var.set(optical_candidates[0])
    if removable_candidates and not manual_windows_usb_path_var.get().strip():
        manual_windows_usb_path_var.set(removable_candidates[0])

    manual_windows_test_status_var.set(
        f"Detected optical: {optical_candidates[0] if optical_candidates else 'none'} | "
        f"USB/SD: {removable_candidates[0] if removable_candidates else 'none'}"
    )


def parse_usb_port_count():
    value = manual_windows_usb_expected_var.get().strip()
    if not value:
        return -1
    try:
        count = int(value)
    except ValueError:
        messagebox.showerror("Manual Tests", "USB ports must be a number")
        return None
    if count < 0:
        messagebox.showerror("Manual Tests", "USB ports must be 0 or greater")
        return None
    return count


def test_item(name, status, details=""):
    return {"name": name, "status": status, "details": details}


def usb_device_snapshot():
    try:
        from register_client import usb_device_listing

        return set(usb_device_listing())
    except Exception:
        return set()


def wait_for_usb_change(baseline, seconds=3.0):
    latest = set(baseline)
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        remaining = max(0.0, deadline - time.monotonic())
        manual_windows_test_status_var.set(f"Waiting for Windows USB detection... {remaining:.1f}s")
        root.update_idletasks()
        root.update()
        latest = usb_device_snapshot()
        if latest != baseline:
            break
        time.sleep(0.35)
    return latest


def run_fast_usb_detection_test(expected_count):
    try:
        from register_client import pc_auditor_usb_detection_test

        return pc_auditor_usb_detection_test(
            expected_usb_devices=expected_count,
            non_interactive=True,
        )
    except Exception as exc:
        return test_item("Test USB device detection", "UNKNOWN", f"USB detection failed: {exc}")


def run_usb_ports_one_by_one_gui(port_count):
    messagebox.showinfo(
        "USB one-by-one test",
        "Use one known-good USB mouse or flash drive.\n\n"
        "The app will test each USB port separately. Remove the USB device before every port baseline.",
    )

    items = []
    passed = 0
    failed = 0
    skipped = 0

    for port_index in range(1, port_count + 1):
        name = f"Test USB port {port_index}"
        manual_windows_test_status_var.set(f"Testing USB port {port_index}/{port_count}")

        if not messagebox.askokcancel(
            name,
            f"Step 1/{port_count}\n\nRemove the USB test device from all USB ports, then press OK.",
        ):
            items.append(test_item(name, "UNKNOWN", "Skipped before baseline"))
            skipped += 1
            continue

        baseline = usb_device_snapshot()

        if not messagebox.askokcancel(
            name,
            f"Plug the known-good USB device into USB port {port_index}.\n\n"
            "Press OK after Windows detects it, or immediately after plugging it in.",
        ):
            items.append(test_item(name, "UNKNOWN", "Skipped before device insertion"))
            skipped += 1
            continue

        manual_windows_test_status_var.set(f"Checking USB port {port_index}...")
        after = usb_device_snapshot()
        if after == baseline:
            after = wait_for_usb_change(baseline)
        added = sorted(after - baseline)
        removed = sorted(baseline - after)

        if added or len(after) > len(baseline):
            detail_parts = [f"Windows detected USB change on port {port_index}"]
            if added:
                detail_parts.append("added: " + "; ".join(added[:5]))
            if removed:
                detail_parts.append("removed: " + "; ".join(removed[:3]))
            items.append(test_item(name, "OK", " | ".join(detail_parts)))
            passed += 1
            continue

        answer = messagebox.askyesnocancel(
            name,
            "Windows did not report a new USB device automatically.\n\n"
            "Did the USB device work physically in this port?",
        )
        if answer is True:
            items.append(test_item(name, "OK", "Technician confirmed device worked; Windows USB list did not change"))
            passed += 1
        elif answer is False:
            items.append(test_item(name, "FAIL", "No USB detection and technician reported the port did not work"))
            failed += 1
        else:
            items.append(test_item(name, "UNKNOWN", "No USB detection; technician skipped manual confirmation"))
            skipped += 1

    if failed:
        summary_status = "FAIL"
    elif skipped:
        summary_status = "WARN"
    else:
        summary_status = "OK"

    items.insert(
        0,
        test_item(
            "Test USB ports one by one",
            summary_status,
            f"Ports tested: {port_count} | OK: {passed} | FAIL: {failed} | UNKNOWN: {skipped}",
        ),
    )
    return items


def post_client_registration(server_url, payload, timeout=8):
    url = server_url.rstrip("/") + "/clients/register"
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urlrequest.Request(
        url,
        data=body,
        method="POST",
        headers={"Content-Type": "application/json; charset=utf-8", "Accept": "application/json"},
    )

    try:
        with urlrequest.urlopen(req, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urlerror.HTTPError as exc:
        raw = exc.read().decode("utf-8", errors="replace")
        try:
            payload = json.loads(raw)
            message = payload.get("error", raw)
        except json.JSONDecodeError:
            message = raw or str(exc)
        raise RuntimeError(message) from exc
    except urlerror.URLError as exc:
        raise RuntimeError(str(exc.reason)) from exc
    except TimeoutError as exc:
        raise RuntimeError("Network request timed out") from exc


def run_windows_sound_test_gui():
    try:
        import tempfile
        from register_client import play_wav_file, status_item, write_stereo_channel_wav

        fd, path = tempfile.mkstemp(prefix="manual_windows_sound_", suffix=".wav")
        os.close(fd)
        try:
            write_stereo_channel_wav(path)
            played, playback_details = play_wav_file(path)
        finally:
            try:
                os.remove(path)
            except OSError:
                pass
    except Exception as exc:
        return {"name": "Test sound channels", "status": "UNKNOWN", "details": f"Could not play sound test: {exc}"}

    if not played:
        return {"name": "Test sound channels", "status": "UNKNOWN", "details": playback_details}

    answer = messagebox.askyesnocancel(
        "Sound left/right channel test",
        "Did you hear the sound correctly?\n\nFirst LEFT channel, then RIGHT channel.",
    )
    if answer is True:
        status = "OK"
        details = "User confirmed left/right sound output"
    elif answer is False:
        status = "FAIL"
        details = "User reported a left/right sound problem"
    else:
        status = "UNKNOWN"
        details = "User skipped sound confirmation"
    return status_item("Test sound channels", status, f"{details} | {playback_details}")


def summarize_test_items(items):
    summary = {"OK": 0, "WARN": 0, "FAIL": 0, "UNKNOWN": 0}
    for item in items:
        status = item.get("status", "UNKNOWN")
        summary[status] = summary.get(status, 0) + 1
    return summary


def set_windows_test_controls_running(running):
    state = "disabled" if running else "normal"
    if "manual_windows_run_btn" in globals():
        manual_windows_run_btn.configure(state=state)


def run_windows_device_tests_from_tab():
    if not require_login():
        return

    if not platform.system().lower().startswith("win"):
        messagebox.showwarning("Manual Tests", "Windows device tests are Windows-only")
        return

    selected_tests = {
        "usb_ports": manual_windows_usb_ports_var.get(),
        "fast_usb": manual_windows_fast_usb_var.get(),
        "optical": manual_windows_optical_var.get(),
        "removable": manual_windows_removable_var.get(),
        "sound": manual_windows_sound_var.get(),
        "upload": manual_windows_upload_var.get(),
    }
    if not any(selected_tests[key] for key in ("usb_ports", "optical", "removable", "sound")):
        messagebox.showerror("Manual Tests", "Select at least one Windows device test")
        return

    if selected_tests["upload"] and not save_network_settings(show_message=False):
        return

    usb_port_count = None
    if selected_tests["usb_ports"]:
        usb_port_count = parse_usb_port_count()
        if usb_port_count is None:
            return

    server_url = network_base_url()
    client_name = workstation_name_var.get().strip()
    optical_path = manual_windows_optical_path_var.get().strip()
    usb_path = manual_windows_usb_path_var.get().strip()

    set_windows_test_controls_running(True)
    items = []
    if selected_tests["usb_ports"]:
        if selected_tests["fast_usb"]:
            manual_windows_test_status_var.set("Running fast USB detection test...")
            items.append(run_fast_usb_detection_test(usb_port_count))
        else:
            if usb_port_count <= 0:
                messagebox.showerror("Manual Tests", "USB ports must be at least 1 when Fast USB is off")
                set_windows_test_controls_running(False)
                return
            manual_windows_test_status_var.set("Starting USB port one-by-one test...")
            items.extend(run_usb_ports_one_by_one_gui(usb_port_count))
    manual_windows_test_status_var.set("Running optical and removable media tests...")

    def worker():
        try:
            from types import SimpleNamespace
            from register_client import (
                pc_auditor_optical_read_test,
                pc_auditor_removable_read_test,
            )

            args = SimpleNamespace(
                non_interactive_tests=True,
                optical_test_path=optical_path,
                usb_test_path=usb_path,
            )
            worker_items = list(items)
            if selected_tests["optical"]:
                worker_items.append(pc_auditor_optical_read_test(args))
            if selected_tests["removable"]:
                worker_items.append(pc_auditor_removable_read_test(args))
            root.after(
                0,
                lambda: finish_windows_device_tests_from_tab(
                    worker_items,
                    server_url,
                    client_name,
                    run_sound=selected_tests["sound"],
                    upload=selected_tests["upload"],
                ),
            )
        except Exception as exc:
            root.after(0, lambda exc=exc: handle_windows_device_tests_error(exc))

    threading.Thread(target=worker, daemon=True).start()


def finish_windows_device_tests_from_tab(items, server_url, client_name, run_sound=True, upload=True):
    if run_sound:
        manual_windows_test_status_var.set("Playing left/right sound test...")
        items.append(run_windows_sound_test_gui())

    diagnostics = {"items": items, "summary": summarize_test_items(items)}
    inventory = collect_system_inventory()
    inventory["Diagnostics"] = diagnostics
    payload = {
        "name": client_name,
        "username": current_user or "",
        "role": current_role or "",
        "app_version": APP_VERSION,
        "platform": client_platform_label(),
        "inventory": inventory,
    }

    if not upload:
        complete_windows_device_tests_local(payload)
        return

    manual_windows_test_status_var.set("Uploading Windows test results...")

    def upload_worker():
        try:
            result = post_client_registration(server_url, payload, timeout=8)
            root.after(0, lambda: complete_windows_device_tests_upload(payload, result))
        except Exception as exc:
            root.after(0, lambda exc=exc: complete_windows_device_tests_local(payload, upload_error=exc))

    threading.Thread(target=upload_worker, daemon=True).start()


def apply_windows_device_tests_payload(payload):
    global network_clients

    client_name = payload.get("name", "")
    if client_name:
        network_clients[client_name] = {
            "name": client_name,
            "username": payload.get("username", ""),
            "role": payload.get("role", ""),
            "app_version": payload.get("app_version", ""),
            "platform": payload.get("platform", ""),
            "inventory": payload.get("inventory", {}),
            "last_seen": result.get("last_seen", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        }

    populate_device_results()
    populate_manual_tests()
    set_windows_test_controls_running(False)
    return payload.get("inventory", {}).get("Diagnostics", {})


def complete_windows_device_tests_upload(payload, result):
    diagnostics = apply_windows_device_tests_payload(payload)
    diagnostics = payload.get("inventory", {}).get("Diagnostics", {})
    summary = diagnostics.get("summary", {}) if isinstance(diagnostics, dict) else {}
    manual_windows_test_status_var.set(
        f"Windows tests saved | OK: {summary.get('OK', 0)} | WARN: {summary.get('WARN', 0)} | "
        f"FAIL: {summary.get('FAIL', 0)} | UNKNOWN: {summary.get('UNKNOWN', 0)}"
    )
    log_action("WINDOWS_DEVICE_TESTS", f"Saved Windows device tests for {client_name}")
    messagebox.showinfo("Manual Tests", "Windows device tests completed and saved")


def complete_windows_device_tests_local(payload, upload_error=None):
    diagnostics = apply_windows_device_tests_payload(payload)
    summary = diagnostics.get("summary", {}) if isinstance(diagnostics, dict) else {}
    suffix = f" | upload failed: {upload_error}" if upload_error else " | local only"
    manual_windows_test_status_var.set(
        f"Windows tests local | OK: {summary.get('OK', 0)} | WARN: {summary.get('WARN', 0)} | "
        f"FAIL: {summary.get('FAIL', 0)} | UNKNOWN: {summary.get('UNKNOWN', 0)}{suffix}"
    )
    client_name = payload.get("name", "")
    log_action("WINDOWS_DEVICE_TESTS_LOCAL", f"Saved local Windows device tests for {client_name}{suffix}")
    if upload_error:
        messagebox.showwarning("Manual Tests", f"Windows tests completed locally, but upload failed:\n{upload_error}")
    else:
        messagebox.showinfo("Manual Tests", "Windows device tests completed locally")


def handle_windows_device_tests_error(exc):
    set_windows_test_controls_running(False)
    manual_windows_test_status_var.set(f"Windows tests failed: {exc}")
    messagebox.showerror("Manual Tests", f"Windows device tests failed: {exc}")


def manual_rows_from_clients():
    rows = []
    for row in diagnostic_rows_from_clients():
        device = str(row.get("device", ""))
        details = str(row.get("details", ""))
        if device == "Diagnostics" and details == "No diagnostics reported by this client":
            continue
        rows.append(row)
    return rows


def set_manual_details(text):
    if "manual_details_text" not in globals():
        return
    manual_details_text.configure(state="normal")
    manual_details_text.delete("1.0", tk.END)
    manual_details_text.insert("1.0", text)
    manual_details_text.configure(state="disabled")


def populate_manual_tests():
    if "manual_tests_tree" not in globals():
        return

    clear_tree(manual_tests_tree)
    status_filter = manual_status_filter_var.get().strip()
    rows = manual_rows_from_clients()
    shown = 0
    counts = {"OK": 0, "WARN": 0, "FAIL": 0, "UNKNOWN": 0}

    for row in rows:
        status = row["status"] or "UNKNOWN"
        counts[status] = counts.get(status, 0) + 1

        if status_filter != "All" and status != status_filter:
            continue

        manual_tests_tree.insert(
            "",
            "end",
            values=(row["client"], row["device"], status, row["details"], row["last_seen"]),
        )
        shown += 1

    manual_summary_var.set(
        f"Test results: {shown}/{len(rows)} | OK: {counts.get('OK', 0)} | "
        f"WARN: {counts.get('WARN', 0)} | FAIL: {counts.get('FAIL', 0)} | "
        f"UNKNOWN: {counts.get('UNKNOWN', 0)}"
    )


def refresh_manual_tests():
    fetch_network_clients(show_message=False)
    load_manual_requirement(show_message=False)
    populate_manual_tests()


def load_manual_requirement(show_message=True):
    if not require_login():
        return

    if not save_network_settings(show_message=False):
        return

    try:
        result = network_request("GET", "/settings", timeout=8)
    except RuntimeError as exc:
        if show_message:
            messagebox.showerror("Manual Tests", f"Could not load requirement: {exc}")
        return

    settings = result.get("settings", {})
    manual_required_var.set(bool(settings.get("manual_tests_required", False)))

    if show_message:
        state = "required" if manual_required_var.get() else "optional"
        messagebox.showinfo("Manual Tests", f"Manual tests are now {state}")


def save_manual_requirement():
    if not require_admin():
        return

    if not save_network_settings(show_message=False):
        return

    required = bool(manual_required_var.get())

    try:
        network_request("POST", "/settings", data={"manual_tests_required": required}, timeout=8)
    except RuntimeError as exc:
        messagebox.showerror("Manual Tests", f"Could not save requirement: {exc}")
        return

    log_action("MANUAL_REQUIREMENT", f"Manual tests required={required}")
    state = "required" if required else "optional"
    messagebox.showinfo("Manual Tests", f"Manual tests are now {state}")


def on_manual_test_select(event=None):
    item = manual_tests_tree.focus()
    if not item:
        return
    values = manual_tests_tree.item(item).get("values", [])
    if not values:
        return

    client_name, test_name, status, note, last_seen = values
    detail_payload = {
        "client": client_name,
        "test": test_name,
        "status": status,
        "note": note,
        "last_seen": last_seen,
    }
    set_manual_details(json.dumps(detail_payload, ensure_ascii=False, indent=2))


def populate_device_results():
    if "device_results_tree" not in globals():
        return

    clear_tree(device_results_tree)
    status_filter = device_status_filter_var.get().strip()
    rows = diagnostic_rows_from_clients()
    shown = 0
    counts = {"OK": 0, "WARN": 0, "FAIL": 0, "UNKNOWN": 0}

    for row in rows:
        status = row["status"] or "UNKNOWN"
        counts[status] = counts.get(status, 0) + 1

        if status_filter != "All" and status != status_filter:
            continue

        device_results_tree.insert(
            "",
            "end",
            values=(row["client"], row["device"], status, row["details"], row["last_seen"]),
        )
        shown += 1

    total_clients = len(network_clients)
    device_summary_var.set(
        f"Clients: {total_clients} | Results: {shown}/{len(rows)} | OK: {counts.get('OK', 0)} | "
        f"WARN: {counts.get('WARN', 0)} | FAIL: {counts.get('FAIL', 0)} | UNKNOWN: {counts.get('UNKNOWN', 0)}"
    )


def refresh_device_results():
    fetch_network_clients(show_message=False)
    populate_device_results()


def on_device_result_select(event=None):
    item = device_results_tree.focus()
    if not item:
        return
    values = device_results_tree.item(item).get("values", [])
    if not values:
        return

    client_name, device, status, details, last_seen = values
    client = network_clients.get(client_name, {})
    inventory = client.get("inventory", {}) if isinstance(client, dict) else {}
    diagnostics = inventory.get("Diagnostics", {}) if isinstance(inventory, dict) else {}

    detail_payload = {
        "client": client_name,
        "device": device,
        "status": status,
        "details": details,
        "last_seen": last_seen,
        "diagnostics_summary": diagnostics.get("summary", {}) if isinstance(diagnostics, dict) else {},
    }
    set_device_details(json.dumps(detail_payload, ensure_ascii=False, indent=2))


def on_network_client_select(event=None):
    item = network_client_tree.focus()
    if not item:
        return

    values = network_client_tree.item(item).get("values", [])
    if not values:
        return

    client_name = values[0]
    if client_name:
        job_assignee_var.set(client_name)
        client = network_clients.get(client_name, {})
        set_network_details(json.dumps(client, ensure_ascii=False, indent=2))


def selected_network_client():
    item = network_client_tree.focus()
    if not item:
        return None
    values = network_client_tree.item(item).get("values", [])
    if not values:
        return None
    return network_clients.get(values[0])


def load_selected_client_info_to_form():
    client = selected_network_client()
    if not client:
        messagebox.showwarning("Network Jobs", "Select a client first")
        return

    inventory = client.get("inventory", {})
    if not inventory:
        messagebox.showwarning("Network Jobs", "Selected client has not sent inventory details yet")
        return

    clear_form()
    for index, field_name in enumerate(AUDITOR_FIELDS):
        value = inventory.get(field_name, "")
        if value:
            set_field(entries[index], value)

    serial = inventory.get("Serial number", "")
    if serial:
        record_xenoglossi_entry.insert(0, serial)
        record_username_entry.insert(0, serial)
    else:
        record_xenoglossi_entry.insert(0, client.get("name", ""))
        record_username_entry.insert(0, client.get("name", ""))

    notebook.select(form_frame)


def collect_current_form_payload():
    values = [get_field(entry) for entry in entries]
    return {
        "xenoglossi": record_xenoglossi_entry.get().strip(),
        "record_username": record_username_entry.get().strip(),
        "dropdown": dropdown.get().strip(),
        "checkbox": bool(checkbox_var.get()),
        "inputs": values,
        "fields": {field_name: values[index] for index, field_name in enumerate(AUDITOR_FIELDS)},
    }


def load_payload_to_form(payload):
    clear_form()

    xenoglossi = str(payload.get("xenoglossi", "")).strip()
    record_username = str(payload.get("record_username", "")).strip()
    dropdown_value = str(payload.get("dropdown", "")).strip()
    inputs = payload.get("inputs")

    if not isinstance(inputs, list):
        fields = payload.get("fields", {})
        inputs = [fields.get(field_name, "") for field_name in AUDITOR_FIELDS] if isinstance(fields, dict) else []
    inputs = normalize_input_values(inputs)

    if xenoglossi:
        record_xenoglossi_entry.insert(0, xenoglossi)
    if record_username:
        record_username_entry.insert(0, record_username)
    if dropdown_value:
        dropdown.set(dropdown_value)

    checkbox_var.set(bool(payload.get("checkbox", False)))

    for index, value in enumerate(inputs):
        set_field(entries[index], value)

    notebook.select(form_frame)


def selected_network_job():
    item = network_tree.focus()
    if not item:
        return None
    values = network_tree.item(item).get("values", [])
    if not values:
        return None
    return network_jobs.get(str(values[0]))


def show_selected_network_job_details(event=None):
    job = selected_network_job()
    if not job:
        set_network_details("")
        return

    set_network_details(json.dumps(job, ensure_ascii=False, indent=2))


def fetch_network_jobs(show_message=True):
    global network_jobs

    if not require_login():
        return

    if not save_network_settings(show_message=False):
        return

    register_this_client(show_message=False)

    assignee = workstation_name_var.get().strip()
    status = job_status_filter_var.get().strip() or "Queued"

    try:
        result = network_request(
            "GET",
            "/jobs",
            params={"assignee": assignee, "status": status},
            auth=True,
            timeout=8,
        )
    except RuntimeError as exc:
        if show_message:
            messagebox.showerror("Network Jobs", f"Could not fetch jobs: {exc}")
        else:
            set_network_details(f"Network unavailable: {exc}")
        return

    jobs = result.get("jobs", [])
    network_jobs = {str(job["id"]): job for job in jobs}
    clear_tree(network_tree)

    for job in jobs:
        network_tree.insert(
            "",
            "end",
            values=(
                job.get("id"),
                job.get("title", ""),
                job.get("assignee", ""),
                job.get("status", ""),
                job.get("created_by", ""),
                job.get("created_at", ""),
                job.get("updated_at", ""),
            ),
        )

    set_network_details(f"Fetched {len(jobs)} job(s) for {assignee}.")

    if show_message:
        messagebox.showinfo("Network Jobs", f"Fetched {len(jobs)} job(s)")


def publish_current_form_job():
    if not require_admin():
        return

    if not save_network_settings(show_message=False):
        return

    payload = collect_current_form_payload()
    xenoglossi = payload["xenoglossi"]
    filled_values = [value for value in payload["inputs"] if value]

    if not xenoglossi:
        messagebox.showerror("Network Jobs", "Asset Tag is required before publishing")
        return

    if not filled_values:
        messagebox.showerror("Network Jobs", "Fill at least one PC Auditor field before publishing")
        return

    title = job_title_var.get().strip() or f"PC Auditor - {xenoglossi}"
    assignee = job_assignee_var.get().strip() or "ALL"

    try:
        result = network_request(
            "POST",
            "/jobs",
            data={
                "title": title,
                "assignee": assignee,
                "created_by": current_user,
                "payload": payload,
            },
            auth=True,
            timeout=8,
        )
    except RuntimeError as exc:
        messagebox.showerror("Network Jobs", f"Could not publish job: {exc}")
        return

    job_id = result.get("job_id")
    log_action("NETWORK_JOB_CREATE", f"Published job #{job_id} for {assignee}")
    messagebox.showinfo("Network Jobs", f"Job #{job_id} published")
    fetch_network_clients(show_message=False)
    fetch_network_jobs(show_message=False)


def update_network_job_status(job_id, status, notes=""):
    result = network_request(
        "POST",
        f"/jobs/{job_id}/status",
        data={"status": status, "actor": current_user or workstation_name_var.get().strip(), "notes": notes},
        auth=True,
        timeout=8,
    )
    log_action("NETWORK_JOB_STATUS", f"Job #{job_id} -> {status}")
    return result


def load_selected_network_job(mark_loaded=True):
    job = selected_network_job()
    if not job:
        messagebox.showwarning("Network Jobs", "Select a job first")
        return None

    load_payload_to_form(job.get("payload", {}))

    if mark_loaded and job.get("status") == "Queued":
        try:
            update_network_job_status(job["id"], "Loaded", f"Loaded on {workstation_name_var.get().strip()}")
        except RuntimeError as exc:
            messagebox.showwarning("Network Jobs", f"Loaded locally, but status update failed: {exc}")

    fetch_network_jobs(show_message=False)
    return job


def apply_selected_network_job():
    job = load_selected_network_job(mark_loaded=True)
    if not job:
        return
    apply_keyboard_actions()


def complete_selected_network_job():
    job = selected_network_job()
    if not job:
        messagebox.showwarning("Network Jobs", "Select a job first")
        return

    try:
        update_network_job_status(job["id"], "Completed", f"Completed on {workstation_name_var.get().strip()}")
    except RuntimeError as exc:
        messagebox.showerror("Network Jobs", f"Could not complete job: {exc}")
        return

    messagebox.showinfo("Network Jobs", "Job marked completed")
    fetch_network_jobs(show_message=False)


def fail_selected_network_job():
    job = selected_network_job()
    if not job:
        messagebox.showwarning("Network Jobs", "Select a job first")
        return

    try:
        update_network_job_status(job["id"], "Failed", f"Failed on {workstation_name_var.get().strip()}")
    except RuntimeError as exc:
        messagebox.showerror("Network Jobs", f"Could not mark job failed: {exc}")
        return

    messagebox.showinfo("Network Jobs", "Job marked failed")
    fetch_network_jobs(show_message=False)


network_tree.bind("<<TreeviewSelect>>", show_selected_network_job_details)
network_client_tree.bind("<<TreeviewSelect>>", on_network_client_select)
device_results_tree.bind("<<TreeviewSelect>>", on_device_result_select)
audit_archive_tree.bind("<<TreeviewSelect>>", on_audit_archive_select)
manual_tests_tree.bind("<<TreeviewSelect>>", on_manual_test_select)


# =========================
# USER FUNCTIONS
# =========================
def clear_user_form():
    global selected_user_id
    selected_user_id = None
    user_username_var.set("")
    user_password_entry.delete(0, tk.END)
    user_role_var.set("User")
    user_active_var.set(True)
    if user_tree.selection():
        user_tree.selection_remove(user_tree.selection())


def load_users():
    if "user_tree" not in globals():
        return

    clear_tree(user_tree)
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT id, username, role, status, created_at, last_login FROM app_users ORDER BY username")
    for row in c.fetchall():
        user_tree.insert("", "end", values=(row[0], row[1], row[2], row[3], row[4] or "", row[5] or ""))
    conn.close()


def on_user_select(event):
    global selected_user_id

    item = user_tree.focus()
    if not item:
        return

    values = user_tree.item(item)["values"]
    selected_user_id = values[0]
    user_username_var.set(values[1])
    user_role_var.set(values[2])
    user_active_var.set(values[3] == "Active")
    user_password_entry.delete(0, tk.END)


user_tree.bind("<<TreeviewSelect>>", on_user_select)


def add_user():
    if not require_admin():
        return

    username = user_username_var.get().strip()
    password = user_password_entry.get().strip()
    role = user_role_var.get().strip()
    status = "Active" if user_active_var.get() else "Disabled"

    if not username or not password:
        messagebox.showerror("Users", "Username and password are required")
        return

    try:
        conn = get_conn()
        c = conn.cursor()
        c.execute(
            "INSERT INTO app_users (username, password, role, status, created_at) VALUES (?, ?, ?, ?, ?)",
            (username, hash_password(password), role, status, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
        conn.commit()
        conn.close()
    except sqlite3.IntegrityError:
        messagebox.showerror("Users", "Username already exists")
        return

    log_action("USER_CREATE", f"Created user {username}")
    load_users()
    refresh_analytics()
    clear_user_form()
    messagebox.showinfo("Users", "User created")


def update_user():
    if not require_admin():
        return

    if not selected_user_id:
        messagebox.showwarning("Users", "Please select a user first")
        return

    username = user_username_var.get().strip()
    role = user_role_var.get().strip()
    status = "Active" if user_active_var.get() else "Disabled"
    password = user_password_entry.get().strip()

    if not username:
        messagebox.showerror("Users", "Username is required")
        return

    if username == current_user and status != "Active":
        messagebox.showwarning("Users", "You cannot disable the current account")
        return

    try:
        conn = get_conn()
        c = conn.cursor()
        if password:
            c.execute(
                "UPDATE app_users SET username=?, password=?, role=?, status=? WHERE id=?",
                (username, hash_password(password), role, status, selected_user_id),
            )
        else:
            c.execute(
                "UPDATE app_users SET username=?, role=?, status=? WHERE id=?",
                (username, role, status, selected_user_id),
            )
        conn.commit()
        conn.close()
    except sqlite3.IntegrityError:
        messagebox.showerror("Users", "Username already exists")
        return

    log_action("USER_UPDATE", f"Updated user {username}")
    load_users()
    refresh_analytics()
    clear_user_form()
    messagebox.showinfo("Users", "User updated")


def delete_user():
    if not require_admin():
        return

    if not selected_user_id:
        messagebox.showwarning("Users", "Please select a user first")
        return

    username = user_username_var.get().strip()
    if username == current_user:
        messagebox.showwarning("Users", "You cannot delete the current account")
        return

    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM app_users WHERE role='Admin' AND status='Active' AND id<>?", (selected_user_id,))
    remaining_admins = c.fetchone()[0]
    if remaining_admins == 0:
        conn.close()
        messagebox.showwarning("Users", "At least one active admin account is required")
        return

    confirm = messagebox.askyesno("Users", f"Delete user {username}?")
    if not confirm:
        conn.close()
        return

    c.execute("DELETE FROM app_users WHERE id=?", (selected_user_id,))
    conn.commit()
    conn.close()

    log_action("USER_DELETE", f"Deleted user {username}")
    load_users()
    refresh_analytics()
    clear_user_form()
    messagebox.showinfo("Users", "User deleted")


# =========================
# SETTINGS FUNCTIONS
# =========================
def load_settings_ui():
    if "setting_system_name_var" not in globals():
        return

    setting_system_name_var.set(get_setting("system_name", APP_TITLE))
    setting_export_folder_var.set(get_setting("default_export_folder", BASE_DIR))
    setting_keyboard_delay_var.set(get_setting("keyboard_delay", "80"))
    setting_language_var.set(language_label(get_setting("language", UI_LANGUAGE)))
    setting_auto_refresh_var.set(get_setting("auto_refresh", "1") == "1")
    load_batch_files()


def choose_export_folder():
    folder = filedialog.askdirectory(initialdir=get_default_export_folder(), title="Choose export folder")
    if folder:
        setting_export_folder_var.set(folder)


def license_status_label(value):
    labels = {
        "0": "Unlicensed",
        "1": "Licensed",
        "2": "OOB grace",
        "3": "OOT grace",
        "4": "Non-genuine grace",
        "5": "Notification",
        "6": "Extended grace",
    }
    return labels.get(str(value).strip(), str(value or "").strip())


def windows_activation_xpr_status():
    script = os.path.join(os.environ.get("SystemRoot", r"C:\Windows"), "System32", "slmgr.vbs")
    if not os.path.exists(script):
        return ""
    output = run_command(["cscript", "//nologo", script, "/xpr"], timeout=10)
    lower = output.lower()
    if "permanent" in lower or "activated" in lower:
        return "Licensed"
    if output.strip():
        return "Activated / see Windows activation details"
    return ""


def first_pipe_parts(value, count):
    line = first_line(value)
    parts = [part.strip() for part in line.split("|")] if line else []
    while len(parts) < count:
        parts.append("")
    return parts[:count]


def read_windows_license_info():
    if not platform.system().lower().startswith("win"):
        raise RuntimeError("Windows license reading is available only on Windows")

    product_name, edition_id, display_version, current_build = first_pipe_parts(
        powershell_value(
            "$cv=Get-ItemProperty 'HKLM:\\SOFTWARE\\Microsoft\\Windows NT\\CurrentVersion'; "
            "$cv.ProductName + '|' + $cv.EditionID + '|' + $cv.DisplayVersion + '|' + $cv.CurrentBuild"
        ),
        4,
    )
    caption = first_line(powershell_value("(Get-CimInstance Win32_OperatingSystem).Caption"))
    if caption:
        product_name = caption
    elif product_name.startswith("Microsoft Windows 10"):
        try:
            if int(current_build) >= 22000:
                product_name = product_name.replace("Windows 10", "Windows 11", 1)
        except ValueError:
            pass
    firmware_key = first_line(powershell_value("(Get-CimInstance SoftwareLicensingService).OA3xOriginalProductKey"))
    backup_key = first_line(
        powershell_value(
            "(Get-ItemProperty 'HKLM:\\SOFTWARE\\Microsoft\\Windows NT\\CurrentVersion\\SoftwareProtectionPlatform').BackupProductKeyDefault"
        )
    )
    product_key = firmware_key or backup_key or "Not available"

    license_query = (
        "$products=Get-CimInstance SoftwareLicensingProduct | "
        "Where-Object { $_.Name -like '*Windows*' -and $null -ne $_.LicenseStatus }; "
        "$p=$products | Sort-Object @{Expression={$_.LicenseStatus -eq 1};Descending=$true}, LicenseStatus -Descending | "
        "Select-Object -First 1; "
        "if ($p) { $p.Name + '|' + $p.Description + '|' + $p.LicenseStatus + '|' + $p.PartialProductKey }"
    )
    license_name, license_description, license_status, partial_key = first_pipe_parts(
        powershell_value(license_query),
        4,
    )
    if not license_status:
        license_name, license_description, license_status, partial_key = first_pipe_parts(
            powershell_value(
                "$p=Get-CimInstance SoftwareLicensingProduct | "
                "Where-Object { $_.ApplicationID -eq '55c92734-d682-4d71-983e-d6ec3f16059f' } | "
                "Sort-Object @{Expression={$_.LicenseStatus -eq 1};Descending=$true}, LicenseStatus -Descending | "
                "Select-Object -First 1; "
                "if ($p) { $p.Name + '|' + $p.Description + '|' + $p.LicenseStatus + '|' + $p.PartialProductKey }"
            ),
            4,
        )
    if not license_status:
        license_status = windows_activation_xpr_status()

    if license_status and license_status.isdigit():
        license_status = license_status_label(license_status)

    if not partial_key:
        partial_key = first_line(
            powershell_value(
                "$p=Get-CimInstance SoftwareLicensingProduct | Where-Object { $_.PartialProductKey } | "
                "Select-Object -First 1; if ($p) { $p.PartialProductKey }"
            )
        )

    old_license_name, old_license_description, old_license_status, old_partial_key = first_pipe_parts(
        powershell_value(
            "$p=Get-CimInstance SoftwareLicensingProduct | "
            "Where-Object { $_.ApplicationID -eq '55c92734-d682-4d71-983e-d6ec3f16059f' -and $_.PartialProductKey } | "
            "Select-Object -First 1; "
            "if ($p) { $p.Name + '|' + $p.Description + '|' + $p.LicenseStatus + '|' + $p.PartialProductKey }"
        ),
        4,
    )
    license_name = license_name or old_license_name
    license_description = license_description or old_license_description
    partial_key = partial_key or old_partial_key
    channel = ""
    description_upper = license_description.upper()
    for candidate in ("OEM", "RETAIL", "VOLUME", "MAK", "KMS"):
        if candidate in description_upper:
            channel = candidate
            break

    return {
        "edition": edition_id or caption,
        "product_name": product_name or caption or license_name,
        "display_version": display_version,
        "build": current_build,
        "product_key": product_key,
        "license_status": license_status or "Unknown",
        "partial_key": partial_key,
        "channel": channel or license_description or "Unknown",
    }


def read_windows_license_to_settings():
    try:
        info = read_windows_license_info()
    except Exception as exc:
        messagebox.showerror("Windows License", str(exc))
        return

    version_bits = []
    if info.get("display_version"):
        version_bits.append(info["display_version"])
    if info.get("build"):
        version_bits.append("build " + info["build"])
    edition = info.get("edition", "")
    if version_bits:
        edition = f"{edition} ({', '.join(version_bits)})"

    status = info.get("license_status", "")
    if info.get("partial_key"):
        status = f"{status} | partial key {info['partial_key']}"

    windows_license_edition_var.set(edition or "Unknown")
    windows_license_product_var.set(info.get("product_name", "") or "Unknown")
    windows_license_key_var.set(info.get("product_key", "") or "Not available")
    windows_license_status_var.set(status or "Unknown")
    windows_license_channel_var.set(info.get("channel", "") or "Unknown")
    log_action("WINDOWS_LICENSE_READ", f"Read Windows license info: {windows_license_product_var.get()}")


def save_settings():
    if not require_admin():
        return

    system_name = setting_system_name_var.get().strip() or APP_TITLE
    export_folder = setting_export_folder_var.get().strip() or BASE_DIR
    delay = setting_keyboard_delay_var.get().strip() or "80"
    language = language_code(setting_language_var.get())

    if not os.path.isdir(export_folder):
        messagebox.showerror("Settings", "Export folder does not exist")
        return

    try:
        delay_int = int(delay)
        if delay_int < 0 or delay_int > 5000:
            raise ValueError
    except ValueError:
        messagebox.showerror("Settings", "Keyboard delay must be a number from 0 to 5000")
        return

    set_setting("system_name", system_name)
    set_setting("default_export_folder", export_folder)
    set_setting("keyboard_delay", delay_int)
    set_setting("language", language)
    set_setting("auto_refresh", "1" if setting_auto_refresh_var.get() else "0")

    root.title(system_name)
    title_label.config(text=system_name)
    log_action("SETTINGS_SAVE", "Saved application settings")
    restart_note = "\n\nRestart the app to apply language changes." if language != UI_LANGUAGE else ""
    messagebox.showinfo("Settings", "Settings saved" + restart_note)


def reset_settings():
    if not require_admin():
        return

    confirm = messagebox.askyesno("Settings", "Reset settings to defaults?")
    if not confirm:
        return

    for key, value in DEFAULT_SETTINGS.items():
        set_setting(key, value)
    load_settings_ui()
    root.title(APP_TITLE)
    title_label.config(text=APP_TITLE)
    log_action("SETTINGS_RESET", "Reset application settings")
    messagebox.showinfo("Settings", "Defaults restored")


def available_batch_files():
    files = []
    try:
        names = os.listdir(BASE_DIR)
    except OSError:
        return files

    for name in sorted(names, key=str.lower):
        path = os.path.join(BASE_DIR, name)
        if os.path.isfile(path) and name.lower().endswith(".bat"):
            files.append(name)
    return files


def load_batch_files():
    if "batch_tree" not in globals():
        return

    clear_tree(batch_tree)
    files = available_batch_files()
    for name in files:
        path = os.path.join(BASE_DIR, name)
        try:
            stat = os.stat(path)
            modified = datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            size = f"{stat.st_size} bytes"
        except OSError:
            modified = ""
            size = ""
        batch_tree.insert("", "end", values=(name, modified, size))

    batch_status_var.set(f"{len(files)} batch file(s) found")


def selected_batch_file_names():
    names = []
    if "batch_tree" not in globals():
        return names
    for item in batch_tree.selection():
        values = batch_tree.item(item).get("values", [])
        if values:
            names.append(str(values[0]))
    return names


def resolve_batch_path(name):
    base = os.path.abspath(BASE_DIR)
    path = os.path.abspath(os.path.join(base, name))
    if not path.lower().endswith(".bat"):
        raise ValueError("Only .bat files can be launched")
    if os.path.commonpath([base, path]) != base:
        raise ValueError("Batch file must be inside the application folder")
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    return path


def launch_batch_file(name, env_updates=None):
    if not platform.system().lower().startswith("win"):
        raise RuntimeError("Batch launcher is available only on Windows")

    path = resolve_batch_path(name)
    cmd = os.environ.get("COMSPEC", "cmd.exe")
    creationflags = subprocess.CREATE_NO_WINDOW if hasattr(subprocess, "CREATE_NO_WINDOW") else 0
    env = os.environ.copy()
    env["RUN_BACKGROUND"] = "1"
    if env_updates:
        env.update({str(key): str(value) for key, value in env_updates.items()})
    process = subprocess.Popen(
        [cmd, "/c", path],
        cwd=BASE_DIR,
        creationflags=creationflags,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env=env,
    )
    return path, process.pid


def run_windows_device_tests_batch():
    if not require_login():
        return

    if not platform.system().lower().startswith("win"):
        messagebox.showwarning("Manual Tests", "Windows device tests are available only on Windows")
        return

    if not save_network_settings(show_message=False):
        return

    batch_name = "windows_device_tests.bat"
    env_updates = {
        "JOB_SERVER_URL": network_base_url(),
        "CLIENT_NAME": workstation_name_var.get().strip(),
    }
    try:
        path, pid = launch_batch_file(batch_name, env_updates=env_updates)
    except Exception as exc:
        messagebox.showerror("Manual Tests", f"Could not start {batch_name}:\n{exc}")
        return

    manual_windows_test_status_var.set(f"Started {os.path.basename(path)} in background pid={pid}")
    log_action("WINDOWS_DEVICE_TESTS_BATCH", f"Started {os.path.basename(path)} pid={pid}")
    messagebox.showinfo("Manual Tests", f"Started {os.path.basename(path)} in background")


def run_pc_audit_batch():
    global pc_audit_batch_started

    if not require_login():
        return

    batch_name = "start_pc_audit.bat"
    try:
        path, pid = launch_batch_file(batch_name)
    except Exception as exc:
        messagebox.showerror("PC Audit", f"Could not start {batch_name}:\n{exc}")
        return

    pc_audit_batch_started = True
    log_action("PC_AUDIT_BATCH_START", f"Started {os.path.basename(path)} pid={pid}")
    messagebox.showinfo("PC Audit", f"Started {os.path.basename(path)} in background")


def run_batch_file_names(names):
    if not names:
        messagebox.showwarning("Batch Launcher", "Select at least one batch file")
        return

    launched = []
    errors = []
    for name in names:
        try:
            path, pid = launch_batch_file(name)
            launched.append(f"{os.path.basename(path)} pid={pid}")
            time.sleep(0.2)
        except Exception as exc:
            errors.append(f"{name}: {exc}")

    if launched:
        batch_status_var.set(f"Started background: {', '.join(launched)}")
        log_action("BATCH_LAUNCH_BACKGROUND", f"Started background: {', '.join(launched)}")

    if errors:
        messagebox.showerror("Batch Launcher", "\n".join(errors[:8]))
    elif launched:
        messagebox.showinfo("Batch Launcher", f"Started {len(launched)} batch file(s) in background")


def run_selected_batch_files():
    if not require_admin():
        return
    run_batch_file_names(selected_batch_file_names())


def run_all_batch_files():
    if not require_admin():
        return

    names = available_batch_files()
    if not names:
        messagebox.showwarning("Batch Launcher", "No batch files found")
        return

    confirm = messagebox.askyesno(
        "Batch Launcher",
        f"This will start {len(names)} hidden background batch process(es).\n\nContinue?",
    )
    if not confirm:
        return

    run_batch_file_names(names)


def backup_database():
    if not require_admin():
        return

    if not os.path.exists(DB_NAME):
        messagebox.showwarning("Backup", "Database file does not exist yet")
        return

    path = ask_save_path(
        ".db",
        [("SQLite database", "*.db"), ("All files", "*.*")],
        "Backup Database",
        f"app_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db",
    )

    if path:
        shutil.copy2(DB_NAME, path)
        set_setting("last_backup", path)
        log_action("DB_BACKUP", f"Database backup saved to {path}")
        messagebox.showinfo("Backup", "Database backup completed")


# =========================
# LOG FUNCTIONS
# =========================
def get_log_rows():
    query = "SELECT id, timestamp, username, action, details FROM app_logs"
    clauses = []
    params = []

    user = log_user_var.get().strip()
    action = log_action_var.get().strip()

    if user:
        clauses.append("username LIKE ?")
        params.append(f"%{user}%")
    if action:
        clauses.append("action LIKE ?")
        params.append(f"%{action}%")

    if clauses:
        query += " WHERE " + " AND ".join(clauses)

    query += " ORDER BY id DESC"

    conn = get_conn()
    c = conn.cursor()
    c.execute(query, params)
    rows = c.fetchall()
    conn.close()
    return rows


def load_logs():
    if "log_tree" not in globals():
        return

    clear_tree(log_tree)
    for row in get_log_rows():
        log_tree.insert("", "end", values=row)


def export_logs_excel():
    if not require_admin():
        return

    rows = get_log_rows()
    if not rows:
        messagebox.showwarning("Logs", "No logs to export")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Logs"
    ws.append(["ID", "Timestamp", "User", "Action", "Details"])
    for row in rows:
        ws.append(list(row))

    path = ask_save_path(
        ".xlsx",
        [("Excel files", "*.xlsx")],
        "Export Logs",
        f"logs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
    )

    if path:
        wb.save(path)
        log_action("EXPORT_LOGS", f"Exported logs to {path}")
        load_logs()
        refresh_analytics()
        messagebox.showinfo("Logs", "Logs exported successfully")


def clear_logs():
    if not require_admin():
        return

    confirm = messagebox.askyesno("Logs", "Clear all logs?")
    if not confirm:
        return

    conn = get_conn()
    c = conn.cursor()
    c.execute("DELETE FROM app_logs")
    conn.commit()
    conn.close()

    log_action("LOGS_CLEAR", "Cleared audit log")
    load_logs()
    refresh_analytics()
    messagebox.showinfo("Logs", "Logs cleared")


# =========================
# KEYBOARD AUTOMATION
# =========================
def clipboard_text():
    try:
        return True, root.clipboard_get()
    except tk.TclError:
        return False, ""


def restore_clipboard(had_text, text):
    try:
        root.clipboard_clear()
        if had_text:
            root.clipboard_append(text)
        root.update()
    except tk.TclError:
        pass


def send_text_by_paste(value):
    try:
        root.clipboard_clear()
        root.clipboard_append(value)
        root.update()
        pyautogui.hotkey("ctrl", "v")
        time.sleep(max(0.04, pyautogui.PAUSE))
    except Exception:
        pyautogui.write(value, interval=0.03)


def apply_keyboard_actions(event=None):
    if not require_login():
        return

    if pyautogui is None:
        messagebox.showerror(
            "Missing library",
            "pyautogui is not installed. Run: pip install pyautogui",
        )
        return

    raw_values = [
        get_field(entry)
        for field_name, entry in zip(AUDITOR_FIELDS, entries)
        if field_name not in KEYBOARD_SKIP_FIELDS
    ]
    filled_values = [value for value in raw_values if value]
    values = raw_values if tab_empty_var.get() else filled_values

    if not filled_values:
        messagebox.showwarning("Apply", "There are no filled input fields")
        return

    messagebox.showinfo(
        "Apply",
        "Press OK, then click the first target field in PC Auditor within 2 seconds.",
    )

    root.iconify()
    time.sleep(2)

    pyautogui.FAILSAFE = True
    try:
        pyautogui.PAUSE = int(get_setting("keyboard_delay", "80")) / 1000
    except ValueError:
        pyautogui.PAUSE = 0.08

    had_clipboard_text, original_clipboard_text = clipboard_text()

    try:
        for value in values:
            if value:
                action = value.lower().replace(" ", "")

                if "+" in action:
                    keys = action.split("+")
                    pyautogui.hotkey(*keys)
                elif action in pyautogui.KEYBOARD_KEYS:
                    pyautogui.press(action)
                else:
                    send_text_by_paste(value)
            pyautogui.press("tab")
            time.sleep(max(0.04, pyautogui.PAUSE))
        log_action("APPLY_KEYS", f"Applied {len(values)} keyboard actions")
    except Exception as exc:
        messagebox.showerror("Apply error", str(exc))
    finally:
        restore_clipboard(had_clipboard_text, original_clipboard_text)
        root.deiconify()


root.bind("<F1>", apply_keyboard_actions)


# =========================
# NOTEBOOK EVENTS
# =========================
def on_tab_changed(event):
    selected = notebook.select()
    if selected == str(analytics_frame):
        refresh_analytics()
    elif selected == str(reports_frame):
        build_report(show_message=False)
    elif selected == str(network_frame):
        load_network_settings_ui()
    elif selected == str(device_results_frame):
        load_network_settings_ui()
        refresh_device_results()
    elif selected == str(audit_archive_frame):
        load_network_settings_ui()
        populate_audit_client_filter()
        fetch_audit_archive(show_message=False)
    elif selected == str(manual_tests_frame):
        load_network_settings_ui()
        refresh_manual_tests()
    elif selected == str(users_frame):
        load_users()
    elif selected == str(settings_frame):
        load_settings_ui()
    elif selected == str(logs_frame):
        load_logs()


notebook.bind("<<NotebookTabChanged>>", on_tab_changed)


# =========================
# START
# =========================
init_db()
load_settings_ui()
load_network_settings_ui()
root.title(get_setting("system_name", APP_TITLE))
title_label.config(text=get_setting("system_name", APP_TITLE))
load_data()
load_users()
load_logs()
refresh_analytics()
set_protected_tabs("hidden")
apply_ui_language(root)

root.mainloop()
